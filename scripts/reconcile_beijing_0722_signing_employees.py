#!/usr/bin/env python3
"""Generate a fail-closed production SQL import from the Beijing 0722 signing sheet."""

from __future__ import annotations

import argparse
import datetime as dt
import decimal
import hashlib
import os
import re
from pathlib import Path

import openpyxl


SHEET = "签约数据"
EXPECTED_ROWS = 27
OPERATOR = "employee_xlsx_0722"
LOCK_NAME = "erp:employee:beijing:0722"

REQUIRED_HEADERS = (
    "序号", "姓名", "在职状态", "员工状态", "1级部门\n（金英）",
    "2级部门\n（灵韵）", "3级部门", "4级门店", "3/4级部门负责人",
    "签约公司", "社保类型", "员工岗位", "城市等级", "法定代表人",
    "注册地", "合同类型", "员工职级", "工作地", "身份证号", "家庭住址",
    "联系电话", "入职时间", "劳动合同期限形式", "劳动合同起始日期",
    "劳动合同结束日期", "试用期开始日期", "试用期结束日期", "工时制度",
    "综合工资", "底薪", "综合岗位津贴", "综合驻外补贴", "月度绩效津贴",
    "房补", "备注",
)

POSTS = {"区域运营总监", "店长助理", "茶艺师", "驻店经理", "店长", "实习生"}

# The combined Conrad/Hilton source assignment has no standalone Hilton node in
# production.  Preserve the exact source label in store_name, use the regional
# operating department as primary, and retain the existing Conrad shop scope.
ORG = {
    "北京柏悦": (1176, "北京柏悦", ((1176, "Y"),)),
    "北京丽思卡尔顿": (1173, "北京丽思卡尔顿", ((1173, "Y"),)),
    "北京瑞吉": (1178, "北京瑞吉", ((1178, "Y"),)),
    "华彬费尔蒙": (1263, "费尔蒙", ((1263, "Y"),)),
    "通州皇冠": (1183, "通州皇冠", ((1183, "Y"),)),
    "北京康莱德+北京希尔顿逸林": (
        1156, "北京区域运营", ((1156, "Y"), (1174, "N"))
    ),
    "马会": (1184, "马会", ((1184, "Y"),)),
    "万达文华": (1171, "万达文华", ((1171, "Y"),)),
    "西安君悦&淮扬1949": (
        1157, "西安区域运营", ((1157, "Y"), (1185, "N"), (1186, "N"))
    ),
}

CONTRACT_TYPES = {"劳动合同": "LABOR_CONTRACT", "劳务合同": "SERVICE_CONTRACT"}
CONTRACT_TERMS = {"固定期限": "FIXED_TERM", "无固定期限": "OPEN_ENDED"}
SOCIAL_TYPES = {"A": "SOCIAL_UNINSURED", "B": "SOCIAL_INSURED"}


class ImportError(RuntimeError):
    pass


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def date_value(value: object) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m-%d")
    raw = text(value)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return dt.datetime.strptime(raw, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    raise ImportError(f"unsupported date: {raw!r}")


def money(value: object, *, blank_zero: bool = False) -> decimal.Decimal:
    raw = text(value).replace(",", "")
    if not raw and blank_zero:
        return decimal.Decimal(0)
    try:
        return decimal.Decimal(raw)
    except decimal.InvalidOperation as exc:
        raise ImportError(f"unsupported money value: {value!r}") from exc


def sql(value: object) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, decimal.Decimal):
        return format(value, "f")
    if isinstance(value, int):
        return str(value)
    raw = str(value).replace("\\", "\\\\").replace("'", "''")
    return f"'{raw}'"


def load_rows(path: Path, expected_sha: str) -> tuple[list[dict[str, object]], str]:
    payload = path.read_bytes()
    digest = hashlib.sha256(payload).hexdigest()
    if expected_sha and digest.lower() != expected_sha.lower():
        raise ImportError(f"workbook sha256 changed: expected {expected_sha}, got {digest}")

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if SHEET not in workbook.sheetnames:
        raise ImportError(f"sheet not found: {SHEET}")
    sheet = workbook[SHEET]
    headers = tuple(text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
    if headers != REQUIRED_HEADERS:
        raise ImportError(f"unexpected signing headers: {headers!r}")

    rows: list[dict[str, object]] = []
    names: set[str] = set()
    phones: set[str] = set()
    ids: set[str] = set()
    for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not text(values[1]):
            continue
        record = dict(zip(headers, values, strict=True))
        name = text(record["姓名"])
        phone = re.sub(r"\D", "", text(record["联系电话"]))
        id_number = text(record["身份证号"]).upper()
        if name in names or phone in phones or id_number in ids:
            raise ImportError(f"duplicate identity at Excel row {excel_row}")
        if not re.fullmatch(r"1\d{10}", phone):
            raise ImportError(f"invalid phone at Excel row {excel_row}")
        if not re.fullmatch(r"\d{17}[0-9X]", id_number):
            raise ImportError(f"invalid ID number at Excel row {excel_row}")
        names.add(name); phones.add(phone); ids.add(id_number)

        raw_store = text(record["4级门店"]).replace("\r\n", "\n").replace("\r", "\n")
        store = re.sub(r"\s*\n\s*", "+", raw_store)
        store = re.sub(r"[ \t]+", "", store)
        if store not in ORG:
            raise ImportError(f"unsupported store at Excel row {excel_row}: {store!r}")
        post = text(record["员工岗位"])
        if post not in POSTS:
            raise ImportError(f"unsupported post at Excel row {excel_row}: {post!r}")
        contract_label = text(record["合同类型"])
        term_label = text(record["劳动合同期限形式"])
        social_label = text(record["社保类型"])
        if contract_label not in CONTRACT_TYPES or term_label not in CONTRACT_TERMS:
            raise ImportError(f"unsupported contract values at Excel row {excel_row}")
        if social_label not in SOCIAL_TYPES:
            raise ImportError(f"unsupported social type at Excel row {excel_row}")

        total = money(record["综合工资"])
        parts = [
            money(record["底薪"]), money(record["综合岗位津贴"]),
            money(record["综合驻外补贴"]), money(record["月度绩效津贴"]),
            money(record["房补"], blank_zero=True),
        ]
        if sum(parts, decimal.Decimal(0)) != total:
            raise ImportError(f"salary components do not equal total at Excel row {excel_row}")

        if contract_label == "劳动合同":
            category = "FULL_TIME"
        elif post == "实习生":
            category = "INTERN"
        else:
            category = "RETIRED_REHIRE"

        target_dept_id, target_dept_name, scopes = ORG[store]
        rows.append({
            "row_no": excel_row,
            "name": name,
            "phone": phone,
            "id_number": id_number,
            "employee_state": text(record["在职状态"]),
            "employee_status": text(record["员工状态"]),
            "level1": text(record["1级部门\n（金英）"]),
            "level2": text(record["2级部门\n（灵韵）"]),
            "level3": text(record["3级部门"]),
            "store_name": store,
            "department_supervisor": text(record["3/4级部门负责人"]) or None,
            "company_name": text(record["签约公司"]),
            "social_type": SOCIAL_TYPES[social_label],
            "position_name": post,
            "city_level": text(record["城市等级"]),
            "legal_representative": text(record["法定代表人"]),
            "registered_address": text(record["注册地"]),
            "contract_type": CONTRACT_TYPES[contract_label],
            "job_grade": text(record["员工职级"]),
            "work_location": text(record["工作地"]),
            "current_address": text(record["家庭住址"]),
            "entry_date": date_value(record["入职时间"]),
            "contract_term": CONTRACT_TERMS[term_label],
            "contract_start": date_value(record["劳动合同起始日期"]),
            "contract_end": date_value(record["劳动合同结束日期"]),
            "probation_start": date_value(record["试用期开始日期"]),
            "probation_end": date_value(record["试用期结束日期"]),
            "attendance_method": text(record["工时制度"]),
            "salary_total": total,
            "base_salary": parts[0],
            "post_salary": parts[1],
            "field_allowance": parts[2],
            "performance_salary": parts[3],
            "employee_category": category,
            "target_dept_id": target_dept_id,
            "target_dept_name": target_dept_name,
            "scopes": scopes,
        })

    if len(rows) != EXPECTED_ROWS:
        raise ImportError(f"expected {EXPECTED_ROWS} rows, got {len(rows)}")
    if {str(row["company_name"]) for row in rows} != {"舟山茗汇文化传播有限公司"}:
        raise ImportError("unexpected signing company set")
    if {str(row["legal_representative"]) for row in rows} != {"杜翠香"}:
        raise ImportError("unexpected legal representative set")
    if len({str(row["registered_address"]) for row in rows}) != 1:
        raise ImportError("registered address is not consistent")
    return rows, digest


def source_insert(row: dict[str, object]) -> str:
    values = [
        row["row_no"], row["name"], row["phone"], row["id_number"],
        row["employee_state"], row["employee_status"], row["level1"], row["level2"],
        row["level3"], row["store_name"], row["department_supervisor"],
        row["company_name"], row["social_type"], row["position_name"], row["city_level"],
        row["legal_representative"], row["registered_address"], row["contract_type"],
        row["job_grade"], row["work_location"], row["current_address"], row["entry_date"],
        row["contract_term"], row["contract_start"], row["contract_end"],
        row["probation_start"], row["probation_end"], row["attendance_method"],
        row["salary_total"], row["base_salary"], row["post_salary"],
        row["field_allowance"], row["performance_salary"], row["employee_category"],
        row["target_dept_id"], row["target_dept_name"],
    ]
    return "(" + ",".join(sql(value) for value in values) + ")"


def build_sql(rows: list[dict[str, object]], digest: str) -> str:
    procedure = f"tmp_reconcile_beijing_0722_{digest[:12]}"
    salary_version = f"BJ0722-{digest[:16]}"
    source_columns = (
        "row_no,name,phone,id_number,employee_state,employee_status,level1,level2,level3,"
        "store_name,department_supervisor,company_name,social_type,position_name,city_level,"
        "legal_representative,registered_address,contract_type,job_grade,work_location,"
        "current_address,entry_date,contract_term,contract_start,contract_end,probation_start,"
        "probation_end,attendance_method,salary_total,base_salary,post_salary,field_allowance,"
        "performance_salary,employee_category,target_dept_id,target_dept_name"
    )
    source_values = ",\n".join(source_insert(row) for row in rows)
    scope_values = ",\n".join(
        f"({row['row_no']},{dept_id},{sql(default_flag)})"
        for row in rows for dept_id, default_flag in row["scopes"]
    )
    source_schema = """
row_no int primary key,name varchar(64) not null,phone varchar(20) not null,id_number varchar(32) not null,
employee_state varchar(32),employee_status varchar(32),level1 varchar(100),level2 varchar(100),
level3 varchar(100),store_name varchar(160),department_supervisor varchar(64),company_name varchar(160),
social_type varchar(32),position_name varchar(100),city_level varchar(64),legal_representative varchar(64),
registered_address varchar(255),contract_type varchar(32),job_grade varchar(32),work_location varchar(100),
current_address varchar(255),entry_date date,contract_term varchar(32),contract_start date,contract_end date,
probation_start date,probation_end date,attendance_method varchar(64),salary_total decimal(16,2),
base_salary decimal(16,2),post_salary decimal(16,2),field_allowance decimal(16,2),
performance_salary decimal(16,2),employee_category varchar(32),target_dept_id bigint,target_dept_name varchar(100)
""".strip()

    return f"""SET NAMES utf8mb4 COLLATE utf8mb4_general_ci;
DROP TEMPORARY TABLE IF EXISTS tmp_bj0722_source;
DROP TEMPORARY TABLE IF EXISTS tmp_bj0722_scope;
DROP TEMPORARY TABLE IF EXISTS tmp_bj0722_map;
CREATE TEMPORARY TABLE tmp_bj0722_source ({source_schema}) ENGINE=InnoDB;
INSERT INTO tmp_bj0722_source({source_columns}) VALUES
{source_values};
CREATE TEMPORARY TABLE tmp_bj0722_scope(row_no int not null,dept_id bigint not null,is_default char(1) not null,primary key(row_no,dept_id)) ENGINE=InnoDB;
INSERT INTO tmp_bj0722_scope(row_no,dept_id,is_default) VALUES
{scope_values};

DROP PROCEDURE IF EXISTS {procedure};
DELIMITER $$
CREATE PROCEDURE {procedure}()
main: BEGIN
  DECLARE v_lock int DEFAULT 0;
  DECLARE v_count int DEFAULT 0;
  DECLARE v_legal_entity_id bigint DEFAULT NULL;
  DECLARE v_legal_entity_code varchar(64) DEFAULT NULL;
  DECLARE v_started_at datetime DEFAULT NOW();
  DECLARE EXIT HANDLER FOR SQLEXCEPTION
  BEGIN
    ROLLBACK;
    DO RELEASE_LOCK({sql(LOCK_NAME)});
    RESIGNAL;
  END;

  SELECT GET_LOCK({sql(LOCK_NAME)},30) INTO v_lock;
  IF v_lock <> 1 THEN SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='Beijing 0722 import lock unavailable'; END IF;
  START TRANSACTION;

  IF (SELECT COUNT(*) FROM tmp_bj0722_source) <> {EXPECTED_ROWS} THEN
    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='source row count changed';
  END IF;
  CREATE TEMPORARY TABLE tmp_bj0722_map AS
  SELECT s.row_no,MIN(u.user_id) AS user_id,COUNT(DISTINCT u.user_id) AS candidate_count
  FROM tmp_bj0722_source s
  LEFT JOIN sys_user u ON u.del_flag='0' AND (
       CAST(u.phonenumber AS BINARY)=CAST(s.phone AS BINARY)
    OR CAST(u.user_name AS BINARY)=CAST(s.phone AS BINARY)
    OR EXISTS (SELECT 1 FROM sys_user_profile px WHERE px.user_id=u.user_id
               AND CAST(px.id_number AS BINARY)=CAST(s.id_number AS BINARY))
  )
  GROUP BY s.row_no;
  ALTER TABLE tmp_bj0722_map ADD PRIMARY KEY(row_no),ADD KEY idx_bj0722_map_user(user_id);
  IF EXISTS (SELECT 1 FROM tmp_bj0722_map WHERE candidate_count<>1 OR user_id IS NULL) THEN
    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='identity is missing or ambiguous';
  END IF;
  IF (SELECT COUNT(DISTINCT user_id) FROM tmp_bj0722_map) <> {EXPECTED_ROWS} THEN
    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='multiple source rows resolve to one user';
  END IF;

  SELECT COUNT(*) INTO v_count FROM sys_user u JOIN tmp_bj0722_map m ON m.user_id=u.user_id FOR UPDATE;
  SELECT COUNT(*) INTO v_count FROM sys_user_profile p JOIN tmp_bj0722_map m ON m.user_id=p.user_id FOR UPDATE;
  SELECT COUNT(*) INTO v_count FROM sys_user_post up JOIN tmp_bj0722_map m ON m.user_id=up.user_id FOR UPDATE;
  SELECT COUNT(*) INTO v_count FROM sys_user_role ur JOIN tmp_bj0722_map m ON m.user_id=ur.user_id FOR UPDATE;
  SELECT COUNT(*) INTO v_count FROM sys_user_shop us JOIN tmp_bj0722_map m ON m.user_id=us.user_id FOR UPDATE;

  IF (SELECT COUNT(*) FROM sys_user_profile p JOIN tmp_bj0722_map m ON m.user_id=p.user_id) <> {EXPECTED_ROWS} THEN
    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='formal employee profile missing';
  END IF;
  IF EXISTS (
    SELECT 1 FROM sys_user other JOIN tmp_bj0722_source s
    JOIN tmp_bj0722_map m ON m.row_no=s.row_no
    WHERE CAST(other.user_name AS BINARY)=CAST(s.phone AS BINARY) AND other.user_id<>m.user_id
  ) THEN SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='target login name is already owned'; END IF;
  IF EXISTS (
    SELECT 1 FROM sys_user_profile other JOIN tmp_bj0722_source s
    JOIN tmp_bj0722_map m ON m.row_no=s.row_no
    WHERE CAST(other.id_number AS BINARY)=CAST(s.id_number AS BINARY) AND other.user_id<>m.user_id
  ) THEN SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='target ID number is already owned'; END IF;
  IF EXISTS (
    SELECT 1 FROM tmp_bj0722_source s LEFT JOIN sys_dept d ON d.dept_id=s.target_dept_id
    WHERE d.dept_id IS NULL OR d.status<>'0' OR d.del_flag<>'0'
       OR CAST(d.dept_name AS BINARY)<>CAST(s.target_dept_name AS BINARY)
  ) THEN SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='primary department mapping changed'; END IF;
  IF EXISTS (
    SELECT 1 FROM tmp_bj0722_scope sc LEFT JOIN sys_dept d ON d.dept_id=sc.dept_id
    WHERE d.dept_id IS NULL OR d.status<>'0' OR d.del_flag<>'0'
  ) THEN SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='scope department mapping changed'; END IF;
  IF EXISTS (
    SELECT s.position_name FROM tmp_bj0722_source s LEFT JOIN sys_post p
      ON CAST(p.post_name AS BINARY)=CAST(s.position_name AS BINARY) AND p.status='0'
    GROUP BY s.position_name HAVING COUNT(DISTINCT p.post_id)<>1
  ) THEN SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='post mapping is missing or ambiguous'; END IF;
  IF EXISTS (
    SELECT s.position_name FROM tmp_bj0722_source s LEFT JOIN sys_role r
      ON CAST(r.role_name AS BINARY)=CAST(s.position_name AS BINARY) AND r.status='0' AND r.del_flag='0'
    GROUP BY s.position_name HAVING COUNT(DISTINCT r.role_id)<>1
  ) THEN SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='role mapping is missing or ambiguous'; END IF;

  IF (SELECT COUNT(*) FROM sys_legal_entity e JOIN (SELECT DISTINCT company_name FROM tmp_bj0722_source) s
      ON CAST(e.legal_entity_name AS BINARY)=CAST(s.company_name AS BINARY) WHERE e.status='0') <> 1 THEN
    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='legal entity mapping is missing or ambiguous';
  END IF;
  SELECT e.legal_entity_id,e.legal_entity_code INTO v_legal_entity_id,v_legal_entity_code
  FROM sys_legal_entity e JOIN (SELECT DISTINCT company_name FROM tmp_bj0722_source) s
    ON CAST(e.legal_entity_name AS BINARY)=CAST(s.company_name AS BINARY)
  WHERE e.status='0' FOR UPDATE;
  UPDATE sys_legal_entity e JOIN (
    SELECT DISTINCT company_name,legal_representative,registered_address FROM tmp_bj0722_source
  ) s ON e.legal_entity_id=v_legal_entity_id
  SET e.legal_representative=s.legal_representative,e.registered_address=s.registered_address,
      e.update_by={sql(OPERATOR)},e.update_time=NOW();

  INSERT INTO hr_employee_position_no_history(
    user_id,employee_no,old_position_no,new_position_no,old_post_id,new_post_id,
    old_post_code,new_post_code,effective_date,change_source,change_by,changed_time,remark
  )
  SELECT p.user_id,p.employee_no,p.position_no,CONCAT(UPPER(tp.post_code),'-',p.employee_no),
         oldp.old_post_id,tp.post_id,op.post_code,tp.post_code,COALESCE(s.entry_date,CURDATE()),
         'POSITION_CHANGE',{sql(OPERATOR)},NOW(),
         CONCAT('北京区域0722签约数据;源行',s.row_no,';sha256={digest[:16]}')
  FROM tmp_bj0722_source s JOIN tmp_bj0722_map m ON m.row_no=s.row_no
  JOIN sys_user_profile p ON p.user_id=m.user_id
  JOIN sys_post tp ON CAST(tp.post_name AS BINARY)=CAST(s.position_name AS BINARY) AND tp.status='0'
  LEFT JOIN (SELECT up.user_id,MIN(up.post_id) old_post_id FROM sys_user_post up GROUP BY up.user_id) oldp
    ON oldp.user_id=m.user_id
  LEFT JOIN sys_post op ON op.post_id=oldp.old_post_id
  WHERE NOT (CAST(p.position_no AS BINARY) <=> CAST(CONCAT(UPPER(tp.post_code),'-',p.employee_no) AS BINARY))
     OR NOT (oldp.old_post_id <=> tp.post_id);

  UPDATE sys_user u JOIN tmp_bj0722_map m ON m.user_id=u.user_id
  JOIN tmp_bj0722_source s ON s.row_no=m.row_no
  SET u.dept_id=s.target_dept_id,u.user_name=s.phone,u.nick_name=s.name,u.phonenumber=s.phone,
      u.status='0',u.del_flag='0',u.update_by={sql(OPERATOR)},u.update_time=NOW(),
      u.remark=CONCAT_WS('; ',NULLIF(u.remark,''),CONCAT('北京区域0722签约数据权威同步;源行',s.row_no,';sha256={digest[:16]}'));
  IF ROW_COUNT() <> {EXPECTED_ROWS} THEN SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='user update count mismatch'; END IF;

  UPDATE sys_user_profile p JOIN tmp_bj0722_map m ON m.user_id=p.user_id
  JOIN tmp_bj0722_source s ON s.row_no=m.row_no
  JOIN sys_post tp ON CAST(tp.post_name AS BINARY)=CAST(s.position_name AS BINARY) AND tp.status='0'
  SET p.dept_id=s.target_dept_id,p.position_no=CONCAT(UPPER(tp.post_code),'-',p.employee_no),
      p.company_name=s.company_name,p.dept_level1_name=s.level1,p.dept_level2_name=s.level2,
      p.dept_level3_name=s.level3,p.store_name=s.store_name,p.position_names=s.position_name,
      p.job_grade=s.job_grade,p.department_supervisor=s.department_supervisor,
      p.employee_status=s.employee_status,p.employee_category=s.employee_category,
      p.id_type='居民身份证',p.id_number=s.id_number,p.current_address=s.current_address,
      p.entry_date=s.entry_date,p.probation_start_date=s.probation_start,p.probation_end_date=s.probation_end,
      p.contract_start_date=s.contract_start,p.contract_end_date=s.contract_end,
      p.contract_type=s.contract_type,p.contract_term=s.contract_term,p.work_location=s.work_location,
      p.work_city_level=s.city_level,p.attendance_method=s.attendance_method,p.social_type=s.social_type,
      p.legal_entity=s.company_name,p.legal_entity_id=v_legal_entity_id,p.legal_entity_code=v_legal_entity_code,
      p.base_salary=s.base_salary,p.post_salary=s.post_salary,p.field_allowance=s.field_allowance,
      p.performance_salary=s.performance_salary,p.salary_total=s.salary_total,
      p.salary_version={sql(salary_version)},p.update_by={sql(OPERATOR)},p.update_time=NOW();
  IF ROW_COUNT() <> {EXPECTED_ROWS} THEN SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='profile update count mismatch'; END IF;

  DELETE up FROM sys_user_post up JOIN tmp_bj0722_map m ON m.user_id=up.user_id;
  INSERT INTO sys_user_post(user_id,post_id)
  SELECT m.user_id,p.post_id FROM tmp_bj0722_map m JOIN tmp_bj0722_source s ON s.row_no=m.row_no
  JOIN sys_post p ON CAST(p.post_name AS BINARY)=CAST(s.position_name AS BINARY) AND p.status='0';

  DELETE ur FROM sys_user_role ur JOIN tmp_bj0722_map m ON m.user_id=ur.user_id
  WHERE ur.role_id IN (
    SELECT DISTINCT r.role_id FROM sys_role r JOIN tmp_bj0722_source s
      ON CAST(r.role_name AS BINARY)=CAST(s.position_name AS BINARY)
    WHERE r.status='0' AND r.del_flag='0'
  );
  INSERT INTO sys_user_role(user_id,role_id)
  SELECT m.user_id,r.role_id FROM tmp_bj0722_map m JOIN tmp_bj0722_source s ON s.row_no=m.row_no
  JOIN sys_role r ON CAST(r.role_name AS BINARY)=CAST(s.position_name AS BINARY)
    AND r.status='0' AND r.del_flag='0'
  ON DUPLICATE KEY UPDATE role_id=VALUES(role_id);

  DELETE us FROM sys_user_shop us JOIN tmp_bj0722_map m ON m.user_id=us.user_id WHERE us.is_default='Y';
  INSERT INTO sys_user_shop(user_id,dept_id,is_default,create_by,create_time)
  SELECT m.user_id,sc.dept_id,sc.is_default,{sql(OPERATOR)},NOW()
  FROM tmp_bj0722_scope sc JOIN tmp_bj0722_map m ON m.row_no=sc.row_no
  ON DUPLICATE KEY UPDATE is_default=VALUES(is_default);

  INSERT INTO sys_security_session_outbox(
    event_id,event_type,user_id,reason_code,status,attempts,available_at,created_at
  )
  SELECT CONCAT('BJ0722-',SUBSTR(SHA2(CONCAT(UUID(),':',m.user_id,':',m.row_no),256),1,57)),
         'INVALIDATE_USER_SESSIONS',m.user_id,'USER_ROLES_CHANGED','PENDING',0,NOW(),NOW()
  FROM tmp_bj0722_map m;

  IF (SELECT COUNT(*) FROM tmp_bj0722_source s JOIN tmp_bj0722_map m ON m.row_no=s.row_no
      JOIN sys_user u ON u.user_id=m.user_id JOIN sys_user_profile p ON p.user_id=m.user_id
      JOIN sys_post tp ON CAST(tp.post_name AS BINARY)=CAST(s.position_name AS BINARY) AND tp.status='0'
      WHERE CAST(u.user_name AS BINARY)=CAST(s.phone AS BINARY)
        AND CAST(u.nick_name AS BINARY)=CAST(s.name AS BINARY)
        AND CAST(u.phonenumber AS BINARY)=CAST(s.phone AS BINARY)
        AND u.dept_id=s.target_dept_id AND u.status='0' AND u.del_flag='0'
        AND p.dept_id=s.target_dept_id AND CAST(p.id_number AS BINARY)=CAST(s.id_number AS BINARY)
        AND CAST(p.store_name AS BINARY)=CAST(s.store_name AS BINARY)
        AND CAST(p.position_names AS BINARY)=CAST(s.position_name AS BINARY)
        AND CAST(p.position_no AS BINARY)=CAST(CONCAT(UPPER(tp.post_code),'-',p.employee_no) AS BINARY)
        AND CAST(p.employee_status AS BINARY) <=> CAST(s.employee_status AS BINARY)
        AND CAST(p.employee_category AS BINARY) <=> CAST(s.employee_category AS BINARY)
        AND p.entry_date <=> s.entry_date AND p.contract_start_date <=> s.contract_start
        AND p.contract_end_date <=> s.contract_end AND p.probation_start_date <=> s.probation_start
        AND p.probation_end_date <=> s.probation_end AND p.salary_total <=> s.salary_total
        AND p.base_salary <=> s.base_salary AND p.post_salary <=> s.post_salary
        AND p.field_allowance <=> s.field_allowance AND p.performance_salary <=> s.performance_salary
        AND p.legal_entity_id=v_legal_entity_id) <> {EXPECTED_ROWS} THEN
    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='final employee/profile verification failed';
  END IF;
  IF EXISTS (
    SELECT m.user_id FROM tmp_bj0722_map m JOIN tmp_bj0722_source s ON s.row_no=m.row_no
    LEFT JOIN sys_user_post up ON up.user_id=m.user_id
    LEFT JOIN sys_post p ON p.post_id=up.post_id
    GROUP BY m.user_id,s.position_name
    HAVING COUNT(up.post_id)<>1 OR SUM(CAST(p.post_name AS BINARY)=CAST(s.position_name AS BINARY))<>1
  ) THEN SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='final post verification failed'; END IF;
  IF EXISTS (
    SELECT m.user_id FROM tmp_bj0722_map m JOIN tmp_bj0722_source s ON s.row_no=m.row_no
    JOIN sys_role target ON CAST(target.role_name AS BINARY)=CAST(s.position_name AS BINARY)
      AND target.status='0' AND target.del_flag='0'
    LEFT JOIN sys_user_role ur ON ur.user_id=m.user_id AND ur.role_id=target.role_id
    GROUP BY m.user_id,target.role_id HAVING COUNT(ur.role_id)<>1
  ) THEN SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='final role verification failed'; END IF;
  IF EXISTS (
    SELECT m.user_id FROM tmp_bj0722_map m LEFT JOIN sys_user_shop us
      ON us.user_id=m.user_id AND us.is_default='Y'
    GROUP BY m.user_id HAVING COUNT(us.dept_id)<>1
  ) THEN SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='final default scope verification failed'; END IF;
  IF EXISTS (
    SELECT 1 FROM tmp_bj0722_scope sc JOIN tmp_bj0722_map m ON m.row_no=sc.row_no
    LEFT JOIN sys_user_shop us ON us.user_id=m.user_id AND us.dept_id=sc.dept_id
      AND CAST(us.is_default AS BINARY)=CAST(sc.is_default AS BINARY)
    WHERE us.user_id IS NULL
  ) THEN SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='final shop scope verification failed'; END IF;
  IF (SELECT COUNT(*) FROM sys_security_session_outbox o JOIN tmp_bj0722_map m ON m.user_id=o.user_id
      WHERE CAST(o.event_id AS BINARY) LIKE CAST('BJ0722-%' AS BINARY)
        AND o.created_at>=v_started_at) <> {EXPECTED_ROWS} THEN
    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='session invalidation verification failed';
  END IF;

  COMMIT;
  DO RELEASE_LOCK({sql(LOCK_NAME)});
END$$
DELIMITER ;
CALL {procedure}();
DROP PROCEDURE IF EXISTS {procedure};
SELECT 'IMPORT_OK' AS result,COUNT(*) AS employees,
       SUM(CAST(u.user_name AS BINARY)=CAST(s.phone AS BINARY)) AS phone_login_accounts,
       {sql(digest)} AS workbook_sha256
FROM tmp_bj0722_map m JOIN sys_user u ON u.user_id=m.user_id
JOIN tmp_bj0722_source s ON s.row_no=m.row_no;
"""


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--expected-sha", default="")
    args = parser.parse_args()

    excel = Path(args.excel).expanduser().resolve()
    output = Path(args.output).expanduser().resolve()
    if not excel.is_file():
        raise ImportError(f"Excel file not found: {excel}")
    rows, digest = load_rows(excel, args.expected_sha)
    generated = build_sql(rows, digest)
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor = os.open(output, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
        handle.write(generated)
    print(f"workbook_sha256={digest}")
    print(f"source_sheet={SHEET}")
    print(f"source_rows={len(rows)}")
    print(f"unique_phones={len({row['phone'] for row in rows})}")
    print(f"unique_ids={len({row['id_number'] for row in rows})}")
    print(f"output={output}")
    print(f"output_bytes={output.stat().st_size}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except ImportError as exc:
        print(f"ERROR: {exc}", file=os.sys.stderr)
        raise SystemExit(2)
