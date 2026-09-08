#!/usr/bin/env python3
"""Generate the audited product-master refresh SQL from the 2026 workbook.

Only visible product rows from the ``供应链平台产品2026版`` sheet are active.
Hidden product rows are retained in the database for historical references but
are soft-disabled by the generated migration.
"""

from __future__ import annotations

import argparse
import hashlib
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


SHEET_NAME = "供应链平台产品2026版"
EXPECTED_SOURCE_SHA256 = "4716cc71bca2319880a83d8dfdd942673a9003388173bb124bc61ca7547263bf"
EXPECTED_VISIBLE_ROWS = 158
EXPECTED_HIDDEN_PRODUCT_ROWS = (52, 53, 54, 55, 59, 60, 61, 62, 63)
EXPECTED_VISIBLE_CATEGORIES = 21
EXPECTED_VISIBLE_SUPPLIERS = 82
HEADER_NAMES = {"产品类别名称", "产品售卖名称", "品名", "名称"}
SUPPLIER_PREVIOUS_NAMES = {"杭州塞纳茶叶有限公司132": "杭州塞纳茶叶有限公司32"}
MONEY_QUANTUM = Decimal("0.01")


@dataclass(frozen=True)
class ProductRow:
    source_row: int
    category: str
    name: str
    grade: str | None
    spec: str | None
    description: str | None
    unit: str | None
    reference_cost: Decimal | None
    supplier: str | None
    phone: str | None
    retail_price: Decimal | None


def text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    normalized = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    return normalized or None


def money(value: object, source_row: int, column: str) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise ValueError(f"row {source_row} column {column} is not a money value: {value!r}")
    try:
        return Decimal(str(value)).quantize(MONEY_QUANTUM, rounding=ROUND_HALF_UP)
    except Exception as exc:
        raise ValueError(f"row {source_row} column {column} is not numeric: {value!r}") from exc


def sql_text(value: str | None) -> str:
    if value is None:
        return "NULL"
    return "'" + value.replace("'", "''") + "'"


def sql_money(value: Decimal | None) -> str:
    return "NULL" if value is None else format(value, ".2f")


def chunks(items: list[str], size: int) -> Iterable[list[str]]:
    for index in range(0, len(items), size):
        yield items[index:index + size]


def parse_workbook(path: Path) -> tuple[list[ProductRow], list[int]]:
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    if source_hash != EXPECTED_SOURCE_SHA256:
        raise ValueError(
            f"unexpected workbook SHA-256 {source_hash}; expected {EXPECTED_SOURCE_SHA256}"
        )

    workbook = load_workbook(path, data_only=True, read_only=False)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"missing sheet {SHEET_NAME!r}")
    sheet = workbook[SHEET_NAME]

    visible: list[ProductRow] = []
    hidden_product_rows: list[int] = []
    current_category: str | None = None
    for row_number in range(4, sheet.max_row + 1):
        category_cell = text(sheet.cell(row_number, 2).value)
        if category_cell not in {None, "上线分类", "品类"}:
            current_category = category_cell

        name = text(sheet.cell(row_number, 3).value)
        if name is None or name in HEADER_NAMES:
            continue
        if sheet.row_dimensions[row_number].hidden:
            hidden_product_rows.append(row_number)
            continue
        if current_category is None:
            raise ValueError(f"visible product row {row_number} has no category")

        visible.append(
            ProductRow(
                source_row=row_number,
                category=current_category,
                name=name,
                grade=text(sheet.cell(row_number, 4).value),
                spec=text(sheet.cell(row_number, 5).value),
                description=text(sheet.cell(row_number, 6).value),
                unit=text(sheet.cell(row_number, 7).value),
                reference_cost=money(sheet.cell(row_number, 8).value, row_number, "H"),
                supplier=text(sheet.cell(row_number, 9).value),
                phone=text(sheet.cell(row_number, 10).value),
                retail_price=money(sheet.cell(row_number, 11).value, row_number, "K"),
            )
        )

    if len(visible) != EXPECTED_VISIBLE_ROWS:
        raise ValueError(f"expected {EXPECTED_VISIBLE_ROWS} visible products, got {len(visible)}")
    if tuple(hidden_product_rows) != EXPECTED_HIDDEN_PRODUCT_ROWS:
        raise ValueError(
            f"unexpected hidden product rows {hidden_product_rows}; "
            f"expected {list(EXPECTED_HIDDEN_PRODUCT_ROWS)}"
        )
    categories = {row.category for row in visible}
    if len(categories) != EXPECTED_VISIBLE_CATEGORIES:
        raise ValueError(
            f"expected {EXPECTED_VISIBLE_CATEGORIES} visible categories, got {len(categories)}"
        )
    suppliers = {row.supplier for row in visible if row.supplier is not None}
    if len(suppliers) != EXPECTED_VISIBLE_SUPPLIERS:
        raise ValueError(
            f"expected {EXPECTED_VISIBLE_SUPPLIERS} visible suppliers, got {len(suppliers)}"
        )
    return visible, hidden_product_rows


def supplier_rows(products: list[ProductRow]) -> list[tuple[str, str | None, str | None]]:
    phones: dict[str, set[str | None]] = {}
    for product in products:
        if product.supplier is None:
            continue
        phones.setdefault(product.supplier, set()).add(product.phone)
    conflicts = {name: values for name, values in phones.items() if len(values) != 1}
    if conflicts:
        raise ValueError(f"supplier phone conflicts: {conflicts}")
    return sorted(
        (
            name,
            SUPPLIER_PREVIOUS_NAMES.get(name),
            next(iter(values)),
        )
        for name, values in phones.items()
    )


def values_sql(rows: list[str], prefix: str = "") -> list[str]:
    statements: list[str] = []
    for group in chunks(rows, 50):
        statements.append(prefix + ",\n".join(group) + ";")
    return statements


def build_refresh_sql(
    products: list[ProductRow],
    hidden_rows: list[int],
    source_path: Path,
    database: str,
    shop_dept_id: int,
) -> str:
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    product_values = [
        "(" + ", ".join(
            [
                str(row.source_row),
                sql_text(row.category),
                sql_text(row.name),
                sql_text(row.grade),
                sql_text(row.spec),
                sql_text(row.description),
                sql_text(row.unit),
                sql_money(row.reference_cost),
                sql_text(row.supplier),
                sql_text(row.phone),
                sql_money(row.retail_price),
            ]
        ) + ")"
        for row in products
    ]
    supplier_values = [
        f"({sql_text(name)}, {sql_text(previous)}, {sql_text(phone)})"
        for name, previous, phone in supplier_rows(products)
    ]
    category_values = [f"({sql_text(name)})" for name in sorted({row.category for row in products})]
    hidden_values = [f"({row})" for row in hidden_rows]

    lines = [
        f"-- Generated from {source_path.name}",
        f"-- Source SHA-256: {source_hash}",
        f"-- Source sheet: {SHEET_NAME}; visible products: {len(products)}; hidden products soft-disabled: {len(hidden_rows)}",
        "-- This migration preserves product_id, product_code, stock, stock logs, and all historical business references.",
        "SET NAMES utf8mb4;",
        f"USE `{database.replace('`', '``')}`;",
        "",
        "DROP TEMPORARY TABLE IF EXISTS tmp_product_master_20260711;",
        "CREATE TEMPORARY TABLE tmp_product_master_20260711 (",
        "    source_row int NOT NULL PRIMARY KEY,",
        "    category_name varchar(128) NOT NULL,",
        "    product_name varchar(128) NOT NULL,",
        "    grade varchar(32) NULL,",
        "    spec varchar(256) NULL,",
        "    product_description text NULL,",
        "    unit varchar(32) NULL,",
        "    reference_cost decimal(16,2) NULL,",
        "    supplier_name varchar(128) NULL,",
        "    supplier_phone varchar(32) NULL,",
        "    retail_price decimal(16,2) NULL",
        ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;",
        *values_sql(product_values, "INSERT INTO tmp_product_master_20260711 VALUES\n"),
        "",
        "DROP TEMPORARY TABLE IF EXISTS tmp_product_hidden_20260711;",
        "CREATE TEMPORARY TABLE tmp_product_hidden_20260711 (source_row int NOT NULL PRIMARY KEY);",
        *values_sql(hidden_values, "INSERT INTO tmp_product_hidden_20260711 VALUES\n"),
        "",
        "DROP TEMPORARY TABLE IF EXISTS tmp_product_category_20260711;",
        "CREATE TEMPORARY TABLE tmp_product_category_20260711 (category_name varchar(128) NOT NULL PRIMARY KEY)",
        "ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;",
        *values_sql(category_values, "INSERT INTO tmp_product_category_20260711 VALUES\n"),
        "",
        "DROP TEMPORARY TABLE IF EXISTS tmp_product_supplier_20260711;",
        "CREATE TEMPORARY TABLE tmp_product_supplier_20260711 (",
        "    supplier_name varchar(128) NOT NULL PRIMARY KEY,",
        "    previous_name varchar(128) NULL,",
        "    contact_phone varchar(32) NULL",
        ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;",
        *values_sql(supplier_values, "INSERT INTO tmp_product_supplier_20260711 VALUES\n"),
        "",
        "CREATE TABLE IF NOT EXISTS backup_inv_product_before_20260711_refresh LIKE inv_product;",
        "INSERT INTO backup_inv_product_before_20260711_refresh",
        "SELECT source.* FROM inv_product source",
        "WHERE NOT EXISTS (SELECT 1 FROM backup_inv_product_before_20260711_refresh LIMIT 1);",
        "CREATE TABLE IF NOT EXISTS backup_inv_product_category_before_20260711_refresh LIKE inv_product_category;",
        "INSERT INTO backup_inv_product_category_before_20260711_refresh",
        "SELECT source.* FROM inv_product_category source",
        "WHERE NOT EXISTS (SELECT 1 FROM backup_inv_product_category_before_20260711_refresh LIMIT 1);",
        "CREATE TABLE IF NOT EXISTS backup_inv_supplier_before_20260711_refresh LIKE inv_supplier;",
        "INSERT INTO backup_inv_supplier_before_20260711_refresh",
        "SELECT source.* FROM inv_supplier source",
        "WHERE NOT EXISTS (SELECT 1 FROM backup_inv_supplier_before_20260711_refresh LIMIT 1);",
        "",
        "DROP PROCEDURE IF EXISTS apply_product_master_20260711;",
        "DELIMITER $$",
        "CREATE PROCEDURE apply_product_master_20260711()",
        "BEGIN",
        "    DECLARE current_products int DEFAULT 0;",
        "    DECLARE visible_matches int DEFAULT 0;",
        "    DECLARE hidden_matches int DEFAULT 0;",
        "    DECLARE category_matches int DEFAULT 0;",
        "    DECLARE supplier_matches int DEFAULT 0;",
        "    DECLARE mismatch_count int DEFAULT 0;",
        "    DECLARE EXIT HANDLER FOR SQLEXCEPTION",
        "    BEGIN",
        "        ROLLBACK;",
        "        RESIGNAL;",
        "    END;",
        "",
        f"    SELECT COUNT(*) INTO current_products FROM inv_product WHERE shop_dept_id = {shop_dept_id};",
        "    IF current_products <> 167 THEN",
        "        SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT = 'preflight failed: expected 167 current products';",
        "    END IF;",
        "    IF (SELECT COUNT(*) FROM backup_inv_product_before_20260711_refresh) <> 167",
        "       OR (SELECT COUNT(*) FROM backup_inv_product_category_before_20260711_refresh) <> 23",
        "       OR (SELECT COUNT(*) FROM backup_inv_supplier_before_20260711_refresh) <> 87 THEN",
        "        SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT = 'preflight failed: backup row counts do not match';",
        "    END IF;",
        "",
        "    SELECT COUNT(DISTINCT p.product_id) INTO visible_matches",
        "    FROM inv_product p",
        "    JOIN tmp_product_master_20260711 source",
        "      ON CAST(SUBSTRING_INDEX(p.remark, 'row=', -1) AS UNSIGNED) = source.source_row",
        f"    WHERE p.shop_dept_id = {shop_dept_id};",
        f"    IF visible_matches <> {len(products)} THEN",
        "        SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT = 'preflight failed: visible product row mapping is not one-to-one';",
        "    END IF;",
        "",
        "    SELECT COUNT(DISTINCT p.product_id) INTO hidden_matches",
        "    FROM inv_product p",
        "    JOIN tmp_product_hidden_20260711 source",
        "      ON CAST(SUBSTRING_INDEX(p.remark, 'row=', -1) AS UNSIGNED) = source.source_row",
        f"    WHERE p.shop_dept_id = {shop_dept_id};",
        f"    IF hidden_matches <> {len(hidden_rows)} THEN",
        "        SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT = 'preflight failed: hidden product row mapping is not one-to-one';",
        "    END IF;",
        "",
        "    SELECT COUNT(DISTINCT c.category_id) INTO category_matches",
        "    FROM inv_product_category c",
        "    JOIN tmp_product_category_20260711 source ON source.category_name = c.category_name",
        f"    WHERE c.shop_dept_id = {shop_dept_id};",
        f"    IF category_matches <> {EXPECTED_VISIBLE_CATEGORIES} THEN",
        "        SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT = 'preflight failed: category mapping is incomplete';",
        "    END IF;",
        "",
        "    SELECT COUNT(DISTINCT s.supplier_id) INTO supplier_matches",
        "    FROM inv_supplier s",
        "    JOIN tmp_product_supplier_20260711 source",
        "      ON s.supplier_name = source.supplier_name OR s.supplier_name = source.previous_name",
        f"    WHERE s.shop_dept_id = {shop_dept_id};",
        f"    IF supplier_matches <> {EXPECTED_VISIBLE_SUPPLIERS} THEN",
        "        SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT = 'preflight failed: supplier mapping is incomplete';",
        "    END IF;",
        "",
        "    START TRANSACTION;",
        "",
        "    UPDATE inv_supplier s",
        "    JOIN tmp_product_supplier_20260711 source",
        "      ON s.supplier_name = source.supplier_name OR s.supplier_name = source.previous_name",
        "    SET s.supplier_name = source.supplier_name,",
        "        s.contact_person = '',",
        "        s.contact_phone = COALESCE(source.contact_phone, ''),",
        "        s.status = '0',",
        "        s.cooperation_status = '0',",
        "        s.update_by = 'product_master_20260711',",
        "        s.update_time = NOW(),",
        "        s.remark = '';",
        "",
        "    UPDATE inv_supplier s",
        "    LEFT JOIN tmp_product_supplier_20260711 source ON source.supplier_name = s.supplier_name",
        "    SET s.status = '1',",
        "        s.cooperation_status = '2',",
        "        s.update_by = 'product_master_20260711',",
        "        s.update_time = NOW()",
        f"    WHERE s.shop_dept_id = {shop_dept_id}",
        "      AND s.create_by = 'xlsx_reimport'",
        "      AND source.supplier_name IS NULL;",
        "",
        "    UPDATE inv_product_category c",
        "    JOIN tmp_product_category_20260711 source ON source.category_name = c.category_name",
        "    SET c.status = '0',",
        "        c.del_flag = '0',",
        "        c.update_by = 'product_master_20260711',",
        "        c.update_time = NOW()",
        f"    WHERE c.shop_dept_id = {shop_dept_id};",
        "",
        "    UPDATE inv_product_category c",
        "    LEFT JOIN tmp_product_category_20260711 source ON source.category_name = c.category_name",
        "    SET c.status = '1',",
        "        c.del_flag = '2',",
        "        c.update_by = 'product_master_20260711',",
        "        c.update_time = NOW()",
        f"    WHERE c.shop_dept_id = {shop_dept_id}",
        "      AND c.create_by = 'xlsx_reimport'",
        "      AND source.category_name IS NULL;",
        "",
        "    UPDATE inv_product p",
        "    JOIN tmp_product_master_20260711 source",
        "      ON CAST(SUBSTRING_INDEX(p.remark, 'row=', -1) AS UNSIGNED) = source.source_row",
        "    JOIN inv_product_category c",
        "      ON c.category_name = source.category_name",
        f"     AND c.shop_dept_id = {shop_dept_id}",
        "    SET p.product_name = source.product_name,",
        "        p.category_id = c.category_id,",
        "        p.grade = source.grade,",
        "        p.sku = NULL,",
        "        p.spec = source.spec,",
        "        p.size = NULL,",
        "        p.unit = source.unit,",
        "        p.purchase_price = source.reference_cost,",
        "        p.sales_price = source.retail_price,",
        "        p.sale_price_250g = NULL,",
        "        p.sale_price_500g = NULL,",
        "        p.cost_price = source.reference_cost,",
        "        p.supplier_name = source.supplier_name,",
        "        p.supplier_phone = source.supplier_phone,",
        "        p.supplier_remark = NULL,",
        "        p.internal_tea_name = NULL,",
        "        p.product_description = source.product_description,",
        "        p.barcode = NULL,",
        "        p.image_url = NULL,",
        "        p.package_image_url = NULL,",
        "        p.dry_tea_image_url = NULL,",
        "        p.tea_soup_image_url = NULL,",
        "        p.leaf_bottom_image_url = NULL,",
        "        p.extra_image_url = NULL,",
        "        p.status = '0',",
        "        p.del_flag = '0',",
        "        p.update_by = 'product_master_20260711',",
        "        p.update_time = NOW(),",
        "        p.remark = CONCAT('Imported from 供应链平台产品20260711更新版.xlsx sheet=供应链平台产品2026版 row=', source.source_row)",
        f"    WHERE p.shop_dept_id = {shop_dept_id};",
        "",
        "    UPDATE inv_product p",
        "    JOIN tmp_product_hidden_20260711 source",
        "      ON CAST(SUBSTRING_INDEX(p.remark, 'row=', -1) AS UNSIGNED) = source.source_row",
        "    SET p.status = '1',",
        "        p.del_flag = '2',",
        "        p.update_by = 'product_master_20260711',",
        "        p.update_time = NOW(),",
        "        p.remark = CONCAT('Excluded hidden row from 供应链平台产品20260711更新版.xlsx sheet=供应链平台产品2026版 row=', source.source_row)",
        f"    WHERE p.shop_dept_id = {shop_dept_id};",
        "",
        "    SELECT COUNT(*) INTO mismatch_count",
        "    FROM inv_product p",
        "    JOIN tmp_product_master_20260711 source",
        "      ON CAST(SUBSTRING_INDEX(p.remark, 'row=', -1) AS UNSIGNED) = source.source_row",
        "    JOIN inv_product_category c ON c.category_id = p.category_id",
        "    WHERE NOT (",
        "        p.product_name <=> source.product_name",
        "        AND c.category_name <=> source.category_name",
        "        AND p.grade <=> source.grade",
        "        AND p.spec <=> source.spec",
        "        AND p.product_description <=> source.product_description",
        "        AND p.unit <=> source.unit",
        "        AND p.purchase_price <=> source.reference_cost",
        "        AND p.cost_price <=> source.reference_cost",
        "        AND p.sales_price <=> source.retail_price",
        "        AND p.supplier_name <=> source.supplier_name",
        "        AND p.supplier_phone <=> source.supplier_phone",
        "        AND p.status = '0'",
        "        AND p.del_flag = '0'",
        "    );",
        "    IF mismatch_count <> 0 THEN",
        "        SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT = 'postflight failed: source/product mismatch remains';",
        "    END IF;",
        f"    IF (SELECT COUNT(*) FROM inv_product WHERE shop_dept_id = {shop_dept_id} AND del_flag = '0') <> {len(products)} THEN",
        f"        SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT = 'postflight failed: active product count is not {len(products)}';",
        "    END IF;",
        f"    IF (SELECT COUNT(*) FROM inv_product WHERE shop_dept_id = {shop_dept_id} AND del_flag = '2' AND status = '1') <> {len(hidden_rows)} THEN",
        f"        SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT = 'postflight failed: hidden product count is not {len(hidden_rows)}';",
        "    END IF;",
        "",
        "    COMMIT;",
        "END$$",
        "DELIMITER ;",
        "CALL apply_product_master_20260711();",
        "DROP PROCEDURE apply_product_master_20260711;",
        "",
        f"SELECT 'active_products' AS metric, COUNT(*) AS value FROM inv_product WHERE shop_dept_id = {shop_dept_id} AND del_flag = '0'",
        f"UNION ALL SELECT 'soft_disabled_hidden_products', COUNT(*) FROM inv_product WHERE shop_dept_id = {shop_dept_id} AND del_flag = '2' AND status = '1'",
        f"UNION ALL SELECT 'active_categories', COUNT(*) FROM inv_product_category WHERE shop_dept_id = {shop_dept_id} AND del_flag = '0' AND status = '0'",
        f"UNION ALL SELECT 'active_suppliers', COUNT(*) FROM inv_supplier WHERE shop_dept_id = {shop_dept_id} AND status = '0' AND cooperation_status = '0';",
        "",
    ]
    return "\n".join(lines)


def build_rollback_sql(database: str, shop_dept_id: int) -> str:
    product_columns = [
        "product_name", "product_code", "category_id", "grade", "sku", "spec", "size", "unit",
        "purchase_price", "sales_price", "sale_price_250g", "sale_price_500g", "cost_price",
        "safety_stock_min", "safety_stock_max", "supplier_name", "supplier_phone", "supplier_remark",
        "internal_tea_name", "product_description", "barcode", "image_url", "package_image_url",
        "dry_tea_image_url", "tea_soup_image_url", "leaf_bottom_image_url", "extra_image_url",
        "shop_dept_id", "status", "del_flag", "create_by", "create_time", "update_by", "update_time", "remark",
    ]
    supplier_columns = [
        "supplier_name", "supplier_code", "contact_person", "contact_phone", "contact_email", "address",
        "settlement_method", "cooperation_status", "shop_dept_id", "status", "create_by", "create_time",
        "update_by", "update_time", "remark",
    ]
    category_columns = [
        "parent_id", "ancestors", "category_name", "category_code", "order_num", "shop_dept_id",
        "status", "del_flag", "create_by", "create_time", "update_by", "update_time", "remark",
    ]

    def assignments(columns: list[str]) -> str:
        return ",\n        ".join(f"target.{column} = backup.{column}" for column in columns)

    return f"""-- Roll back erp_inventory_product_master_20260711.sql from its in-database snapshots.
SET NAMES utf8mb4;
USE `{database.replace('`', '``')}`;

DROP PROCEDURE IF EXISTS rollback_product_master_20260711;
DELIMITER $$
CREATE PROCEDURE rollback_product_master_20260711()
BEGIN
    DECLARE EXIT HANDLER FOR SQLEXCEPTION
    BEGIN
        ROLLBACK;
        RESIGNAL;
    END;
    IF (SELECT COUNT(*) FROM backup_inv_product_before_20260711_refresh) <> 167
       OR (SELECT COUNT(*) FROM backup_inv_product_category_before_20260711_refresh) <> 23
       OR (SELECT COUNT(*) FROM backup_inv_supplier_before_20260711_refresh) <> 87 THEN
        SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT = 'rollback failed: backup row counts do not match';
    END IF;

    START TRANSACTION;
    UPDATE inv_product target
    JOIN backup_inv_product_before_20260711_refresh backup ON backup.product_id = target.product_id
    SET {assignments(product_columns)}
    WHERE target.shop_dept_id = {shop_dept_id};

    UPDATE inv_product_category target
    JOIN backup_inv_product_category_before_20260711_refresh backup ON backup.category_id = target.category_id
    SET {assignments(category_columns)}
    WHERE target.shop_dept_id = {shop_dept_id};

    UPDATE inv_supplier target
    JOIN backup_inv_supplier_before_20260711_refresh backup ON backup.supplier_id = target.supplier_id
    SET {assignments(supplier_columns)}
    WHERE target.shop_dept_id = {shop_dept_id};
    COMMIT;
END$$
DELIMITER ;
CALL rollback_product_master_20260711();
DROP PROCEDURE rollback_product_master_20260711;
"""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path("sql/erp_inventory_product_master_20260711.sql"))
    parser.add_argument(
        "--rollback-output",
        type=Path,
        default=Path("sql/erp_inventory_product_master_20260711_rollback.sql"),
    )
    parser.add_argument("--database", default="BossERP_NEW")
    parser.add_argument("--shop-dept-id", type=int, default=100)
    args = parser.parse_args()

    products, hidden_rows = parse_workbook(args.workbook)
    args.output.write_text(
        build_refresh_sql(products, hidden_rows, args.workbook, args.database, args.shop_dept_id),
        encoding="utf-8",
    )
    args.rollback_output.write_text(
        build_rollback_sql(args.database, args.shop_dept_id),
        encoding="utf-8",
    )
    print(f"generated {args.output}: visible_products={len(products)} hidden_products={len(hidden_rows)}")
    print(f"generated {args.rollback_output}")


if __name__ == "__main__":
    main()
