#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path
from typing import Iterable, List, Tuple

import openpyxl

import erp_employee_importer as importer


def sql_quote(value: str) -> str:
    return "'" + (value or "").replace("\\", "\\\\").replace("'", "''") + "'"


def values_rows(rows: Iterable[Tuple[str, str, str, str, str, str]]) -> str:
    return ",\n".join(
        "(" + ", ".join(sql_quote(item) for item in row) + ")"
        for row in rows
    )


def missing_phone_rows(missing_excel: Path, plan: importer.ImportPlan) -> List[Tuple[str, str, str, str, str, str]]:
    wb = openpyxl.load_workbook(missing_excel, data_only=True)
    ws = wb["无手机号明细"]
    headers = [cell.value for cell in ws[1]]
    rows: List[Tuple[str, str, str, str, str, str]] = []
    for raw_row in ws.iter_rows(min_row=2, values_only=True):
        data = dict(zip(headers, raw_row))
        name = importer.clean_name(data.get("姓名", ""))
        if not name:
            continue
        post_name = importer.normalize_post(data.get("职位", ""))
        role_name = plan.role_name_for_post(post_name) if post_name else ""
        rows.append(("missing_name", name, name, post_name, role_name, "无手机号员工明细_含负责人.xlsx"))
    return rows


def build_expected_rows(employee_excel: Path, missing_excel: Path) -> List[Tuple[str, str, str, str, str, str]]:
    plan = importer.build_plan_from_workbook(str(employee_excel))
    rows: List[Tuple[str, str, str, str, str, str]] = []
    for phone, user in sorted(plan.users_by_phone.items()):
        rows.append(("phone", phone, user.name, user.post_name, plan.role_name_for_post(user.post_name), "员工Excel"))
    rows.extend(missing_phone_rows(missing_excel, plan))
    return rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--employee-excel", default=importer.DEFAULT_EXCEL_PATH)
    parser.add_argument("--missing-excel", default="无手机号员工明细_含负责人.xlsx")
    args = parser.parse_args()

    rows = build_expected_rows(Path(args.employee_excel), Path(args.missing_excel))
    print("#!/usr/bin/env bash")
    print("set -euo pipefail")
    print('DB="BossERP_NEW"')
    print('MYSQL_ROOT_PASS="$(cat /root/.erp-mysql-root-pass)"')
    print('mysql --protocol=tcp -uroot -p"${MYSQL_ROOT_PASS}" "${DB}" --batch --raw <<\'SQL\'')
    print("CREATE TEMPORARY TABLE expected_user_role_post (")
    print("  match_type varchar(32) not null,")
    print("  match_key varchar(128) not null,")
    print("  expected_name varchar(128) not null,")
    print("  expected_post varchar(128) not null,")
    print("  expected_role varchar(128) not null,")
    print("  source_name varchar(128) not null")
    print(");")
    print("INSERT INTO expected_user_role_post(match_type, match_key, expected_name, expected_post, expected_role, source_name) VALUES")
    print(values_rows(rows) + ";")
    print("CREATE TEMPORARY TABLE expected_user_role_post_lookup AS SELECT * FROM expected_user_role_post;")
    print(r"""
select 'expected_counts' as section;
select count(*) as expected_rows,
       sum(match_type='phone') as expected_phone_rows,
       sum(match_type='missing_name') as expected_missing_name_rows
from expected_user_role_post;

select 'role_post_mismatches' as section;
select e.source_name, e.match_type, e.match_key, e.expected_name,
       coalesce(u.user_id, 0) as user_id,
       coalesce(u.nick_name, '<missing>') as actual_name,
       e.expected_post,
       coalesce(a.actual_posts, '<none>') as actual_posts,
       e.expected_role,
       coalesce(a.actual_roles, '<none>') as actual_roles,
       coalesce(a.post_count, 0) as post_count,
       coalesce(a.role_count, 0) as role_count
from expected_user_role_post e
left join sys_user u
  on u.del_flag='0'
 and (
      (e.match_type='phone' and (u.phonenumber=e.match_key or u.user_name=e.match_key))
      or (e.match_type='missing_name' and u.create_by='missing_phone_reseed' and u.nick_name=e.match_key)
 )
left join (
    select u.user_id,
           group_concat(distinct p.post_name order by p.post_sort, p.post_id separator '|') as actual_posts,
           group_concat(distinct r.role_name order by r.role_sort, r.role_id separator '|') as actual_roles,
           count(distinct p.post_id) as post_count,
           count(distinct r.role_id) as role_count
    from sys_user u
    left join sys_user_post up on up.user_id=u.user_id
    left join sys_post p on p.post_id=up.post_id
    left join sys_user_role ur on ur.user_id=u.user_id
    left join sys_role r on r.role_id=ur.role_id
    where u.del_flag='0'
    group by u.user_id
) a on a.user_id=u.user_id
where u.user_id is null
   or coalesce(a.post_count, 0) <> 1
   or coalesce(a.role_count, 0) <> 1
   or (e.expected_post <> '' and coalesce(a.actual_posts, '') <> e.expected_post)
   or (e.expected_role <> '' and coalesce(a.actual_roles, '') <> e.expected_role)
order by e.source_name, e.match_type, e.expected_name, e.match_key;

select 'active_import_users_not_in_expected' as section;
select u.user_id, u.nick_name, u.user_name, u.phonenumber, u.create_by,
       coalesce(a.actual_posts, '<none>') as actual_posts,
       coalesce(a.actual_roles, '<none>') as actual_roles
from sys_user u
left join (
    select u.user_id,
           group_concat(distinct p.post_name order by p.post_sort, p.post_id separator '|') as actual_posts,
           group_concat(distinct r.role_name order by r.role_sort, r.role_id separator '|') as actual_roles
    from sys_user u
    left join sys_user_post up on up.user_id=u.user_id
    left join sys_post p on p.post_id=up.post_id
    left join sys_user_role ur on ur.user_id=u.user_id
    left join sys_role r on r.role_id=ur.role_id
    where u.del_flag='0'
    group by u.user_id
) a on a.user_id=u.user_id
where u.del_flag='0'
  and u.create_by in ('employee_xls_import','missing_phone_reseed')
  and not exists (
      select 1 from expected_user_role_post_lookup e
      where (e.match_type='phone' and (u.phonenumber=e.match_key or u.user_name=e.match_key))
         or (e.match_type='missing_name' and u.create_by='missing_phone_reseed' and u.nick_name=e.match_key)
  )
order by u.create_by, u.nick_name, u.user_id;
SQL
""")


if __name__ == "__main__":
    main()
