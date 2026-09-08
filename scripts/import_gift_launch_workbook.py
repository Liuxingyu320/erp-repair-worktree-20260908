#!/usr/bin/env python3
"""Import the July 2026 gift/OE launch workbook into local BossERP_NEW.

The importer is intentionally deterministic and idempotent:

* duplicated regular rows copied into the spring sheet are merged;
* visually distinct packaging rows receive a packaging suffix;
* every embedded image is extracted once by SHA-256, with the first image used
  as the catalog main image and remaining URLs retained in the remark;
* stable source-based item codes make a rerun update the same records.

Run without --apply for a read-only validation pass.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
from collections import OrderedDict, defaultdict
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


DEFAULT_WORKBOOK = Path(
    "/Users/liuxingyu/Library/Containers/com.tencent.xinWeChat/Data/Documents/"
    "xwechat_files/wxid_el0e9x2hyy9022_b7f6/msg/file/2026-07/礼盒上线.xlsx"
)
REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_IMAGE_DIR = REPO_ROOT / "uploadPath/public/2026/08/03/gift-launch"
DEFAULT_PUBLIC_URL = "http://localhost:8080/file/public/2026/08/03/gift-launch"
DEFAULT_DATABASE = "BossERP_NEW"
IMPORT_BATCH = "gift-launch-20260803"
UI_IMAGE_UPLOAD_LIMIT = 5 * 1024 * 1024

GIFT_CATEGORY_CODES = OrderedDict(
    [
        ("乌龙茶", "WLC"),
        ("红茶", "HC"),
        ("小青柑", "XQG"),
        ("白茶", "BC"),
        ("中秋礼盒", "ZQLH"),
        ("春茶礼盒", "CCLH"),
    ]
)

OE_CATEGORY_CODES = OrderedDict(
    [
        ("瓷器类", "CQL"),
        ("玻璃类", "BLL"),
        ("茶具辅料", "CJFL"),
        ("电器类", "DQL"),
        ("制服", "ZF"),
    ]
)

# These rows share product/spec data but contain visibly different package art.
# The names are read from the package fronts in the workbook images.
REGULAR_PACKAGE_STYLE = {
    5: "如愿",
    6: "如愿",
    7: "如愿",
    8: "万事顺",
    9: "万事顺",
    10: "万事顺",
    12: "马上有礼",
    13: "马上有礼",
    25: "如愿",
    26: "万事顺",
    35: "如愿",
    36: "万事顺",
    37: "马上有礼",
    40: "福禄",
    41: "满庭芳",
    43: "如愿",
    44: "万事顺",
    45: "马上有礼",
}


@dataclass(frozen=True)
class ImageAsset:
    digest: str
    extension: str
    data: bytes

    @property
    def filename(self) -> str:
        return f"{self.digest[:24]}.{self.extension}"


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def decimal_value(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    return Decimal(str(value)).quantize(Decimal("0.01"))


def normalized(value: Any) -> str:
    text = clean_text(value) or ""
    return "".join(text.split())


def merged_value(ws, row: int, column: int) -> Any:
    value = ws.cell(row, column).value
    if value not in (None, ""):
        return value
    for cell_range in ws.merged_cells.ranges:
        if (
            cell_range.min_row <= row <= cell_range.max_row
            and cell_range.min_col <= column <= cell_range.max_col
        ):
            return ws.cell(cell_range.min_row, cell_range.min_col).value
    return value


def load_sheet_images(workbook) -> dict[str, dict[int, list[ImageAsset]]]:
    result: dict[str, dict[int, list[ImageAsset]]] = {}
    for ws in workbook.worksheets:
        row_images: dict[int, list[ImageAsset]] = defaultdict(list)
        for image in ws._images:
            if not hasattr(image.anchor, "_from"):
                continue
            data = image._data()
            digest = hashlib.sha256(data).hexdigest()
            image_format = (getattr(image, "format", "png") or "png").lower()
            extension = "jpg" if image_format in {"jpg", "jpeg"} else "png"
            row_images[image.anchor._from.row + 1].append(
                ImageAsset(digest=digest, extension=extension, data=data)
            )
        result[ws.title] = dict(row_images)
    return result


def images_for_row(ws, row_images: dict[int, list[ImageAsset]], row: int, image_column: int) -> list[ImageAsset]:
    if row_images.get(row):
        return row_images[row]
    for cell_range in ws.merged_cells.ranges:
        if (
            cell_range.min_row <= row <= cell_range.max_row
            and cell_range.min_col <= image_column <= cell_range.max_col
        ):
            return row_images.get(cell_range.min_row, [])
    return []


def image_urls(assets: Iterable[ImageAsset], public_url: str) -> list[str]:
    prefix = public_url.rstrip("/")
    return [f"{prefix}/{asset.filename}" for asset in assets]


def build_remark(parts: Iterable[str | None]) -> str | None:
    text = "；".join(part.strip("； ") for part in parts if part and part.strip("； "))
    if not text:
        return None
    if len(text) <= 500:
        return text
    return text[:497] + "..."


def parse_gifts(value_book, image_book, sheet_images, public_url: str) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    selections = [
        ("常规礼盒上新", range(2, 49), "CG"),
        ("中秋礼盒", range(2, 15), "ZQ"),
        # Rows 15-42 are exact image/data copies of 常规礼盒上新!21:48.
        ("春茶礼盒", range(2, 15), "CC"),
    ]
    for sheet_name, rows, source_code in selections:
        ws = value_book[sheet_name]
        image_ws = image_book[sheet_name]
        last_category: str | None = None
        for row in rows:
            category = clean_text(merged_value(ws, row, 2))
            if category:
                last_category = category
            else:
                category = last_category
            name = clean_text(merged_value(ws, row, 3))
            if not name:
                continue
            style = REGULAR_PACKAGE_STYLE.get(row) if sheet_name == "常规礼盒上新" else None
            catalog_name = f"{name}（{style}）" if style else name
            description = clean_text(merged_value(ws, row, 6))
            source_refs = [f"{sheet_name}!{row}"]
            if sheet_name == "常规礼盒上新" and 21 <= row <= 48:
                spring_row = row - 6
                origin = clean_text(merged_value(value_book["春茶礼盒"], spring_row, 6))
                if origin and normalized(origin) != normalized(description):
                    description = f"产地：{origin}\n{description or ''}".strip()
                source_refs.append(f"春茶礼盒!{spring_row}")
            unit = clean_text(merged_value(ws, row, 7))
            if sheet_name == "春茶礼盒" and not unit:
                unit = "盒"
            assets = images_for_row(image_ws, sheet_images[sheet_name], row, 13)
            urls = image_urls(assets, public_url)
            raw_image_cell = clean_text(merged_value(ws, row, 13))
            original_remark = clean_text(merged_value(ws, row, 12))
            remark = build_remark(
                [
                    f"导入批次：{IMPORT_BATCH}",
                    f"来源：{','.join(source_refs)}",
                    f"包装款式：{style}" if style else None,
                    f"原备注：{original_remark}" if original_remark else None,
                    f"图片列原值：{raw_image_cell}" if raw_image_cell else None,
                    f"附图：{','.join(urls[1:])}" if len(urls) > 1 else None,
                ]
            )
            category_code = GIFT_CATEGORY_CODES.get(category or "")
            if not category_code:
                raise ValueError(f"{sheet_name}!{row} 未识别礼盒分类：{category!r}")
            records.append(
                {
                    "gift_code": f"LH-{category_code}-{source_code}{row:03d}",
                    "category": category,
                    "category_code": category_code,
                    "gift_name": catalog_name,
                    "grade": clean_text(merged_value(ws, row, 4)),
                    "spec": clean_text(merged_value(ws, row, 5)),
                    "product_description": description,
                    "replenishment_unit": unit,
                    "cost_price": decimal_value(merged_value(ws, row, 8)),
                    "guide_price_1": decimal_value(merged_value(ws, row, 9)),
                    "guide_price_2": decimal_value(merged_value(ws, row, 10)),
                    "supplier_name": clean_text(merged_value(ws, row, 11)),
                    "image_url": urls[0] if urls else None,
                    "remark": remark,
                    "assets": assets,
                    "source": source_refs,
                }
            )
    return records


def parse_oe(value_book, image_book, sheet_images, public_url: str) -> list[dict[str, Any]]:
    ws = value_book["OE配置表"]
    image_ws = image_book["OE配置表"]
    records: list[dict[str, Any]] = []
    last_category: str | None = None
    for row in range(2, 34):
        category = clean_text(merged_value(ws, row, 2))
        if category:
            last_category = category
        else:
            category = last_category
        oe_type_name = clean_text(merged_value(ws, row, 3))
        oe_item_name = clean_text(merged_value(ws, row, 4)) or oe_type_name
        if not oe_item_name:
            continue
        category_code = OE_CATEGORY_CODES.get(category or "")
        if not category_code:
            raise ValueError(f"OE配置表!{row} 未识别OE分类：{category!r}")
        assets = images_for_row(image_ws, sheet_images["OE配置表"], row, 11)
        urls = image_urls(assets, public_url)
        source_code = clean_text(merged_value(ws, row, 5))
        remark = build_remark(
            [
                f"导入批次：{IMPORT_BATCH}",
                f"来源：OE配置表!{row}",
                f"附图：{','.join(urls[1:])}" if len(urls) > 1 else None,
            ]
        )
        records.append(
            {
                "oe_item_code": source_code or f"OE-{category_code}-{row:03d}",
                "category": category,
                "category_code": category_code,
                "oe_type_name": oe_type_name,
                "oe_item_name": oe_item_name,
                "item_description": clean_text(merged_value(ws, row, 6)),
                "order_unit": clean_text(merged_value(ws, row, 7)),
                "cost_price": decimal_value(merged_value(ws, row, 8)),
                "supplier_name": clean_text(merged_value(ws, row, 9)),
                "supplier_phone": clean_text(merged_value(ws, row, 10)),
                "image_url": urls[0] if urls else None,
                "remark": remark,
                "assets": assets,
                "source": [f"OE配置表!{row}"],
            }
        )
    return records


def validate_records(gifts: list[dict[str, Any]], oe_items: list[dict[str, Any]]) -> None:
    if len(gifts) != 73:
        raise ValueError(f"礼盒记录应为73条，实际为{len(gifts)}条")
    if len(oe_items) != 32:
        raise ValueError(f"OE记录应为32条，实际为{len(oe_items)}条")
    gift_codes = {record["gift_code"] for record in gifts}
    oe_codes = {record["oe_item_code"] for record in oe_items}
    if len(gift_codes) != len(gifts):
        raise ValueError("礼盒编码不唯一")
    if len(oe_codes) != len(oe_items):
        raise ValueError("OE编码不唯一")
    natural_keys = {
        (
            normalized(record["category"]),
            normalized(record["gift_name"]),
            normalized(record["grade"]),
            normalized(record["spec"]),
        )
        for record in gifts
    }
    if len(natural_keys) != len(gifts):
        raise ValueError(
            f"包装款式命名后仍有礼盒自然键重复：{len(gifts) - len(natural_keys)}条"
        )
    if any(len(record["gift_name"]) > 128 for record in gifts):
        raise ValueError("礼盒名称超过128字符")
    if any(record["product_description"] and len(record["product_description"]) > 1000 for record in gifts):
        raise ValueError("礼盒描述超过1000字符")
    if any(record["remark"] and len(record["remark"]) > 500 for record in gifts + oe_items):
        raise ValueError("备注超过500字符")
    if sum(bool(record["image_url"]) for record in gifts) != len(gifts):
        raise ValueError("存在没有主图的礼盒记录")


def all_assets(records: Iterable[dict[str, Any]]) -> OrderedDict[str, ImageAsset]:
    assets: OrderedDict[str, ImageAsset] = OrderedDict()
    for record in records:
        for asset in record["assets"]:
            assets.setdefault(asset.digest, asset)
    return assets


def write_assets(records: Iterable[dict[str, Any]], output_dir: Path) -> int:
    assets = all_assets(records)
    output_dir.mkdir(parents=True, exist_ok=True)
    written = 0
    for asset in assets.values():
        target = output_dir / asset.filename
        if target.exists():
            if hashlib.sha256(target.read_bytes()).hexdigest() != asset.digest:
                raise ValueError(f"图片目标已存在但内容不同：{target}")
            continue
        target.write_bytes(asset.data)
        written += 1
    return written


def sql_text(value: Any) -> str:
    if value is None:
        return "NULL"
    data = str(value).encode("utf-8")
    return f"CONVERT(0x{data.hex()} USING utf8mb4)"


def sql_decimal(value: Decimal | None) -> str:
    return "NULL" if value is None else format(value, "f")


def gift_category_sql(category: str, code: str, order_num: int) -> str:
    return f"""
INSERT INTO inv_gift_category
    (parent_id, ancestors, category_name, category_code, order_num, status, del_flag, create_by, create_time)
VALUES
    (0, '0', {sql_text(category)}, {sql_text(code)}, {order_num}, '0', '0', {sql_text(IMPORT_BATCH)}, NOW())
ON DUPLICATE KEY UPDATE
    category_name = VALUES(category_name), status = '0', del_flag = '0',
    update_by = {sql_text(IMPORT_BATCH)}, update_time = NOW();
"""


def oe_category_sql(category: str, code: str, order_num: int) -> str:
    return f"""
INSERT INTO inv_oe_category
    (parent_id, ancestors, category_name, category_code, order_num, status, del_flag, create_by, create_time)
VALUES
    (0, '0', {sql_text(category)}, {sql_text(code)}, {order_num}, '0', '0', {sql_text(IMPORT_BATCH)}, NOW())
ON DUPLICATE KEY UPDATE
    category_name = VALUES(category_name), status = '0', del_flag = '0',
    update_by = {sql_text(IMPORT_BATCH)}, update_time = NOW();
"""


def gift_sql(record: dict[str, Any]) -> str:
    return f"""
INSERT INTO inv_gift_box
    (gift_code, category_id, gift_name, grade, spec, product_description,
     replenishment_unit, cost_price, guide_price_1, guide_price_2, supplier_name,
     image_url, status, del_flag, create_by, create_time, remark)
VALUES
    ({sql_text(record['gift_code'])},
     (SELECT category_id FROM inv_gift_category WHERE category_code = {sql_text(record['category_code'])} AND del_flag = '0' LIMIT 1),
     {sql_text(record['gift_name'])}, {sql_text(record['grade'])}, {sql_text(record['spec'])},
     {sql_text(record['product_description'])}, {sql_text(record['replenishment_unit'])},
     {sql_decimal(record['cost_price'])}, {sql_decimal(record['guide_price_1'])},
     {sql_decimal(record['guide_price_2'])}, {sql_text(record['supplier_name'])},
     {sql_text(record['image_url'])}, '0', '0', {sql_text(IMPORT_BATCH)}, NOW(), {sql_text(record['remark'])})
ON DUPLICATE KEY UPDATE
    category_id = VALUES(category_id), gift_name = VALUES(gift_name), grade = VALUES(grade),
    spec = VALUES(spec), product_description = VALUES(product_description),
    replenishment_unit = VALUES(replenishment_unit), cost_price = VALUES(cost_price),
    guide_price_1 = VALUES(guide_price_1), guide_price_2 = VALUES(guide_price_2),
    supplier_name = VALUES(supplier_name), image_url = VALUES(image_url), status = '0',
    del_flag = '0', update_by = {sql_text(IMPORT_BATCH)}, update_time = NOW(), remark = VALUES(remark);
"""


def oe_sql(record: dict[str, Any]) -> str:
    return f"""
INSERT INTO inv_oe_item
    (oe_item_code, category_id, oe_type_name, oe_item_name, item_description,
     order_unit, cost_price, supplier_name, supplier_phone, image_url,
     status, del_flag, create_by, create_time, remark)
VALUES
    ({sql_text(record['oe_item_code'])},
     (SELECT category_id FROM inv_oe_category WHERE category_code = {sql_text(record['category_code'])} AND del_flag = '0' LIMIT 1),
     {sql_text(record['oe_type_name'])}, {sql_text(record['oe_item_name'])},
     {sql_text(record['item_description'])}, {sql_text(record['order_unit'])},
     {sql_decimal(record['cost_price'])}, {sql_text(record['supplier_name'])},
     {sql_text(record['supplier_phone'])}, {sql_text(record['image_url'])},
     '0', '0', {sql_text(IMPORT_BATCH)}, NOW(), {sql_text(record['remark'])})
ON DUPLICATE KEY UPDATE
    category_id = VALUES(category_id), oe_type_name = VALUES(oe_type_name),
    oe_item_name = VALUES(oe_item_name), item_description = VALUES(item_description),
    order_unit = VALUES(order_unit), cost_price = VALUES(cost_price),
    supplier_name = VALUES(supplier_name), supplier_phone = VALUES(supplier_phone),
    image_url = VALUES(image_url), status = '0', del_flag = '0',
    update_by = {sql_text(IMPORT_BATCH)}, update_time = NOW(), remark = VALUES(remark);
"""


def mysql_command(database: str, sql: str, capture: bool = False) -> str:
    mysql = shutil.which("mysql")
    if not mysql:
        raise RuntimeError("未找到 mysql 客户端")
    command = [
        mysql,
        "--protocol=TCP",
        "-h127.0.0.1",
        "-uroot",
        "--default-character-set=utf8mb4",
        "--binary-mode",
        database,
    ]
    if capture:
        command.extend(["-N", "-B"])
    completed = subprocess.run(
        command,
        input=sql,
        text=True,
        capture_output=True,
        check=False,
    )
    if completed.returncode != 0:
        raise RuntimeError(completed.stderr.strip() or "mysql 执行失败")
    return completed.stdout.strip()


def require_schema(database: str) -> None:
    result = mysql_command(
        database,
        """
SELECT COUNT(*)
FROM information_schema.columns
WHERE table_schema = DATABASE()
  AND table_name = 'inv_gift_box'
  AND column_name IN ('cost_price', 'supplier_name');
""",
        capture=True,
    )
    if result != "2":
        raise RuntimeError("inv_gift_box 尚未完成成本价/供应商字段迁移")


def apply_records(database: str, gifts: list[dict[str, Any]], oe_items: list[dict[str, Any]]) -> None:
    require_schema(database)
    unrelated = mysql_command(
        database,
        f"""
SELECT
  (SELECT COUNT(*) FROM inv_gift_box WHERE del_flag = '0' AND COALESCE(create_by, '') <> {sql_text(IMPORT_BATCH)}) +
  (SELECT COUNT(*) FROM inv_oe_item WHERE del_flag = '0' AND COALESCE(create_by, '') <> {sql_text(IMPORT_BATCH)});
""",
        capture=True,
    )
    if unrelated != "0":
        raise RuntimeError(f"目标表已有{unrelated}条非本批次数据，已停止以避免覆盖")

    statements = ["SET NAMES utf8mb4;", "START TRANSACTION;"]
    selected_gift_categories = OrderedDict((record["category"], record["category_code"]) for record in gifts)
    selected_oe_categories = OrderedDict((record["category"], record["category_code"]) for record in oe_items)
    statements.extend(
        gift_category_sql(category, code, index)
        for index, (category, code) in enumerate(selected_gift_categories.items(), start=1)
    )
    statements.extend(
        oe_category_sql(category, code, index)
        for index, (category, code) in enumerate(selected_oe_categories.items(), start=1)
    )
    statements.extend(gift_sql(record) for record in gifts)
    statements.extend(oe_sql(record) for record in oe_items)
    statements.append("COMMIT;")
    mysql_command(database, "\n".join(statements))

    gift_codes = ",".join(sql_text(record["gift_code"]) for record in gifts) or "NULL"
    oe_codes = ",".join(sql_text(record["oe_item_code"]) for record in oe_items) or "NULL"
    counts = mysql_command(
        database,
        f"""
SELECT COUNT(*) FROM inv_gift_box WHERE del_flag = '0' AND gift_code IN ({gift_codes});
SELECT COUNT(*) FROM inv_oe_item WHERE del_flag = '0' AND oe_item_code IN ({oe_codes});
""",
        capture=True,
    ).splitlines()
    expected = [str(len(gifts)), str(len(oe_items))]
    if counts != expected:
        raise RuntimeError(f"写入后数量校验失败：实际{counts}，期望{expected}")


def summarize(gifts: list[dict[str, Any]], oe_items: list[dict[str, Any]]) -> dict[str, Any]:
    assets = all_assets(gifts + oe_items)
    return {
        "batch": IMPORT_BATCH,
        "gift_records": len(gifts),
        "oe_records": len(oe_items),
        "gift_categories": sorted({record["category"] for record in gifts}),
        "oe_categories": sorted({record["category"] for record in oe_items}),
        "gift_missing_grade": sum(not record["grade"] for record in gifts),
        "gift_missing_supplier": sum(not record["supplier_name"] for record in gifts),
        "oe_missing_cost": sum(record["cost_price"] is None for record in oe_items),
        "oe_missing_supplier": sum(not record["supplier_name"] for record in oe_items),
        "gift_with_main_image": sum(bool(record["image_url"]) for record in gifts),
        "oe_with_main_image": sum(bool(record["image_url"]) for record in oe_items),
        "unique_image_files": len(assets),
        "image_bytes": sum(len(asset.data) for asset in assets.values()),
        "images_over_ui_upload_limit": sum(
            len(asset.data) > UI_IMAGE_UPLOAD_LIMIT for asset in assets.values()
        ),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--database", default=DEFAULT_DATABASE)
    parser.add_argument("--image-dir", type=Path, default=DEFAULT_IMAGE_DIR)
    parser.add_argument("--public-url", default=DEFAULT_PUBLIC_URL)
    parser.add_argument("--gift-limit", type=int)
    parser.add_argument("--oe-limit", type=int)
    parser.add_argument("--apply", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if not args.workbook.is_file():
        raise FileNotFoundError(args.workbook)
    value_book = load_workbook(args.workbook, data_only=True, read_only=False)
    image_book = load_workbook(args.workbook, data_only=False, read_only=False)
    sheet_images = load_sheet_images(image_book)
    gifts = parse_gifts(value_book, image_book, sheet_images, args.public_url)
    oe_items = parse_oe(value_book, image_book, sheet_images, args.public_url)
    validate_records(gifts, oe_items)
    summary = summarize(gifts, oe_items)

    selected_gifts = gifts[: args.gift_limit] if args.gift_limit is not None else gifts
    selected_oe = oe_items[: args.oe_limit] if args.oe_limit is not None else oe_items
    summary["selected_gift_records"] = len(selected_gifts)
    summary["selected_oe_records"] = len(selected_oe)
    summary["mode"] = "apply" if args.apply else "dry-run"

    if args.apply:
        summary["written_image_files"] = write_assets(selected_gifts + selected_oe, args.image_dir)
        apply_records(args.database, selected_gifts, selected_oe)
    print(json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        print(f"IMPORT_FAILED: {exc}", file=sys.stderr)
        sys.exit(1)
