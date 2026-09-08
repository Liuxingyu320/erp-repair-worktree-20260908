#!/usr/bin/env python3
"""Reconcile the authoritative Beijing 0716 workbook into BossERP_NEW.

The workbook is a contract-signing batch, so this utility only creates or
updates the 27 employees present on the main sheet.  It deliberately does not
infer resignation/deletion from absence in the workbook.
"""

from __future__ import annotations

import argparse
import csv
import dataclasses
import datetime as dt
import decimal
import hashlib
import os
import re
import secrets
import string
import subprocess
import sys
from collections import defaultdict
from pathlib import Path
from typing import Iterable, Optional

import openpyxl
import bcrypt


MAIN_SHEET = "北京区域 (2)"
SIGN_SHEET = "签约数据"
EXPECTED_MAIN_COUNT = 27
EXPECTED_EXISTING_COUNT = 21
EXPECTED_NEW_COUNT = 6
EXPECTED_NEW_NAMES = {"李曼", "马毓谦", "任艳琪", "何顺琪", "刘捷", "苏余玉"}
PRESERVED_CURRENT_NAMES = {"孙静芳", "张影影", "谷佳璐", "刘心雨", "赫亚茹"}
OPERATOR = "employee_xlsx_0716"
SALARY_VERSION = "北京区域0716-2"
LOCK_NAME = "erp:employee:beijing:0716"
PROCEDURE_NAME = "tmp_reconcile_beijing_0716_employees"
LEGAL_ENTITY_CODE = "ZS_MINGHUI"
LEGAL_ENTITY_NAME = "舟山茗汇文化传播有限公司"
LEGAL_REPRESENTATIVE = "杜翠香"
TEMPORARY_PASSWORD_HOURS = 24

CONTRACT_TYPE_CODES = {
    "劳动合同": "LABOR_CONTRACT",
    "劳务合同": "SERVICE_CONTRACT",
}
CONTRACT_TERM_CODES = {
    "固定期限": "FIXED_TERM",
    "无固定期限": "OPEN_ENDED",
}
SOCIAL_TYPE_CODES = {
    "有": "SOCIAL_INSURED",
    "有社保": "SOCIAL_INSURED",
    "无": "SOCIAL_UNINSURED",
    "无社保": "SOCIAL_UNINSURED",
}
EMPLOYEE_CATEGORY_CODES = {
    "全职": "FULL_TIME",
    "正式员工": "FULL_TIME",
    "实习": "INTERN",
    "实习生": "INTERN",
    "退休返聘": "RETIRED_REHIRE",
}
SUPPORTED_EMPLOYEE_STATUSES = {"待入职", "试用", "正式", "待离职", "离职", "停薪留职"}
STATUS_AS_CATEGORY = {"退休返聘", "实习生"}
REQUIRED_CATEGORY_DICTIONARY = {
    "FULL_TIME": (1, "正式员工"),
    "INTERN": (2, "实习生"),
    "RETIRED_REHIRE": (3, "退休返聘"),
}


POST_ROLE = {
    "实习生": "实习生",
    "茶艺师": "茶艺师",
    "店长助理": "店长助理",
    "店长": "店长",
    "驻店经理": "驻店经理",
    "区域运营总监": "区域运营总监",
}

# Column O is the workbook's target store grouping.  Columns L onward are an
# older employee export carried in the same sheet for reference, so they must
# not override the target grouping.  Existing and new employees use this map.
SOURCE_STORE_ORG_SUFFIXES = {
    "北京柏悦": (
        ("金英灵韵", "北京区域", "北京区域运营", "北京柏悦"),
        [("金英灵韵", "北京区域", "北京区域运营", "北京柏悦")],
    ),
    "北京丽思卡尔顿": (
        ("金英灵韵", "北京区域", "北京区域运营", "北京丽思卡尔顿"),
        [("金英灵韵", "北京区域", "北京区域运营", "北京丽思卡尔顿")],
    ),
    "北京瑞吉": (
        ("金英灵韵", "北京区域", "北京区域运营", "北京瑞吉"),
        [("金英灵韵", "北京区域", "北京区域运营", "北京瑞吉")],
    ),
    "华彬费尔蒙": (
        ("金英灵韵", "北京区域", "北京区域运营", "费尔蒙"),
        [("金英灵韵", "北京区域", "北京区域运营", "费尔蒙")],
    ),
    "通州皇冠": (
        ("金英灵韵", "北京区域", "北京区域运营", "通州皇冠"),
        [("金英灵韵", "北京区域", "北京区域运营", "通州皇冠")],
    ),
    "北京康莱德": (
        ("金英灵韵", "北京区域", "北京区域运营", "北京康莱德"),
        [("金英灵韵", "北京区域", "北京区域运营", "北京康莱德")],
    ),
    "马会": (
        ("金英灵韵", "北京区域", "北京区域运营", "马会"),
        [("金英灵韵", "北京区域", "北京区域运营", "马会")],
    ),
    "万达文华": (
        ("金英灵韵", "北京区域", "北京区域运营", "万达文华"),
        [("金英灵韵", "北京区域", "北京区域运营", "万达文华")],
    ),
    "西安君悦&淮扬1949": (
        ("金英灵韵", "北京区域", "西安区域运营"),
        [
            ("金英灵韵", "北京区域", "西安区域运营"),
            ("金英灵韵", "北京区域", "西安区域运营", "淮扬1949"),
            ("金英灵韵", "北京区域", "西安区域运营", "西安君悦"),
        ],
    ),
    "西安区域运营": (
        ("金英灵韵", "北京区域", "西安区域运营"),
        [("金英灵韵", "北京区域", "西安区域运营")],
    ),
}


@dataclasses.dataclass(frozen=True)
class DbUser:
    user_id: int
    user_name: str
    name: str
    phone: str
    del_flag: str
    dept_id: Optional[int]
    employee_no: str
    profile_id: Optional[int]
    id_number: str
    old_post: str
    old_status: str


@dataclasses.dataclass
class SourceEmployee:
    row_number: int
    name: str
    phone: str
    post: str
    employee_status: str
    entry_date: Optional[str]
    profile: dict[str, object]
    existing: Optional[DbUser] = None
    dept_id: Optional[int] = None
    scope_dept_ids: list[int] = dataclasses.field(default_factory=list)
    post_id: Optional[int] = None
    post_code: str = ""
    role_id: Optional[int] = None
    source_store: str = ""
    dept_path: tuple[str, ...] = dataclasses.field(default_factory=tuple)
    temporary_password: str = ""
    temporary_password_hash: str = ""
    generated_employee_no: str = ""
    session_event_id: str = ""


class MigrationError(RuntimeError):
    pass


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def phone(value: object) -> str:
    return re.sub(r"\D", "", text(value))


def date_value(value: object) -> Optional[str]:
    if value in (None, ""):
        return None
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m-%d")
    raw = text(value)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return dt.datetime.strptime(raw, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    raise MigrationError(f"unsupported date value: {raw!r}")


def decimal_value(value: object) -> Optional[decimal.Decimal]:
    if value in (None, ""):
        return None
    try:
        return decimal.Decimal(text(value).replace(",", ""))
    except decimal.InvalidOperation as exc:
        raise MigrationError(f"unsupported numeric value: {value!r}") from exc


def sex_code(value: object, id_number: str) -> Optional[str]:
    raw = text(value)
    if raw == "男":
        return "0"
    if raw == "女":
        return "1"
    if re.fullmatch(r"\d{17}[0-9Xx]", id_number):
        return "0" if int(id_number[16]) % 2 else "1"
    return None


def birth_date(value: object, id_number: str) -> Optional[str]:
    direct = date_value(value)
    if direct:
        return direct
    if re.fullmatch(r"\d{17}[0-9Xx]", id_number):
        raw = id_number[6:14]
        try:
            return dt.datetime.strptime(raw, "%Y%m%d").strftime("%Y-%m-%d")
        except ValueError:
            return None
    return None


def nonempty(target: dict[str, object], key: str, value: object) -> None:
    if value is None:
        return
    if isinstance(value, str) and not value.strip():
        return
    target[key] = value


def coded_value(mapping: dict[str, str], value: object, field: str, employee: str) -> str:
    raw = text(value)
    try:
        return mapping[raw]
    except KeyError as exc:
        raise MigrationError(f"unsupported {field} for {employee}: {raw!r}") from exc


def normalize_store(value: object) -> str:
    return re.sub(r"\s+", "", text(value))


def issue_temporary_credentials(employees: list[SourceEmployee], current_sequence: int) -> None:
    alphabet = string.ascii_letters + string.digits + "!@#$%^&*()-=_+"
    next_sequence = current_sequence
    for source in employees:
        if source.existing:
            continue
        next_sequence += 1
        source.generated_employee_no = f"E{next_sequence:05d}"
        # Prefix guarantees the configured upper/lower/digit/special policy;
        # the remaining characters are independently generated per account.
        raw = "Aa9!" + "".join(secrets.choice(alphabet) for _ in range(12))
        source.temporary_password = raw
        source.temporary_password_hash = bcrypt.hashpw(
            raw.encode("utf-8"), bcrypt.gensalt(rounds=12)
        ).decode("ascii")


def write_credentials(path: Path, employees: list[SourceEmployee]) -> None:
    new_employees = [item for item in employees if not item.existing]
    if not new_employees:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        raise MigrationError(f"credential output already exists: {path}")
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(descriptor, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["姓名", "登录账号", "临时密码", "失效规则"])
        for source in new_employees:
            writer.writerow([
                source.name,
                source.generated_employee_no,
                source.temporary_password,
                f"生成后{TEMPORARY_PASSWORD_HOURS}小时失效，首次登录必须改密",
            ])
    path.chmod(0o600)


def merged_values(ws, column: int) -> dict[int, object]:
    result: dict[int, object] = {}
    for row in range(2, ws.max_row + 1):
        result[row] = ws.cell(row=row, column=column).value
    for merged in ws.merged_cells.ranges:
        if merged.min_col <= column <= merged.max_col:
            value = ws.cell(row=merged.min_row, column=merged.min_col).value
            for row in range(merged.min_row, merged.max_row + 1):
                result[row] = value
    return result


def load_workbook(path: Path) -> tuple[list[SourceEmployee], str, list[str]]:
    sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = openpyxl.load_workbook(path, data_only=True)
    if MAIN_SHEET not in workbook.sheetnames or SIGN_SHEET not in workbook.sheetnames:
        raise MigrationError(f"required sheets not found: {MAIN_SHEET}, {SIGN_SHEET}")

    main = workbook[MAIN_SHEET]
    sign = workbook[SIGN_SHEET]
    store_by_row = merged_values(main, 15)
    department_supervisor_by_row = merged_values(main, 16)
    sign_rows: dict[str, tuple] = {}
    for row in sign.iter_rows(min_row=2, values_only=True):
        name = text(row[1])
        if name:
            if name in sign_rows:
                raise MigrationError(f"duplicate signing name: {name}")
            sign_rows[name] = row

    employees: list[SourceEmployee] = []
    main_names: set[str] = set()
    for row_number, row in enumerate(main.iter_rows(min_row=2, values_only=True), start=2):
        name = text(row[1])
        if not name:
            continue
        if name in main_names:
            raise MigrationError(f"duplicate main-sheet name: {name}")
        main_names.add(name)
        signing = sign_rows.get(name)
        if signing is None:
            raise MigrationError(f"missing signing row for {name}")

        main_phone = phone(row[9])
        signing_phone = phone(signing[13])
        if main_phone and signing_phone and main_phone != signing_phone:
            raise MigrationError(f"phone mismatch between sheets for {name}")
        resolved_phone = main_phone or signing_phone
        if not re.fullmatch(r"1\d{10}", resolved_phone):
            raise MigrationError(f"invalid phone for {name}")

        main_post = text(row[8])
        signing_post = text(signing[3])
        # "暑假工" is the employment note; the signing sheet contains the ERP post.
        resolved_post = signing_post if main_post == "暑假工" else main_post
        if resolved_post not in POST_ROLE:
            raise MigrationError(f"unsupported post for {name}: {resolved_post}")
        if main_post != "暑假工" and signing_post and main_post != signing_post:
            raise MigrationError(f"post mismatch between sheets for {name}: {main_post}/{signing_post}")

        id_number = text(signing[11]).upper() or text(row[34]).upper()
        old_id_number = text(row[34]).upper()
        if id_number and old_id_number and id_number != old_id_number:
            raise MigrationError(f"ID mismatch between sheets for {name}")

        source_status = text(row[6])
        if source_status in STATUS_AS_CATEGORY:
            resolved_status = "试用"
            source_category = source_status
        elif source_status in SUPPORTED_EMPLOYEE_STATUSES:
            resolved_status = source_status
            source_category = text(row[27]) or "全职"
        else:
            raise MigrationError(f"unsupported employee status for {name}: {source_status!r}")
        if source_category not in EMPLOYEE_CATEGORY_CODES:
            raise MigrationError(f"unsupported employee category for {name}: {source_category!r}")
        resolved_category = EMPLOYEE_CATEGORY_CODES[source_category]

        resolved_entry_date = date_value(row[10]) or date_value(signing[15])
        profile: dict[str, object] = {}
        nonempty(profile, "employee_status", resolved_status)
        nonempty(profile, "entry_date", resolved_entry_date)
        nonempty(profile, "employee_category", resolved_category)
        nonempty(profile, "job_grade", text(signing[9]) or text(row[32]))
        nonempty(profile, "department_supervisor", text(department_supervisor_by_row[row_number]))
        if id_number:
            nonempty(profile, "id_type", "居民身份证")
            nonempty(profile, "id_number", id_number)
        sex = sex_code(row[37], id_number)
        nonempty(profile, "sex", sex)
        nonempty(profile, "birth_date", birth_date(row[35], id_number))
        nonempty(profile, "ethnicity", text(row[38]))
        nonempty(profile, "registered_residence", text(row[39]))
        nonempty(profile, "current_address", text(row[45]) or text(signing[12]))
        nonempty(profile, "political_status", text(row[46]))
        nonempty(profile, "marital_status", text(row[41]))
        nonempty(profile, "work_start_date", date_value(row[42]))
        nonempty(profile, "highest_education", text(row[52]))
        nonempty(profile, "highest_graduation_school", text(row[53]))
        nonempty(profile, "highest_graduation_date", date_value(row[54]))
        nonempty(profile, "highest_major", text(row[55]))
        nonempty(profile, "emergency_contact", text(row[66]))
        nonempty(profile, "emergency_contact_relation", text(row[67]))
        nonempty(profile, "emergency_contact_phone", phone(row[68]))
        nonempty(profile, "probation_period", text(row[29]))
        nonempty(profile, "actual_regularization_date", date_value(row[30]))
        nonempty(profile, "planned_regularization_date", date_value(row[31]))
        nonempty(profile, "probation_start_date", date_value(signing[17]))
        nonempty(profile, "probation_end_date", date_value(signing[18]))
        nonempty(profile, "contract_start_date", date_value(signing[15]))
        nonempty(profile, "contract_end_date", date_value(signing[16]))
        nonempty(profile, "contract_type", coded_value(
            CONTRACT_TYPE_CODES, signing[8], "contract type", name
        ))
        nonempty(profile, "contract_term", coded_value(
            CONTRACT_TERM_CODES, signing[14], "contract term", name
        ))
        # The signing-sheet city is authoritative. Column X belongs to the old
        # employee export and can describe a prior shop after reassignment.
        nonempty(profile, "work_location", text(signing[10]))
        nonempty(profile, "work_city_level", text(signing[4]))
        nonempty(profile, "attendance_method", text(signing[19]))
        nonempty(profile, "household_type", text(row[44]))
        social_code = coded_value(SOCIAL_TYPE_CODES, signing[2], "social type", name)
        nonempty(profile, "social_type", social_code)
        nonempty(profile, "social_security_type", social_code)
        nonempty(profile, "bank_account", text(row[56]).replace(" ", ""))
        nonempty(profile, "bank_name", text(row[57]))
        signing_entity = text(signing[5])
        if signing_entity != LEGAL_ENTITY_NAME:
            raise MigrationError(f"unexpected legal entity for {name}: {signing_entity!r}")
        nonempty(profile, "legal_entity", signing_entity)
        salary_total = decimal_value(signing[20])
        salary_components = [decimal_value(signing[index]) for index in (21, 22, 23, 24)]
        if salary_total is None or any(value is None for value in salary_components):
            raise MigrationError(f"salary total/components are incomplete for {name}")
        if sum((value for value in salary_components if value is not None), decimal.Decimal(0)) != salary_total:
            raise MigrationError(f"salary components do not equal total for {name}")
        nonempty(profile, "base_salary", salary_components[0])
        nonempty(profile, "post_salary", salary_components[1])
        nonempty(profile, "field_allowance", salary_components[2])
        nonempty(profile, "performance_salary", salary_components[3])
        nonempty(profile, "salary_total", salary_total)
        nonempty(profile, "salary_version", SALARY_VERSION)
        nonempty(profile, "id_card_portrait_status", text(row[74]))
        nonempty(profile, "id_card_emblem_status", text(row[75]))
        nonempty(profile, "education_certificate_status", text(row[76]))
        nonempty(profile, "degree_certificate_status", text(row[77]))
        nonempty(profile, "resignation_certificate_status", text(row[78]))
        nonempty(profile, "employee_photo_status", text(row[79]))
        source_store = normalize_store(store_by_row[row_number])
        if not source_store and name == "苏余玉":
            source_store = "西安区域运营"
        if source_store not in SOURCE_STORE_ORG_SUFFIXES:
            raise MigrationError(f"unsupported target store for {name}: {source_store!r}")
        nonempty(profile, "_source_store", source_store)

        employees.append(SourceEmployee(
            row_number=row_number,
            name=name,
            phone=resolved_phone,
            post=resolved_post,
            employee_status=resolved_status,
            entry_date=resolved_entry_date,
            profile=profile,
            source_store=source_store,
        ))

    if len(employees) != EXPECTED_MAIN_COUNT:
        raise MigrationError(f"expected {EXPECTED_MAIN_COUNT} main rows, got {len(employees)}")
    phones = [item.phone for item in employees]
    if len(set(phones)) != len(phones):
        raise MigrationError("duplicate normalized phone in main sheet")
    ids = [text(item.profile.get("id_number")) for item in employees if item.profile.get("id_number")]
    if len(set(ids)) != len(ids):
        raise MigrationError("duplicate normalized ID number in main sheet")

    sign_only = sorted(set(sign_rows) - main_names)
    if sign_only != ["段继康"] or "实验数据" not in text(sign_rows["段继康"][25]):
        raise MigrationError(f"unexpected signing-only rows: {sign_only}")
    return employees, sha256, sign_only


def sql_quote(value: object) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, decimal.Decimal):
        return format(value, "f")
    if isinstance(value, (int, float)):
        return str(value)
    raw = str(value).replace("\\", "\\\\").replace("'", "''")
    return f"'{raw}'"


def mysql_command(args: argparse.Namespace) -> list[str]:
    command = [
        "mysql", "--protocol=tcp", f"-h{args.host}", f"-P{args.port}",
        f"-u{args.user}", "--default-character-set=utf8mb4",
    ]
    if args.password:
        command.append(f"-p{args.password}")
    command.append(args.database)
    return command


def mysql_query(args: argparse.Namespace, sql: str) -> list[list[str]]:
    command = mysql_command(args) + ["--batch", "--raw", "--skip-column-names", "-e", sql]
    result = subprocess.run(command, text=True, stdout=subprocess.PIPE,
                            stderr=subprocess.PIPE, check=False)
    if result.returncode:
        raise MigrationError(result.stderr.strip() or "mysql query failed")
    return [line.split("\t") for line in result.stdout.splitlines() if line]


def mysql_execute(args: argparse.Namespace, sql: str) -> str:
    result = subprocess.run(mysql_command(args), input=sql, text=True,
                            stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=False)
    if result.returncode:
        raise MigrationError(result.stderr.strip() or "mysql apply failed")
    return result.stdout


def load_db_users(args: argparse.Namespace) -> list[DbUser]:
    rows = mysql_query(args, """
        select u.user_id,u.user_name,u.nick_name,coalesce(u.phonenumber,''),u.del_flag,
               coalesce(cast(u.dept_id as char),''),coalesce(p.employee_no,''),
               coalesce(cast(p.profile_id as char),''),coalesce(p.id_number,''),
               coalesce(group_concat(distinct sp.post_name order by sp.post_sort separator ','),''),
               coalesce(p.employee_status,'')
        from sys_user u
        left join sys_user_profile p on p.user_id=u.user_id
        left join sys_user_post up on up.user_id=u.user_id
        left join sys_post sp on sp.post_id=up.post_id
        group by u.user_id,u.user_name,u.nick_name,u.phonenumber,u.del_flag,u.dept_id,
                 p.employee_no,p.profile_id,p.id_number,p.employee_status
        order by u.user_id
    """)
    return [DbUser(
        user_id=int(row[0]), user_name=row[1], name=row[2], phone=row[3],
        del_flag=row[4], dept_id=int(row[5]) if row[5] else None,
        employee_no=row[6], profile_id=int(row[7]) if row[7] else None,
        id_number=row[8].upper(), old_post=row[9], old_status=row[10],
    ) for row in rows]


def dept_paths(args: argparse.Namespace) -> dict[tuple[str, ...], int]:
    rows = mysql_query(args, "select dept_id,parent_id,dept_name from sys_dept where del_flag='0'")
    nodes = {int(row[0]): (int(row[1]), row[2]) for row in rows}
    cache: dict[int, tuple[str, ...]] = {}

    def path_for(dept_id: int, seen: set[int]) -> tuple[str, ...]:
        if dept_id in cache:
            return cache[dept_id]
        if dept_id in seen:
            raise MigrationError("cycle in sys_dept")
        parent_id, name = nodes[dept_id]
        if parent_id == 0 or parent_id not in nodes:
            result = (name,)
        else:
            result = path_for(parent_id, seen | {dept_id}) + (name,)
        cache[dept_id] = result
        return result

    return {path_for(dept_id, set()): dept_id for dept_id in nodes}


def resolve_dept(paths: dict[tuple[str, ...], int], suffix: tuple[str, ...]) -> int:
    matches = [dept_id for path, dept_id in paths.items()
               if len(path) >= len(suffix) and path[-len(suffix):] == suffix]
    if len(matches) != 1:
        raise MigrationError(f"department path suffix is not unique: {suffix} -> {matches}")
    return matches[0]


def attach_database_plan(args: argparse.Namespace, employees: list[SourceEmployee]) -> tuple[int, list[int]]:
    users = load_db_users(args)
    by_phone: dict[str, list[DbUser]] = defaultdict(list)
    by_name: dict[str, list[DbUser]] = defaultdict(list)
    by_id: dict[str, list[DbUser]] = defaultdict(list)
    for item in users:
        if item.phone:
            by_phone[phone(item.phone)].append(item)
        if re.fullmatch(r"1\d{10}", item.user_name):
            by_phone[item.user_name].append(item)
        by_name[item.name].append(item)
        if item.id_number:
            by_id[item.id_number].append(item)

    for source in employees:
        phone_matches = {item.user_id: item for item in by_phone.get(source.phone, []) if item.del_flag == "0"}
        name_matches = {item.user_id: item for item in by_name.get(source.name, []) if item.del_flag == "0"}
        if phone_matches and name_matches and set(phone_matches) != set(name_matches):
            raise MigrationError(f"phone/name resolve to different users for {source.name}")
        if not phone_matches:
            conflicting_named = [
                item for item in name_matches.values()
                if item.phone and phone(item.phone) != source.phone
            ]
            if conflicting_named:
                raise MigrationError(f"name fallback has a different non-empty phone for {source.name}")
        candidates = phone_matches or name_matches
        if len(candidates) > 1:
            raise MigrationError(f"ambiguous ERP match for {source.name}: {sorted(candidates)}")
        source.existing = next(iter(candidates.values()), None)
        if source.existing:
            if source.existing.profile_id is None or not source.existing.employee_no:
                raise MigrationError(f"existing employee lacks formal profile/number: {source.name}")
        else:
            deleted_matches = [item for item in by_phone.get(source.phone, []) + by_name.get(source.name, [])
                               if item.del_flag != "0"]
            if deleted_matches:
                raise MigrationError(f"deleted account conflict for {source.name}")

        source_id = text(source.profile.get("id_number")).upper()
        if source_id:
            owners = {item.user_id for item in by_id.get(source_id, []) if item.del_flag == "0"}
            expected = source.existing.user_id if source.existing else None
            if owners and owners != ({expected} if expected is not None else set()):
                raise MigrationError(f"ID number is owned by another ERP user for {source.name}")
        supervisor = text(source.profile.get("department_supervisor"))
        if supervisor:
            supervisor_matches = [item for item in by_name.get(supervisor, []) if item.del_flag == "0"]
            if len(supervisor_matches) != 1:
                raise MigrationError(
                    f"department supervisor is unresolved or ambiguous for {source.name}: {supervisor}"
                )

    existing_count = sum(item.existing is not None for item in employees)
    new_count = len(employees) - existing_count
    if existing_count != EXPECTED_EXISTING_COUNT or new_count != EXPECTED_NEW_COUNT:
        raise MigrationError(
            f"database drift: expected existing/new={EXPECTED_EXISTING_COUNT}/{EXPECTED_NEW_COUNT}, "
            f"got {existing_count}/{new_count}"
        )
    new_names = {item.name for item in employees if not item.existing}
    if new_names != EXPECTED_NEW_NAMES:
        raise MigrationError(f"unexpected new employee set: {sorted(new_names)}")

    paths = dept_paths(args)
    for source in employees:
        main_suffix, scope_suffixes = SOURCE_STORE_ORG_SUFFIXES[source.source_store]
        source.dept_id = resolve_dept(paths, main_suffix)
        source.scope_dept_ids = [resolve_dept(paths, suffix) for suffix in scope_suffixes]
        matching_paths = [path for path, dept_id in paths.items() if dept_id == source.dept_id]
        if len(matching_paths) != 1:
            raise MigrationError(f"target department path not unique for {source.name}")
        source.dept_path = matching_paths[0]

    post_rows = mysql_query(args, """
        select p.post_id,p.post_code,p.post_name,
               coalesce(cast(r.role_id as char),''),coalesce(r.role_name,'')
        from sys_post p
        left join sys_role r on r.role_name=p.post_name and r.status='0' and r.del_flag='0'
        where p.status='0'
    """)
    post_config = {row[2]: (int(row[0]), row[1], int(row[3]) if row[3] else None, row[4])
                   for row in post_rows}
    standard_role_ids: set[int] = set()
    for source in employees:
        config = post_config.get(source.post)
        if not config or config[2] is None or config[3] != POST_ROLE[source.post]:
            raise MigrationError(f"post/role mapping is incomplete for {source.post}")
        source.post_id, source.post_code, source.role_id, _ = config
        standard_role_ids.add(source.role_id)

    sequence_rows = mysql_query(args, """
        select current_value from hr_employee_no_sequence where sequence_key='GLOBAL'
    """)
    if len(sequence_rows) != 1:
        raise MigrationError("GLOBAL employee number sequence is missing")
    return int(sequence_rows[0][0]), sorted(standard_role_ids)


def assignment(column: str, value: object) -> str:
    return f"{column}={sql_quote(value)}"


def profile_values(source: SourceEmployee, dept_id: int) -> dict[str, object]:
    values = dict(source.profile)
    values.pop("_source_store", None)
    level3 = "西安区域运营" if "西安区域运营" in source.dept_path else "北京区域运营"
    values.update({
        "employee_name": source.name,
        "phone_number": source.phone,
        "dept_id": dept_id,
        "dept_level1_name": "金英灵韵运营部",
        "dept_level2_name": "北京区域",
        "dept_level3_name": level3,
        "position_names": source.post,
        "legal_entity_code": LEGAL_ENTITY_CODE,
        "profile_source": "北京区域0716-2.xlsx",
        "last_profile_update_by": OPERATOR,
        "update_by": OPERATOR,
    })
    if source.source_store != "西安区域运营":
        values["store_name"] = source.source_store
    return values


def profile_assignments(source: SourceEmployee, dept_id: int, position_no_sql: str) -> list[str]:
    values = profile_values(source, dept_id)
    result = [assignment(key, value) for key, value in values.items()]
    result.extend([
        f"position_no={position_no_sql}",
        "legal_entity_id=v_legal_entity_id",
        "last_profile_update_time=now()",
        "update_time=now()",
    ])
    return result


def record_verification_lines(source: SourceEmployee, user_sql: str,
                              employee_no_sql: str, position_no_sql: str,
                              standard_roles: str, is_new: bool) -> list[str]:
    assert source.dept_id is not None and source.post_id is not None and source.role_id is not None
    conditions = [
        f"u.user_id={user_sql}",
        f"CAST(u.nick_name AS BINARY)=CAST({sql_quote(source.name)} AS BINARY)",
        f"CAST(u.phonenumber AS BINARY)=CAST({sql_quote(source.phone)} AS BINARY)",
        f"u.dept_id={source.dept_id}",
        "u.status='0'",
        "u.del_flag='0'",
        f"CAST(p.employee_no AS BINARY)=CAST({employee_no_sql} AS BINARY)",
        f"CAST(p.position_no AS BINARY)=CAST({position_no_sql} AS BINARY)",
        "p.legal_entity_id=v_legal_entity_id",
    ]
    sex = source.profile.get("sex")
    if sex is not None:
        conditions.append(f"u.sex <=> {sql_quote(sex)}")
    if is_new:
        conditions.extend([
            "u.credential_state='TEMPORARY'",
            "u.temporary_password_expires_at>NOW()",
        ])
    for column, value in profile_values(source, source.dept_id).items():
        conditions.append(f"p.{column} <=> {sql_quote(value)}")

    result = [
        "  IF (SELECT COUNT(*) FROM sys_user u JOIN sys_user_profile p ON p.user_id=u.user_id "
        f"WHERE {' AND '.join(conditions)}) <> 1 THEN",
        f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='verification failed for {source.name}';",
        "  END IF;",
        f"  IF (SELECT COUNT(*) FROM sys_user_post WHERE user_id={user_sql}) <> 1 OR "
        f"(SELECT COUNT(*) FROM sys_user_post WHERE user_id={user_sql} AND post_id={source.post_id}) <> 1 THEN",
        f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='post verification failed for {source.name}';",
        "  END IF;",
        f"  IF (SELECT COUNT(*) FROM sys_user_role WHERE user_id={user_sql} AND role_id IN ({standard_roles})) <> 1 OR "
        f"(SELECT COUNT(*) FROM sys_user_role WHERE user_id={user_sql} AND role_id={source.role_id}) <> 1 THEN",
        f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='role verification failed for {source.name}';",
        "  END IF;",
        f"  IF (SELECT COUNT(*) FROM sys_user_shop WHERE user_id={user_sql} AND is_default='Y') <> 1 OR "
        f"(SELECT COUNT(*) FROM sys_user_shop WHERE user_id={user_sql} AND dept_id={source.dept_id} AND is_default='Y') <> 1 THEN",
        f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='default scope verification failed for {source.name}';",
        "  END IF;",
    ]
    for dept_id in source.scope_dept_ids:
        is_default = "Y" if dept_id == source.dept_id else "N"
        result.extend([
            f"  IF (SELECT COUNT(*) FROM sys_user_shop WHERE user_id={user_sql} AND dept_id={dept_id} "
            f"AND is_default={sql_quote(is_default)}) <> 1 THEN",
            f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='scope verification failed for {source.name}';",
            "  END IF;",
        ])
    if source.existing and source.existing.dept_id != source.dept_id:
        result.extend([
            f"  IF (SELECT COUNT(*) FROM sys_user_shop WHERE user_id={user_sql} "
            f"AND dept_id={source.existing.dept_id}) <> 0 THEN",
            f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='stale scope remains for {source.name}';",
            "  END IF;",
        ])
    return result


def build_apply_sql(employees: list[SourceEmployee], expected_sequence: int,
                    standard_role_ids: Iterable[int], workbook_sha: str) -> str:
    # DDL happens before the advisory lock can be taken inside the procedure;
    # a per-run name prevents concurrent invocations from replacing each
    # other's procedure (and therefore mismatching temporary credentials).
    procedure_name = f"{PROCEDURE_NAME}_{secrets.token_hex(8)}"
    for source in employees:
        if source.existing and not source.session_event_id:
            source.session_event_id = secrets.token_hex(16)
    standard_roles = ",".join(str(item) for item in standard_role_ids)
    names = ",".join(f"CAST({sql_quote(item.name)} AS BINARY)" for item in employees)
    phones = ",".join(f"CAST({sql_quote(item.phone)} AS BINARY)" for item in employees)
    ids = ",".join(
        f"CAST({sql_quote(text(item.profile.get('id_number')).upper())} AS BINARY)"
        for item in employees if item.profile.get("id_number")
    )
    category_values = ",".join(
        f"CAST({sql_quote(value)} AS BINARY)" for value in REQUIRED_CATEGORY_DICTIONARY
    )
    category_labels = ",".join(
        f"CAST({sql_quote(label)} AS BINARY)"
        for _, label in REQUIRED_CATEGORY_DICTIONARY.values()
    )
    lines = [
        "SET NAMES utf8mb4 COLLATE utf8mb4_general_ci;",
        f"DROP PROCEDURE IF EXISTS {procedure_name};",
        "DELIMITER $$",
        f"CREATE PROCEDURE {procedure_name}()",
        "main: BEGIN",
        "  DECLARE v_lock int DEFAULT 0;",
        "  DECLARE v_sequence int DEFAULT 0;",
        "  DECLARE v_identity_rows int DEFAULT 0;",
        "  DECLARE v_user_id bigint DEFAULT NULL;",
        "  DECLARE v_legal_entity_id bigint DEFAULT NULL;",
        "  DECLARE v_employee_no varchar(64) DEFAULT NULL;",
        "  DECLARE EXIT HANDLER FOR SQLEXCEPTION",
        "  BEGIN",
        "    ROLLBACK;",
        f"    DO RELEASE_LOCK({sql_quote(LOCK_NAME)});",
        "    RESIGNAL;",
        "  END;",
        f"  SELECT GET_LOCK({sql_quote(LOCK_NAME)}, 30) INTO v_lock;",
        "  IF v_lock <> 1 THEN",
        "    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='employee reconciliation lock unavailable';",
        "  END IF;",
        "  START TRANSACTION;",
        # Lock every row that can own a source identity before rechecking it.
        "  SELECT COUNT(*) INTO v_identity_rows FROM sys_user WHERE "
        f"CAST(nick_name AS BINARY) IN ({names}) OR CAST(phonenumber AS BINARY) IN ({phones}) "
        f"OR CAST(user_name AS BINARY) IN ({phones}) FOR UPDATE;",
    ]
    if ids:
        lines.append(
            "  SELECT COUNT(*) INTO v_identity_rows FROM sys_user_profile WHERE "
            f"CAST(id_number AS BINARY) IN ({ids}) FOR UPDATE;"
        )
    lines.extend([
        "  SELECT current_value INTO v_sequence",
        "  FROM hr_employee_no_sequence WHERE sequence_key='GLOBAL' FOR UPDATE;",
        f"  IF v_sequence <> {expected_sequence} THEN",
        "    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='employee number sequence changed after preflight';",
        "  END IF;",
        "  IF EXISTS (SELECT 1 FROM sys_legal_entity WHERE "
        f"CAST(legal_entity_code AS BINARY)=CAST({sql_quote(LEGAL_ENTITY_CODE)} AS BINARY) AND "
        f"CAST(legal_entity_name AS BINARY)<>CAST({sql_quote(LEGAL_ENTITY_NAME)} AS BINARY)) OR "
        "EXISTS (SELECT 1 FROM sys_legal_entity WHERE "
        f"CAST(legal_entity_name AS BINARY)=CAST({sql_quote(LEGAL_ENTITY_NAME)} AS BINARY) AND "
        f"CAST(legal_entity_code AS BINARY)<>CAST({sql_quote(LEGAL_ENTITY_CODE)} AS BINARY)) THEN",
        "    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='legal entity code/name conflict';",
        "  END IF;",
        "  INSERT INTO sys_legal_entity(legal_entity_code,legal_entity_name,legal_representative,status,"
        "create_by,create_time,update_by,update_time,remark) SELECT " + ",".join([
            sql_quote(LEGAL_ENTITY_CODE), sql_quote(LEGAL_ENTITY_NAME), sql_quote(LEGAL_REPRESENTATIVE),
            sql_quote("0"), sql_quote(OPERATOR), "NOW()", sql_quote(OPERATOR), "NOW()",
            sql_quote(f"{SALARY_VERSION}签约主体;sha256={workbook_sha[:16]}")
        ]) + " WHERE NOT EXISTS (SELECT 1 FROM sys_legal_entity WHERE "
        f"CAST(legal_entity_code AS BINARY)=CAST({sql_quote(LEGAL_ENTITY_CODE)} AS BINARY));",
        "  SELECT legal_entity_id INTO v_legal_entity_id FROM sys_legal_entity WHERE "
        f"CAST(legal_entity_code AS BINARY)=CAST({sql_quote(LEGAL_ENTITY_CODE)} AS BINARY) FOR UPDATE;",
        "  UPDATE sys_legal_entity SET " + ",".join([
            assignment("legal_entity_name", LEGAL_ENTITY_NAME),
            assignment("legal_representative", LEGAL_REPRESENTATIVE), assignment("status", "0"),
            assignment("update_by", OPERATOR), "update_time=NOW()",
        ]) + " WHERE legal_entity_id=v_legal_entity_id;",
        "  SELECT COUNT(*) INTO v_identity_rows FROM sys_dict_data WHERE "
        "CAST(dict_type AS BINARY)=CAST('hr_employee_category' AS BINARY) AND ("
        f"CAST(dict_value AS BINARY) IN ({category_values}) OR "
        f"CAST(dict_label AS BINARY) IN ({category_labels})) FOR UPDATE;",
    ])

    for category_value, (dict_sort, category_label) in REQUIRED_CATEGORY_DICTIONARY.items():
        lines.extend([
            "  IF (SELECT COUNT(*) FROM sys_dict_data WHERE "
            "CAST(dict_type AS BINARY)=CAST('hr_employee_category' AS BINARY) AND ("
            f"CAST(dict_value AS BINARY)=CAST({sql_quote(category_value)} AS BINARY) OR "
            f"CAST(dict_label AS BINARY)=CAST({sql_quote(category_label)} AS BINARY))) > 1 THEN",
            f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='category dictionary duplicate for {category_value}';",
            "  END IF;",
            "  IF EXISTS (SELECT 1 FROM sys_dict_data WHERE "
            "CAST(dict_type AS BINARY)=CAST('hr_employee_category' AS BINARY) AND (("
            f"CAST(dict_value AS BINARY)=CAST({sql_quote(category_value)} AS BINARY) AND "
            f"CAST(dict_label AS BINARY)<>CAST({sql_quote(category_label)} AS BINARY)) OR ("
            f"CAST(dict_label AS BINARY)=CAST({sql_quote(category_label)} AS BINARY) AND "
            f"CAST(dict_value AS BINARY)<>CAST({sql_quote(category_value)} AS BINARY)))) THEN",
            f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='category dictionary conflict for {category_value}';",
            "  END IF;",
            "  INSERT INTO sys_dict_data(dict_sort,dict_label,dict_value,dict_type,css_class,list_class,"
            "is_default,status,create_by,create_time,update_by,update_time,remark) SELECT " + ",".join([
                str(dict_sort), sql_quote(category_label), sql_quote(category_value),
                sql_quote("hr_employee_category"), sql_quote(""), sql_quote("default"),
                sql_quote("N"), sql_quote("0"), sql_quote(OPERATOR), "NOW()",
                sql_quote(OPERATOR), "NOW()",
                sql_quote(f"{SALARY_VERSION}员工类别主数据;sha256={workbook_sha[:16]}")
            ]) + " WHERE NOT EXISTS (SELECT 1 FROM sys_dict_data WHERE "
            "CAST(dict_type AS BINARY)=CAST('hr_employee_category' AS BINARY) AND "
            f"CAST(dict_value AS BINARY)=CAST({sql_quote(category_value)} AS BINARY));",
            "  UPDATE sys_dict_data SET " + ",".join([
                assignment("dict_sort", dict_sort), assignment("dict_label", category_label),
                assignment("status", "0"), assignment("update_by", OPERATOR), "update_time=NOW()",
            ]) + " WHERE CAST(dict_type AS BINARY)=CAST('hr_employee_category' AS BINARY) AND "
            f"CAST(dict_value AS BINARY)=CAST({sql_quote(category_value)} AS BINARY);",
            "  IF (SELECT COUNT(*) FROM sys_dict_data WHERE "
            "CAST(dict_type AS BINARY)=CAST('hr_employee_category' AS BINARY) AND "
            f"CAST(dict_value AS BINARY)=CAST({sql_quote(category_value)} AS BINARY) AND "
            f"CAST(dict_label AS BINARY)=CAST({sql_quote(category_label)} AS BINARY) AND status='0') <> 1 THEN",
            f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='category dictionary verify failed for {category_value}';",
            "  END IF;",
        ])

    # Recheck all phone/name/ID ownership inside the same locking transaction.
    for source in employees:
        identity_predicate = (
            f"CAST(nick_name AS BINARY)=CAST({sql_quote(source.name)} AS BINARY) OR "
            f"CAST(phonenumber AS BINARY)=CAST({sql_quote(source.phone)} AS BINARY) OR "
            f"CAST(user_name AS BINARY)=CAST({sql_quote(source.phone)} AS BINARY)"
        )
        source_id = text(source.profile.get("id_number")).upper()
        if source.existing:
            user_id = source.existing.user_id
            lines.extend([
                f"  IF (SELECT COUNT(*) FROM sys_user WHERE ({identity_predicate})) <> 1 OR "
                f"(SELECT COUNT(*) FROM sys_user WHERE user_id={user_id} AND ({identity_predicate}) AND del_flag='0') <> 1 THEN",
                f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='identity changed for {source.name}';",
                "  END IF;",
            ])
            if source_id:
                lines.extend([
                    "  IF EXISTS (SELECT 1 FROM sys_user_profile WHERE "
                    f"user_id<>{user_id} AND CAST(id_number AS BINARY)=CAST({sql_quote(source_id)} AS BINARY)) THEN",
                    f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='ID conflict for {source.name}';",
                    "  END IF;",
                ])
        else:
            lines.extend([
                f"  IF EXISTS (SELECT 1 FROM sys_user WHERE {identity_predicate}) THEN",
                f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='new identity conflict for {source.name}';",
                "  END IF;",
            ])
            if source_id:
                lines.extend([
                    "  IF EXISTS (SELECT 1 FROM sys_user_profile WHERE "
                    f"CAST(id_number AS BINARY)=CAST({sql_quote(source_id)} AS BINARY)) THEN",
                    f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='new ID conflict for {source.name}';",
                    "  END IF;",
                ])

    for source in employees:
        assert source.dept_id is not None and source.post_id is not None and source.role_id is not None
        sex = source.profile.get("sex")
        if source.existing:
            user_sql = str(source.existing.user_id)
            employee_no_sql = sql_quote(source.existing.employee_no)
            position_no_sql = sql_quote(f"{source.post_code.upper()}-{source.existing.employee_no}")
            lines.extend([
                "  UPDATE sys_user SET " + ",".join(filter(None, [
                    assignment("dept_id", source.dept_id), assignment("nick_name", source.name),
                    assignment("phonenumber", source.phone), assignment("status", "0"),
                    assignment("del_flag", "0"), assignment("sex", sex) if sex is not None else "",
                    assignment("update_by", OPERATOR), "update_time=NOW()",
                    assignment("remark", f"{SALARY_VERSION}权威同步;源行{source.row_number};sha256={workbook_sha[:16]}")
                ])) + f" WHERE user_id={user_sql};",
                "  IF ROW_COUNT() <> 1 THEN",
                f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='user update failed for {source.name}';",
                "  END IF;",
                "  UPDATE sys_user_profile SET " + ",".join(
                    profile_assignments(source, source.dept_id, position_no_sql)
                ) + f" WHERE user_id={user_sql};",
                "  IF ROW_COUNT() <> 1 THEN",
                f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='profile update failed for {source.name}';",
                "  END IF;",
                f"  DELETE FROM sys_user_post WHERE user_id={user_sql};",
                f"  INSERT INTO sys_user_post(user_id,post_id) VALUES({user_sql},{source.post_id});",
                f"  DELETE FROM sys_user_role WHERE user_id={user_sql} AND role_id IN ({standard_roles});",
                f"  INSERT INTO sys_user_role(user_id,role_id) VALUES({user_sql},{source.role_id}) "
                "ON DUPLICATE KEY UPDATE role_id=VALUES(role_id);",
                # An old default row is the departed assignment, not an extra
                # manager grant. Remove it; pre-existing non-default scopes are
                # preserved and authoritative scopes are upserted below.
                f"  DELETE FROM sys_user_shop WHERE user_id={user_sql} AND is_default='Y' "
                f"AND dept_id<>{source.dept_id};",
            ])
            for dept_id in source.scope_dept_ids:
                is_default = "Y" if dept_id == source.dept_id else "N"
                lines.append(
                    "  INSERT INTO sys_user_shop(user_id,dept_id,is_default,create_by,create_time) "
                    f"VALUES({user_sql},{dept_id},{sql_quote(is_default)},{sql_quote(OPERATOR)},NOW()) "
                    "ON DUPLICATE KEY UPDATE is_default=VALUES(is_default);"
                )
            lines.extend([
                "  INSERT INTO sys_security_session_outbox(event_id,event_type,user_id,reason_code,status,"
                "attempts,available_at,created_at) VALUES(" + ",".join([
                    sql_quote(source.session_event_id), sql_quote("INVALIDATE_USER_SESSIONS"),
                    user_sql, sql_quote("USER_ROLES_CHANGED"), sql_quote("PENDING"), "0", "NOW()", "NOW()",
                ]) + ");",
                "  IF ROW_COUNT() <> 1 THEN",
                f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='session event insert failed for {source.name}';",
                "  END IF;",
                "  IF (SELECT COUNT(*) FROM sys_security_session_outbox WHERE "
                f"CAST(event_id AS BINARY)=CAST({sql_quote(source.session_event_id)} AS BINARY) AND "
                f"user_id={user_sql} AND event_type='INVALIDATE_USER_SESSIONS' AND "
                "reason_code='USER_ROLES_CHANGED' AND status='PENDING') <> 1 THEN",
                f"    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='session event verify failed for {source.name}';",
                "  END IF;",
            ])
            lines.extend(record_verification_lines(
                source, user_sql, employee_no_sql, position_no_sql, standard_roles, False
            ))
            continue

        if not source.temporary_password_hash:
            raise MigrationError(f"temporary credential not generated for {source.name}")
        lines.extend([
            "  SET v_sequence=v_sequence+1;",
            "  SET v_employee_no=CONCAT('E',LPAD(v_sequence,5,'0'));",
            "  IF EXISTS (SELECT 1 FROM sys_user WHERE CAST(user_name AS BINARY)=CAST(v_employee_no AS BINARY)) OR "
            "EXISTS (SELECT 1 FROM sys_user_profile WHERE CAST(employee_no AS BINARY)=CAST(v_employee_no AS BINARY)) THEN",
            "    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='generated employee number conflict';",
            "  END IF;",
            "  INSERT INTO sys_user(dept_id,user_name,nick_name,user_type,email,phonenumber,sex,avatar,password,"
            "status,del_flag,pwd_update_date,credential_state,temporary_password_expires_at,create_by,create_time,"
            "update_by,update_time,remark) VALUES(" + ",".join([
                str(source.dept_id), "v_employee_no", sql_quote(source.name), sql_quote("00"), sql_quote(""),
                sql_quote(source.phone), sql_quote(sex or "2"), sql_quote(""),
                sql_quote(source.temporary_password_hash), sql_quote("0"), sql_quote("0"), "NULL",
                sql_quote("TEMPORARY"), f"DATE_ADD(NOW(),INTERVAL {TEMPORARY_PASSWORD_HOURS} HOUR)",
                sql_quote(OPERATOR), "NOW()", sql_quote(OPERATOR), "NOW()",
                sql_quote(f"{SALARY_VERSION}新员工;临时凭据;源行{source.row_number};sha256={workbook_sha[:16]}")
            ]) + ");",
            "  SET v_user_id=LAST_INSERT_ID();",
        ])
        insert_values = profile_values(source, source.dept_id)
        insert_values["create_by"] = OPERATOR
        columns = ["user_id", "employee_no", "position_no", "legal_entity_id", *insert_values.keys(),
                   "last_profile_update_time", "create_time", "update_time"]
        position_no_sql = f"CONCAT({sql_quote(source.post_code.upper() + '-')},v_employee_no)"
        values = ["v_user_id", "v_employee_no", position_no_sql, "v_legal_entity_id"]
        values.extend(sql_quote(value) for value in insert_values.values())
        values.extend(["NOW()", "NOW()", "NOW()"])
        lines.extend([
            f"  INSERT INTO sys_user_profile({','.join(columns)}) VALUES({','.join(values)});",
            f"  INSERT INTO sys_user_post(user_id,post_id) VALUES(v_user_id,{source.post_id});",
            f"  INSERT INTO sys_user_role(user_id,role_id) VALUES(v_user_id,{source.role_id});",
        ])
        for dept_id in source.scope_dept_ids:
            is_default = "Y" if dept_id == source.dept_id else "N"
            lines.append(
                "  INSERT INTO sys_user_shop(user_id,dept_id,is_default,create_by,create_time) "
                f"VALUES(v_user_id,{dept_id},{sql_quote(is_default)},{sql_quote(OPERATOR)},NOW());"
            )
        lines.extend(record_verification_lines(
            source, "v_user_id", "v_employee_no", position_no_sql, standard_roles, True
        ))

    lines.extend([
        "  UPDATE hr_employee_no_sequence SET current_value=v_sequence,update_time=NOW() "
        "WHERE sequence_key='GLOBAL';",
        "  IF ROW_COUNT() <> 1 THEN",
        "    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='employee sequence update failed';",
        "  END IF;",
        "  IF (SELECT COUNT(*) FROM sys_user WHERE del_flag='0' AND "
        f"CAST(phonenumber AS BINARY) IN ({phones})) <> {len(employees)} THEN",
        "    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='final source population verification failed';",
        "  END IF;",
        "  COMMIT;",
        f"  DO RELEASE_LOCK({sql_quote(LOCK_NAME)});",
        "END$$",
        "DELIMITER ;",
        f"CALL {procedure_name}();",
        f"DROP PROCEDURE IF EXISTS {procedure_name};",
    ])
    return "\n".join(lines) + "\n"


def print_plan(employees: list[SourceEmployee], sequence: int, sign_only: list[str]) -> None:
    existing = [item for item in employees if item.existing]
    new = [item for item in employees if not item.existing]
    print(f"workbook_rows={len(employees)} existing={len(existing)} new={len(new)} excluded_sign_only={sign_only}")
    print("new_employees=" + ",".join(item.name for item in new))
    changes = []
    for item in existing:
        assert item.existing is not None
        if item.existing.old_post != item.post or item.existing.old_status != item.employee_status:
            changes.append(
                f"{item.name}[post:{item.existing.old_post or '-'}->{item.post};"
                f"status:{item.existing.old_status or '-'}->{item.employee_status}]"
            )
    print("post_or_status_changes=" + (",".join(changes) if changes else "none"))
    department_changes = [
        f"{item.name}[dept:{item.existing.dept_id}->{item.dept_id}]"
        for item in existing
        if item.existing and item.existing.dept_id != item.dept_id
    ]
    print("department_changes=" + (",".join(department_changes) if department_changes else "none"))
    print(f"employee_number_sequence={sequence}->{sequence + len(new)}")
    print("missing_workbook_employees_are_preserved=true")


def verify(args: argparse.Namespace, employees: list[SourceEmployee],
           standard_role_ids: Iterable[int], expected_sequence: int) -> None:
    standard_roles = ",".join(str(item) for item in standard_role_ids)
    phones = ",".join(f"CAST({sql_quote(item.phone)} AS BINARY)" for item in employees)
    rows = mysql_query(args, f"""
        SELECT COUNT(*) FROM sys_user
        WHERE del_flag='0' AND CAST(phonenumber AS BINARY) IN ({phones})
    """)
    if rows != [[str(len(employees))]]:
        raise MigrationError(f"post-apply employee population mismatch: {rows}")

    for source in employees:
        assert source.dept_id is not None and source.post_id is not None and source.role_id is not None
        if source.existing:
            user_sql = str(source.existing.user_id)
            employee_no_sql = sql_quote(source.existing.employee_no)
            position_no_sql = sql_quote(f"{source.post_code.upper()}-{source.existing.employee_no}")
            credential_conditions: list[str] = []
        else:
            if not source.generated_employee_no:
                raise MigrationError(f"generated employee number missing for {source.name}")
            user_sql = (
                "(SELECT user_id FROM sys_user WHERE CAST(user_name AS BINARY)="
                f"CAST({sql_quote(source.generated_employee_no)} AS BINARY))"
            )
            employee_no_sql = sql_quote(source.generated_employee_no)
            position_no_sql = sql_quote(f"{source.post_code.upper()}-{source.generated_employee_no}")
            credential_conditions = [
                "u.credential_state='TEMPORARY'",
                "u.temporary_password_expires_at>NOW()",
            ]

        conditions = [
            f"u.user_id={user_sql}",
            f"CAST(u.nick_name AS BINARY)=CAST({sql_quote(source.name)} AS BINARY)",
            f"CAST(u.phonenumber AS BINARY)=CAST({sql_quote(source.phone)} AS BINARY)",
            f"u.dept_id={source.dept_id}", "u.status='0'", "u.del_flag='0'",
            f"p.employee_no <=> {employee_no_sql}",
            f"p.position_no <=> {position_no_sql}",
            f"e.legal_entity_code <=> {sql_quote(LEGAL_ENTITY_CODE)}",
            "p.legal_entity_id=e.legal_entity_id",
            *credential_conditions,
        ]
        sex = source.profile.get("sex")
        if sex is not None:
            conditions.append(f"u.sex <=> {sql_quote(sex)}")
        for column, value in profile_values(source, source.dept_id).items():
            conditions.append(f"p.{column} <=> {sql_quote(value)}")

        checks = [
            "(SELECT COUNT(*) FROM sys_user u JOIN sys_user_profile p ON p.user_id=u.user_id "
            "JOIN sys_legal_entity e ON e.legal_entity_id=p.legal_entity_id "
            f"WHERE {' AND '.join(conditions)})=1",
            f"(SELECT COUNT(*) FROM sys_user_post WHERE user_id={user_sql})=1",
            f"(SELECT COUNT(*) FROM sys_user_post WHERE user_id={user_sql} AND post_id={source.post_id})=1",
            f"(SELECT COUNT(*) FROM sys_user_role WHERE user_id={user_sql} AND role_id IN ({standard_roles}))=1",
            f"(SELECT COUNT(*) FROM sys_user_role WHERE user_id={user_sql} AND role_id={source.role_id})=1",
            f"(SELECT COUNT(*) FROM sys_user_shop WHERE user_id={user_sql} AND is_default='Y')=1",
            f"(SELECT COUNT(*) FROM sys_user_shop WHERE user_id={user_sql} AND dept_id={source.dept_id} "
            "AND is_default='Y')=1",
        ]
        if source.existing and source.existing.dept_id != source.dept_id:
            checks.append(
                f"(SELECT COUNT(*) FROM sys_user_shop WHERE user_id={user_sql} "
                f"AND dept_id={source.existing.dept_id})=0"
            )
        for dept_id in source.scope_dept_ids:
            is_default = "Y" if dept_id == source.dept_id else "N"
            checks.append(
                f"(SELECT COUNT(*) FROM sys_user_shop WHERE user_id={user_sql} AND dept_id={dept_id} "
                f"AND is_default={sql_quote(is_default)})=1"
            )
        result = mysql_query(args, "SELECT " + ",".join(checks))
        if result != [["1"] * len(checks)]:
            raise MigrationError(f"post-commit exhaustive verification failed for {source.name}: {result}")
        if source.existing:
            session_event = mysql_query(args, f"""
                SELECT COUNT(*) FROM sys_security_session_outbox
                WHERE CAST(event_id AS BINARY)=CAST({sql_quote(source.session_event_id)} AS BINARY)
                  AND user_id={source.existing.user_id}
                  AND event_type='INVALIDATE_USER_SESSIONS'
                  AND reason_code='USER_ROLES_CHANGED'
                  AND status IN ('PENDING','PROCESSING','DONE')
            """)
            if session_event != [["1"]]:
                raise MigrationError(
                    f"session invalidation event verification failed for {source.name}: {session_event}"
                )

    final_sequence = mysql_query(args, """
        SELECT current_value FROM hr_employee_no_sequence WHERE sequence_key='GLOBAL'
    """)
    expected_final_sequence = expected_sequence + sum(item.existing is None for item in employees)
    if final_sequence != [[str(expected_final_sequence)]]:
        raise MigrationError(f"employee sequence verification failed: {final_sequence}")

    for category_value, (_, category_label) in REQUIRED_CATEGORY_DICTIONARY.items():
        category_rows = mysql_query(args, f"""
            SELECT COUNT(*) FROM sys_dict_data
            WHERE CAST(dict_type AS BINARY)=CAST('hr_employee_category' AS BINARY)
              AND CAST(dict_value AS BINARY)=CAST({sql_quote(category_value)} AS BINARY)
              AND CAST(dict_label AS BINARY)=CAST({sql_quote(category_label)} AS BINARY)
              AND status='0'
        """)
        if category_rows != [["1"]]:
            raise MigrationError(
                f"employee category dictionary verification failed for {category_value}: {category_rows}"
            )

    protected_names = sorted(PRESERVED_CURRENT_NAMES | {"段继康"})
    protected_sql = ",".join(sql_quote(name) for name in protected_names)
    touched = mysql_query(args, f"""
        SELECT COUNT(*)
        FROM sys_user u LEFT JOIN sys_user_profile p ON p.user_id=u.user_id
        WHERE u.nick_name IN ({protected_sql})
          AND (u.update_by={sql_quote(OPERATOR)}
               OR p.last_profile_update_by={sql_quote(OPERATOR)}
               OR p.salary_version={sql_quote(SALARY_VERSION)})
    """)
    if touched != [["0"]]:
        raise MigrationError(f"out-of-scope employee invariant failed: {touched}")
    print(
        f"verified_employees={len(employees)} all_written_fields_relations=ok "
        f"preserved_out_of_scope={len(protected_names)} "
        f"session_invalidations={sum(item.existing is not None for item in employees)}"
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", required=True)
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", default="3306")
    parser.add_argument("--user", default="root")
    parser.add_argument("--password", default="")
    parser.add_argument("--database", default="BossERP_NEW")
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--credentials-output")
    args = parser.parse_args()

    excel = Path(args.excel).expanduser().resolve()
    if not excel.is_file():
        raise MigrationError(f"Excel file not found: {excel}")
    employees, workbook_sha, sign_only = load_workbook(excel)
    sequence, standard_role_ids = attach_database_plan(args, employees)
    print(f"workbook_sha256={workbook_sha}")
    print_plan(employees, sequence, sign_only)
    if not args.apply:
        return 0

    if not args.credentials_output:
        raise MigrationError("--credentials-output is required with --apply")
    credential_path = Path(args.credentials_output).expanduser().resolve()
    if credential_path.exists():
        raise MigrationError(f"credential output already exists: {credential_path}")
    issue_temporary_credentials(employees, sequence)
    write_credentials(credential_path, employees)
    apply_sql = build_apply_sql(employees, sequence, standard_role_ids, workbook_sha)
    output = mysql_execute(args, apply_sql)
    if output.strip():
        print(output.strip())
    verify(args, employees, standard_role_ids, sequence)
    print(f"temporary_credentials={credential_path}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except MigrationError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
