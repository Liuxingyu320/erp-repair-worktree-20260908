#!/usr/bin/env python3
"""Build isolated v4 onboarding-signing templates for Beijing release QA.

This builder is deliberately local-only: it never changes the v3 source files,
does not upload anything, and fails if a target artifact already exists.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Final
from zipfile import ZipFile

from docx import Document


REPO_ROOT: Final[Path] = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE_DIR: Final[Path] = (
    REPO_ROOT / "uploadPath/sign-template/onboard-20260714-v3-draft"
)
DEFAULT_OUTPUT_DIR: Final[Path] = (
    REPO_ROOT / "output/sign-release/20260718/templates"
)
DEFAULT_QA_DIR: Final[Path] = REPO_ROOT / "output/sign-release/20260718/qa/pressure"


@dataclass(frozen=True)
class TemplateSpec:
    source_name: str
    output_name: str
    template_type: str
    page_break_marker: str | None
    seal_position: dict[str, int] | None


SPECS: Final[tuple[TemplateSpec, ...]] = (
    TemplateSpec(
        source_name="09_ONBOARD_LABOR_CONTRACT.docx",
        output_name="09_ONBOARD_LABOR_CONTRACT_v4-last-page.docx",
        template_type="ONBOARD_LABOR_CONTRACT",
        page_break_marker="第十六条 附则",
        seal_position={"x": 177, "y": 497, "width": 78, "height": 78},
    ),
    TemplateSpec(
        source_name="13_ONBOARD_SERVICE_CONTRACT.docx",
        output_name="13_ONBOARD_SERVICE_CONTRACT_v4-last-page.docx",
        template_type="ONBOARD_SERVICE_CONTRACT",
        page_break_marker="第十七条 附则",
        seal_position={"x": 177, "y": 537, "width": 78, "height": 78},
    ),
    TemplateSpec(
        source_name="14_ONBOARD_SERVICE_RECEIPT.docx",
        output_name="14_ONBOARD_SERVICE_RECEIPT_v4-last-page.docx",
        template_type="ONBOARD_SERVICE_RECEIPT",
        page_break_marker=None,
        seal_position=None,
    ),
)


def fixed_width(seed: str, width: int) -> str:
    """Repeat and truncate a readable seed to an exact database character width."""
    if not seed or width < 1:
        raise ValueError("seed and width must be non-empty/positive")
    return (seed * ((width // len(seed)) + 1))[:width]


# These strings exercise the persisted field widths in
# sql/erp_oa_sign_package_20260702.sql and the 160-character legal-entity name.
PRESSURE_VALUES: Final[dict[str, str]] = {
    "employeeName": fixed_width("欧阳压力测试姓名", 64),
    "employeePhone": "1" * 32,
    "employeeIdCard": "9" * 32,
    "employeeAddress": fixed_width("北京市朝阳区压力测试地址", 200),
    "companyName": fixed_width("压力测试法律主体公司名称", 160),
    "postName": fixed_width("压力测试岗位", 64),
    "servicePersonType": fixed_width("压力测试劳务人员类型", 32),
    "insuranceType": fixed_width("压力测试商业保险类型", 64),
    "contractStartDate": "9999-12-31",
    "contractEndDate": "9999-12-31",
    "probationStartDate": "9999-12-31",
    "probationEndDate": "9999-12-31",
    "signDate": "9999-12-31",
    "baseSalary": "99999999999999.99",
    "postSalary": "99999999999999.99",
    "fieldAllowance": "99999999999999.99",
    "performanceSalary": "99999999999999.99",
    "salaryTotal": "99999999999999.99",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def ensure_new(paths: list[Path]) -> None:
    existing = [str(path) for path in paths if path.exists()]
    if existing:
        raise RuntimeError(
            "refusing to overwrite existing build artifacts:\n  " + "\n  ".join(existing)
        )


def set_exact_page_break(source: Path, target: Path, marker: str) -> None:
    document = Document(source)
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.strip() == marker]
    if len(matches) != 1:
        raise RuntimeError(
            f"expected one exact paragraph {marker!r} in {source}, found {len(matches)}"
        )
    matches[0].paragraph_format.page_break_before = True
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)

    verification = Document(target)
    verified = [
        paragraph
        for paragraph in verification.paragraphs
        if paragraph.text.strip() == marker
        and paragraph.paragraph_format.page_break_before is True
    ]
    if len(verified) != 1:
        raise RuntimeError(f"active pageBreakBefore verification failed for {target}")


def replace_pressure_values(source: Path, target: Path) -> dict[str, int]:
    document = Document(source)
    replacement_counts = {key: 0 for key in PRESSURE_VALUES}

    def replace_in_paragraph(paragraph) -> None:
        for run in paragraph.runs:
            text = run.text
            for key, value in PRESSURE_VALUES.items():
                placeholder = "${" + key + "}"
                count = text.count(placeholder)
                if count:
                    text = text.replace(placeholder, value)
                    replacement_counts[key] += count
            run.text = text

    def visit_table(table) -> None:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph)
                for nested in cell.tables:
                    visit_table(nested)

    for paragraph in document.paragraphs:
        replace_in_paragraph(paragraph)
    for table in document.tables:
        visit_table(table)
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                replace_in_paragraph(paragraph)
            for table in container.tables:
                visit_table(table)

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    unresolved_parts = []
    with ZipFile(target) as archive:
        for name in archive.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                if b"${" in archive.read(name):
                    unresolved_parts.append(name)
    if unresolved_parts:
        raise RuntimeError(
            f"unresolved placeholders in {target}: {', '.join(unresolved_parts)}"
        )
    return {key: count for key, count in replacement_counts.items() if count}


def build(source_dir: Path, output_dir: Path, qa_dir: Path) -> Path:
    source_dir = source_dir.resolve()
    output_dir = output_dir.resolve()
    qa_dir = qa_dir.resolve()
    manifest_path = output_dir.parent / "manifest-v4-last-page.json"
    output_paths = [output_dir / spec.output_name for spec in SPECS]
    pressure_paths = [qa_dir / spec.output_name.replace(".docx", "_pressure.docx") for spec in SPECS]
    ensure_new(output_paths + pressure_paths + [manifest_path])

    source_hashes_before = {
        spec.source_name: sha256(source_dir / spec.source_name) for spec in SPECS
    }
    artifacts = []
    pressure = []

    for spec, output_path, pressure_path in zip(SPECS, output_paths, pressure_paths):
        source_path = source_dir / spec.source_name
        if not source_path.is_file():
            raise FileNotFoundError(source_path)
        if spec.page_break_marker:
            set_exact_page_break(source_path, output_path, spec.page_break_marker)
        else:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source_path, output_path)

        counts = replace_pressure_values(output_path, pressure_path)
        artifacts.append(
            {
                "file": str(output_path),
                "source": str(source_path),
                "sha256": sha256(output_path),
                "size": output_path.stat().st_size,
                "templateType": spec.template_type,
                "version": "v4-last-page",
                "pageBreakBefore": spec.page_break_marker,
                "companySealRequired": spec.seal_position is not None,
                "sealPlacementMode": "LAST_PAGE" if spec.seal_position else None,
                "recommendedSealPosition": spec.seal_position,
            }
        )
        pressure.append(
            {
                "file": str(pressure_path),
                "sha256": sha256(pressure_path),
                "size": pressure_path.stat().st_size,
                "replacements": counts,
            }
        )

    source_hashes_after = {
        spec.source_name: sha256(source_dir / spec.source_name) for spec in SPECS
    }
    if source_hashes_before != source_hashes_after:
        raise RuntimeError("v3 source files changed during build")

    manifest = {
        "release": "beijing-onboard-sign-20260718-v4-last-page",
        "productionMutation": False,
        "sourceFilesUnchanged": True,
        "sourceSha256": source_hashes_after,
        "artifacts": artifacts,
        "pressureArtifacts": pressure,
        "pressureFieldLengths": {
            key: len(value) for key, value in PRESSURE_VALUES.items()
        },
    }
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return manifest_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--qa-dir", type=Path, default=DEFAULT_QA_DIR)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        manifest_path = build(args.source_dir, args.output_dir, args.qa_dir)
    except Exception as exc:  # fail closed and keep any source untouched
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(manifest_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
