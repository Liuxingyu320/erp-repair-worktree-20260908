#!/usr/bin/env python3
"""Prepare disabled, review-only DOCX candidates from the 2026-07-18 source set.

The builder is intentionally filesystem-only.  It never changes a source DOCX,
does not connect to a database, and does not emit an upload/registration payload.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Final, Iterable
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.text.paragraph import Paragraph


VERSION: Final[str] = "20260724-v8-draft"
SYSTEM_SIGNATURE_COPY: Final[str] = "签名：____________"
SYSTEM_DATE_COPY: Final[str] = "日期：${signDate}"
EVIDENCE_REFERENCE_COPY: Final[str] = (
    "签署时间及文件校验信息见本合同末页《电子签署确认页》。"
)
CHECKBOX_FONT: Final[str] = "DejaVu Sans"
ATTACHMENT_MARKERS: Final[tuple[str, ...]] = (
    "attachmentDormitoryMark",
    "attachmentDutyMark",
    "attachmentHandbookMark",
    "attachmentSalaryMark",
)
SERVICE_PERSON_MARKERS: Final[tuple[str, ...]] = (
    "serviceStudentMark",
    "serviceRetiredMark",
)
INSURANCE_MARKERS: Final[tuple[str, ...]] = (
    "insuranceCommercialAccidentMark",
    "insuranceEmployerLiabilityMark",
)
SERVICE_ATTACHMENT_MARKERS: Final[tuple[str, ...]] = (
    "attachmentDormitoryMark",
    "attachmentDutyMark",
    "attachmentHandbookMark",
    "attachmentServiceReceiptMark",
)
EMPTY_CUSTOM_PROPERTIES_XML: Final[bytes] = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
    b'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"/>'
)

# Values present in the supplied 2026-07-18 source files.  Keeping this map
# explicit prevents an over-broad regex from rewriting employee addresses.
COMPANY_VALUE_REPLACEMENTS: Final[dict[str, str]] = {
    "舟山茗汇文化传播有限公司": "${companyName}",
    "杜翠香": "${companyLegalRepresentative}",
    "浙江省舟山市嵊泗县枸杞乡奇观村育才路9号203室-013工位": "${companyAddress}",
}


@dataclass(frozen=True)
class TemplateSpec:
    source_name: str
    candidate_name: str
    template_type: str
    required: tuple[str, ...]
    editor: Callable[[DocxDocument], None]
    salary_version: str | None = None
    social_type: str | None = None
    employment_type: str | None = None
    post_level_scope: str | None = None
    age_scope: str | None = None
    student_status_scope: str | None = None
    review_notes: tuple[str, ...] = ()
    seal_recommended: bool = False
    supported: bool = True
    blocked_reasons: tuple[str, ...] = ()


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """Replace text while retaining paragraph properties and first-run format."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _set_checkbox_paragraph(paragraph: Paragraph, marker: str, text: str) -> None:
    """Create a replaceable checkbox run without changing the paragraph geometry."""
    first_run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        first_run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)

    marker_run = paragraph.add_run(marker)
    if first_run_properties is not None:
        marker_run._r.insert(0, deepcopy(first_run_properties))
    marker_run.font.name = CHECKBOX_FONT
    marker_fonts = marker_run._r.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        marker_fonts.set(qn("w:" + attribute), CHECKBOX_FONT)

    text_run = paragraph.add_run(text)
    if first_run_properties is not None:
        text_run._r.insert(0, deepcopy(first_run_properties))


def _set_mixed_checkbox_paragraph(
    paragraph: Paragraph, parts: tuple[tuple[str, bool], ...]
) -> None:
    """Replace a paragraph while keeping each inline checkbox in its own glyph-safe run."""
    first_run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        first_run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)

    for text, checkbox in parts:
        run = paragraph.add_run(text)
        if first_run_properties is not None:
            run._r.insert(0, deepcopy(first_run_properties))
        if checkbox:
            run.font.name = CHECKBOX_FONT
            fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn("w:" + attribute), CHECKBOX_FONT)


def _replace_checkbox_first(
    document: DocxDocument, needle: str, marker: str, text: str
) -> None:
    for paragraph in _iter_paragraphs(document):
        if _has_anchor(paragraph.text, needle):
            _set_checkbox_paragraph(paragraph, marker, text)
            return
    raise ValueError(f"required checkbox paragraph anchor not found: {needle!r}")


def _iter_table_paragraphs(table) -> Iterable[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def _iter_paragraphs(document: DocxDocument) -> Iterable[Paragraph]:
    yield from document.paragraphs
    for table in document.tables:
        yield from _iter_table_paragraphs(table)
    for section in document.sections:
        for container in (section.header, section.footer):
            yield from container.paragraphs
            for table in container.tables:
                yield from _iter_table_paragraphs(table)


def _has_anchor(text: str, needle: str) -> bool:
    normalize = lambda value: re.sub(r"\s+", " ", value).strip()
    return normalize(needle) in normalize(text)


def _replace_first(document: DocxDocument, needle: str, text: str) -> None:
    for paragraph in _iter_paragraphs(document):
        if _has_anchor(paragraph.text, needle):
            _set_paragraph_text(paragraph, text)
            return
    raise ValueError(f"required paragraph anchor not found: {needle!r}")


def _replace_all(document: DocxDocument, needle: str, text: str) -> int:
    count = 0
    for paragraph in _iter_paragraphs(document):
        if _has_anchor(paragraph.text, needle):
            _set_paragraph_text(paragraph, text)
            count += 1
    return count


def _replace_exact_all(document: DocxDocument, expected: str, text: str) -> int:
    normalize = lambda value: re.sub(r"\s+", " ", value).strip()
    count = 0
    for paragraph in _iter_paragraphs(document):
        if normalize(paragraph.text) == normalize(expected):
            _set_paragraph_text(paragraph, text)
            count += 1
    return count


def _prefix_first(document: DocxDocument, needle: str, prefix: str) -> None:
    """Prefix a known paragraph while removing only its source fill-in underline."""
    for paragraph in _iter_paragraphs(document):
        if not _has_anchor(paragraph.text, needle):
            continue
        original = paragraph.text.lstrip()
        if not original.startswith(prefix):
            _set_paragraph_text(paragraph, prefix + original)
        for run in paragraph.runs:
            run.underline = False
        return
    raise ValueError(f"required paragraph anchor not found: {needle!r}")


def _fill_underlined_prefix_first(
    document: DocxDocument, needle: str, placeholder: str
) -> Paragraph:
    """Fill a source underline with a placeholder without moving the body text."""
    for paragraph in document.paragraphs:
        if not _has_anchor(paragraph.text, needle):
            continue
        for run in paragraph.runs:
            if run.underline and not run.text.strip():
                run.text = placeholder
                return paragraph
        raise ValueError(
            f"required leading fill-in underline not found for anchor: {needle!r}"
        )
    raise ValueError(f"required paragraph anchor not found: {needle!r}")


def _insert_before(anchor: Paragraph, text: str) -> Paragraph:
    """Insert a style-compatible paragraph immediately before ``anchor``."""
    element = OxmlElement("w:p")
    if anchor._p.pPr is not None:
        element.append(deepcopy(anchor._p.pPr))
    anchor._p.addprevious(element)
    inserted = Paragraph(element, anchor._parent)
    run = inserted.add_run(text)
    if anchor.runs and anchor.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(anchor.runs[0]._r.rPr))
    return inserted


def _insert_before_first(document: DocxDocument, needle: str, text: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if _has_anchor(paragraph.text, needle):
            return _insert_before(paragraph, text)
    raise ValueError(f"required paragraph anchor not found: {needle!r}")


def _require_count(actual: int, expected: int, anchor: str) -> None:
    if actual != expected:
        raise ValueError(
            f"expected {expected} occurrences of {anchor!r}, found {actual}"
        )


def _rewrite_contract_attachments(document: DocxDocument) -> None:
    """Rewrite two receipts and the vertically merged duty-confirmation row."""
    _require_count(
        _replace_all(
            document,
            "乙方签字：",
            f"乙方：${{employeeName}}；{SYSTEM_SIGNATURE_COPY}    {SYSTEM_DATE_COPY}",
        ),
        2,
        "乙方签字：",
    )
    _require_count(
        _replace_all(
            document,
            "确认人（签名及捺印）：",
            f"确认人：${{employeeName}}；{SYSTEM_SIGNATURE_COPY}",
        ),
        1,
        "确认人（签名及捺印）：",
    )
    _require_count(
        _replace_exact_all(
            document, "身份证号码：", "身份证号码：${employeeIdCard}"
        ),
        1,
        "身份证号码：",
    )
    _require_count(
        _replace_exact_all(document, "日期：", SYSTEM_DATE_COPY),
        1,
        "岗位职责附件日期：",
    )


def _clear_company_header_placeholders(document: DocxDocument) -> None:
    """Remove only company-name-only header text from reusable candidates."""
    removable = {"${companyName}", *COMPANY_VALUE_REPLACEMENTS.keys()}
    visited: set[int] = set()
    for section in document.sections:
        header = section.header
        identity = id(header._element)
        if identity in visited:
            continue
        visited.add(identity)
        for paragraph in header.paragraphs:
            if re.sub(r"\s+", "", paragraph.text) in {
                re.sub(r"\s+", "", value) for value in removable
            }:
                _set_paragraph_text(paragraph, "")


def _request_field_update(document: DocxDocument) -> None:
    """Refresh PAGE/NUMPAGES fields when Word or LibreOffice opens the candidate."""
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def _rewrite_dynamic_page_footers(document: DocxDocument) -> None:
    """Remove unsupported WPS fields and reserve a clean PDF-layer footer.

    The supplied contract stores PAGE/NUMPAGES inside a floating WPS text box.
    LibreOffice either ignores that object or retains stale cached values during
    headless conversion.  The runtime therefore stamps the actual page and total
    after PDF conversion; keeping a normal empty footer here reserves the visual
    area without leaking an incorrect cached page number.
    """
    for section in document.sections:
        section.footer.is_linked_to_previous = False
        footer = section.footer
        for child in list(footer._element):
            footer._element.remove(child)
        paragraph = footer.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_together = True


def _remove_redundant_attachment_page_breaks(document: DocxDocument) -> None:
    """Use heading-level pagination only, avoiding a blank overflow page.

    The supplied labor-contract source places a page-break-only paragraph
    immediately before each attachment heading.  Keeping that break while also
    setting ``page_break_before`` on the heading produces a fully blank page
    when a long signature/checklist block flows onto the next page.  Remove
    only those redundant break-only paragraphs; unrelated document page
    breaks remain untouched.
    """
    paragraphs = list(document.paragraphs)
    for index, paragraph in enumerate(paragraphs):
        text = re.sub(r"\s+", " ", paragraph.text).strip()
        if not re.match(r"^附件\s*[12]$", text):
            continue
        previous_index = index - 1
        while previous_index >= 0 and not paragraphs[previous_index].text.strip():
            previous = paragraphs[previous_index]
            page_breaks = [
                element
                for element in previous._p.iter(qn("w:br"))
                if element.get(qn("w:type")) == "page"
            ]
            if page_breaks:
                parent = previous._element.getparent()
                if parent is not None:
                    parent.remove(previous._element)
            previous_index -= 1


def _polish_contract_layout(document: DocxDocument) -> None:
    """Keep short legal units together and prevent orphaned headings/signature rows."""
    _remove_redundant_attachment_page_breaks(document)
    checklist_started = False
    delivery_prefixes = (
        "第十五条",
        "15.1 ",
        "乙方地址：",
        "联系电话：",
        "15.2 ",
        "15.3 ",
    )
    term_prefixes = (
        "1.1 ",
        "${contractTermFixedMark}",
        "${contractTermOpenEndedMark}",
    )
    delivery_paragraphs: list[Paragraph] = []
    term_paragraphs: list[Paragraph] = []
    delivery_started = False
    for paragraph in document.paragraphs:
        text = re.sub(r"\s+", " ", paragraph.text).strip()
        if not text:
            continue
        fmt = paragraph.paragraph_format
        if re.match(r"^第[一二三四五六七八九十百]+条", text):
            fmt.keep_together = True
            fmt.keep_with_next = True
        if text.startswith("第十五条"):
            delivery_started = True
        if delivery_started and text.startswith(delivery_prefixes):
            delivery_paragraphs.append(paragraph)
            fmt.keep_together = True
        elif delivery_started:
            delivery_started = False
        if text.startswith(term_prefixes):
            term_paragraphs.append(paragraph)
            fmt.keep_together = True
        if text.startswith("附件清单"):
            checklist_started = True
            fmt.keep_together = True
            fmt.keep_with_next = True
        elif checklist_started and any(
            text.startswith("${" + marker + "}") for marker in ATTACHMENT_MARKERS
        ):
            fmt.keep_together = True
            fmt.keep_with_next = True
        elif checklist_started:
            checklist_started = False
        if text.startswith(("甲方盖章：", "甲方代表：", EVIDENCE_REFERENCE_COPY)):
            fmt.keep_together = True
            fmt.keep_with_next = True
        if text.startswith("乙方：") and "签名：" in text:
            fmt.keep_together = True
        if re.match(r"^附件\s*[12]$", text):
            fmt.page_break_before = True
            fmt.keep_with_next = True
        if text in ("职工宿舍免责协议书", "员工岗位职责说明书"):
            fmt.keep_together = True
            fmt.keep_with_next = True
    for logical_block in (term_paragraphs, delivery_paragraphs):
        for paragraph in logical_block[:-1]:
            paragraph.paragraph_format.keep_with_next = True
    for table in document.tables:
        _prevent_table_row_splits(table)
        if table.rows and "员工岗位职责说明书" in table.rows[0].cells[0].text:
            _repeat_table_header(table, 2)
    _rewrite_dynamic_page_footers(document)
    _request_field_update(document)


def _prevent_table_row_splits(table) -> None:
    """Keep ordinary rows intact while allowing an oversized note row to flow."""
    for row in table.rows:
        row_text = " ".join(cell.text for cell in row.cells)
        is_oversized_note = bool(row.cells) and (
            row.cells[0].text.strip() == "备注" and len(row_text) > 500
        )
        if not is_oversized_note:
            properties = row._tr.get_or_add_trPr()
            if properties.find(qn("w:cantSplit")) is None:
                properties.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for nested in cell.tables:
                _prevent_table_row_splits(nested)


def _repeat_table_header(table, row_count: int) -> None:
    """Repeat contiguous title/column-header rows on every continuation page."""
    if row_count < 1 or len(table.rows) < row_count:
        raise ValueError("duty table is missing its repeatable header rows")
    for row in table.rows[:row_count]:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:tblHeader")) is None:
            properties.append(OxmlElement("w:tblHeader"))


def _polish_short_document_layout(document: DocxDocument) -> None:
    """Keep short clauses intact and bind legal headings to the next paragraph."""
    for paragraph in document.paragraphs:
        text = re.sub(r"\s+", " ", paragraph.text).strip()
        if not text:
            continue
        paragraph.paragraph_format.keep_together = True
        if re.match(r"^第[一二三四五六七八九十百]+条", text):
            paragraph.paragraph_format.keep_with_next = True


def _polish_confirmation_tables(document: DocxDocument) -> None:
    """Prevent compensation rows from splitting and repeat their column header."""
    _polish_short_document_layout(document)
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs[:-1]):
        text = re.sub(r"\s+", " ", paragraph.text).strip()
        if "签名：" not in text:
            continue
        for next_index in range(index + 1, len(paragraphs)):
            next_text = paragraphs[next_index].text.strip()
            if next_text.startswith("日期："):
                for linked in paragraphs[index:next_index]:
                    linked.paragraph_format.keep_with_next = True
                break
            if next_text:
                break
    for table in document.tables:
        _prevent_table_row_splits(table)
        if table.rows:
            _repeat_table_header(table, 1)
    _request_field_update(document)


def _replace_literals(document: DocxDocument) -> None:
    for paragraph in _iter_paragraphs(document):
        text = paragraph.text
        replaced = text
        for source, placeholder in COMPANY_VALUE_REPLACEMENTS.items():
            replaced = replaced.replace(source, placeholder)
        if replaced != text:
            _set_paragraph_text(paragraph, replaced)


def _insert_after(anchor: Paragraph, text: str) -> Paragraph:
    element = OxmlElement("w:p")
    if anchor._p.pPr is not None:
        element.append(deepcopy(anchor._p.pPr))
    anchor._p.addnext(element)
    inserted = Paragraph(element, anchor._parent)
    run = inserted.add_run(text)
    if anchor.runs and anchor.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(anchor.runs[0]._r.rPr))
    return inserted


def _insert_after_first(document: DocxDocument, needle: str, text: str) -> None:
    for paragraph in document.paragraphs:
        if _has_anchor(paragraph.text, needle):
            _insert_after(paragraph, text)
            return
    raise ValueError(f"required paragraph anchor not found: {needle!r}")


def _remove_empty_paragraphs_before_first(document: DocxDocument, needle: str) -> None:
    """Remove only contiguous source spacer paragraphs before a known anchor."""
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if not _has_anchor(paragraph.text, needle):
            continue
        for spacer in reversed(paragraphs[:index]):
            if spacer.text.strip():
                break
            spacer._element.getparent().remove(spacer._element)
        return
    raise ValueError(f"required paragraph anchor not found: {needle!r}")


def _set_amount_by_label(document: DocxDocument, label: str, placeholder: str) -> None:
    for table in document.tables:
        for row in table.rows:
            if not row.cells or label not in row.cells[0].text:
                continue
            if len(row.cells) < 2:
                raise ValueError(f"amount row has no value cell: {label!r}")
            paragraph = row.cells[1].paragraphs[0]
            _set_paragraph_text(paragraph, "${" + placeholder + "}")
            return
    raise ValueError(f"required amount row not found: {label!r}")


def _edit_commitment(document: DocxDocument) -> None:
    _replace_first(
        document,
        "本人（姓名：",
        "本人（姓名：${employeeName}，身份证号：${employeeIdCard}）郑重承诺：",
    )
    _replace_first(document, "承诺人签字", f"承诺人{SYSTEM_SIGNATURE_COPY}")
    _replace_first(document, "日期：", SYSTEM_DATE_COPY)


def _edit_labor_contract(document: DocxDocument) -> None:
    _replace_first(document, "甲方（用工单位）", "甲方（用工单位）：${companyName}")
    _replace_first(document, "乙方（劳 动 者）", "乙方（劳 动 者）：${employeeName}")
    _replace_first(document, "身 份 证 号", "身 份 证 号 码：${employeeIdCard}")
    _replace_first(document, "签 订 日 期", SYSTEM_DATE_COPY)
    _replace_first(
        document,
        "姓名：",
        "姓名：${employeeName} 公民身份号码：${employeeIdCard}",
    )
    _replace_all(document, "家庭住址：", "家庭住址：${employeeAddress}")
    _replace_all(document, "联系电话：", "联系电话：${employeePhone}")
    _require_count(
        _replace_all(document, "乙方地址：", "乙方地址：${employeeAddress}"),
        1,
        "乙方地址：",
    )
    _replace_first(
        document,
        "1.1 双方协商同意",
        "1.1 双方协商同意，劳动合同期限采取下列第"
        "${contractTermSelection}种形式。",
    )
    _replace_checkbox_first(
        document,
        "A、固定期限",
        "${contractTermFixedMark}",
        " A、固定期限：自${contractStartDate}起至${contractEndDate}止，"
        "其中试用期从${probationStartDate}至${probationEndDate}止。",
    )
    _replace_checkbox_first(
        document,
        "B、无固定期限",
        "${contractTermOpenEndedMark}",
        " B、无固定期限：自_____年____月____日起至法定的劳动合同终止条件出现时止。",
    )
    _replace_first(
        document,
        "3.1 乙方服从甲方工作安排",
        "3.1 乙方服从甲方工作安排，从事${postName}工作岗位。"
        "具体任务、职责及工作标准详见甲方的《岗位职责说明书》、操作规范等文件，"
        "以及甲方管理人员的安排和要求。",
    )
    _replace_first(
        document,
        "5.1 乙方转正后满勤底薪",
        "5.1 乙方转正后满勤底薪为${baseSalary}元/月，试用期底薪为2800元/月。"
        "底薪已包含固定休息日加班的全部补偿。",
    )
    _replace_first(
        document,
        "甲方（盖章）",
        f"甲方盖章：____________    乙方{SYSTEM_SIGNATURE_COPY}",
    )
    _replace_first(
        document,
        "授权代表（签字）",
        f"甲方代表：____________    {SYSTEM_DATE_COPY}",
    )
    _insert_after_first(document, "甲方代表：", EVIDENCE_REFERENCE_COPY)
    for marker, number, title in (
        ("attachmentDormitoryMark", 1, "《职工宿舍免责协议书》"),
        ("attachmentDutyMark", 2, "《岗位职责说明书》"),
        ("attachmentHandbookMark", 3, "《员工手册》"),
        ("attachmentSalaryMark", 4, "《薪酬结构确认书》"),
    ):
        _replace_checkbox_first(
            document, f"□{number}", "${" + marker + "}", f" {number}  {title}"
        )
    dormitory_intro = _fill_underlined_prefix_first(
        document, "为了方便员工", "${companyName}"
    )
    dormitory_intro.paragraph_format.keep_together = True
    _rewrite_contract_attachments(document)
    _polish_contract_layout(document)


def _edit_handbook_receipt(document: DocxDocument) -> None:
    _replace_first(
        document,
        "本人（姓名：",
        "本人（姓名：${employeeName}，身份证号：${employeeIdCard}）确认：",
    )
    _replace_first(document, "员工签字", f"员工{SYSTEM_SIGNATURE_COPY}")
    _replace_first(document, "日期：", SYSTEM_DATE_COPY)


def _edit_salary_confirmation(document: DocxDocument) -> None:
    _replace_first(
        document,
        "本人（姓名：",
        "本人（姓名：${employeeName}（手印），身份证号：${employeeIdCard}），"
        "系贵司员工。现就本人薪酬结构事宜，经与公司平等协商，确认如下：",
    )
    _replace_first(
        document,
        "本人月综合工资标准",
        "本人月综合工资标准为人民币${salaryTotal}元。",
    )
    for label, placeholder in (
        ("底薪", "baseSalary"),
        ("综合岗位津贴", "postSalary"),
        ("综合驻外补贴", "fieldAllowance"),
        ("月度绩效津贴", "performanceSalary"),
    ):
        _set_amount_by_label(document, label, placeholder)
    # A has a total row; B intentionally carries the total in the total-standard
    # paragraph instead.  Populate the row only when the source provides one.
    try:
        _set_amount_by_label(document, "综合工资合计", "salaryTotal")
    except ValueError:
        pass
    _replace_first(document, "员工签字", f"员工{SYSTEM_SIGNATURE_COPY}")
    _replace_first(document, "日期：", SYSTEM_DATE_COPY)
    _remove_empty_paragraphs_before_first(document, "员工签名：")
    _remove_empty_paragraphs_before_first(document, SYSTEM_DATE_COPY)
    _polish_confirmation_tables(document)


def _edit_service_contract(document: DocxDocument) -> None:
    _replace_first(document, "甲方（用人单位）", "甲方（用人单位）：${companyName}")
    _replace_first(document, "乙方（劳务人员）", "乙方（劳务人员）：${employeeName}")
    _replace_first(document, "身 份 证 号", "身 份 证 号 码：${employeeIdCard}")
    _replace_first(document, "签 订 日 期", SYSTEM_DATE_COPY)
    _replace_first(
        document,
        "姓名：",
        "姓名：${employeeName} 公民身份号码：${employeeIdCard}",
    )
    _replace_all(document, "家庭住址：", "家庭住址：${employeeAddress}")
    _replace_all(document, "联系电话：", "联系电话：${employeePhone}")
    person_type = next(
        paragraph for paragraph in _iter_paragraphs(document)
        if _has_anchor(paragraph.text, "乙方为（")
    )
    _set_mixed_checkbox_paragraph(
        person_type,
        (
            ("乙方为（", False),
            ("${serviceStudentMark}", True),
            (" 在校实习生 / ", False),
            ("${serviceRetiredMark}", True),
            (" 退休返聘人员），不具备与甲方建立法定劳动关系的主体资格。", False),
        ),
    )
    _replace_first(
        document,
        "1.1 本协议期限",
        "1.1 本协议期限自${contractStartDate}起至${contractEndDate}止。",
    )
    _replace_first(
        document,
        "3.1 乙方服从甲方工作安排",
        "3.1 乙方服从甲方工作安排，从事${postName}工作岗位。"
        "具体任务、职责及工作标准详见甲方的相关文件及管理人员的安排和要求。",
    )
    _replace_first(
        document,
        "5.1 乙方满勤底薪",
        "5.1 乙方满勤底薪为${baseSalary}元/月。底薪已包含固定休息日加班的全部补偿。",
    )
    _replace_all(document, "劳务协议签收单", "劳务合同书签收单")
    _replace_all(document, "乙方地址：", "乙方地址：${employeeAddress}")
    _replace_first(
        document,
        "甲方（盖章）",
        f"甲方盖章：____________    乙方{SYSTEM_SIGNATURE_COPY}",
    )
    _replace_first(
        document,
        "授权代表（签字）",
        f"甲方代表：____________    {SYSTEM_DATE_COPY}",
    )
    _insert_after_first(document, "甲方代表：", EVIDENCE_REFERENCE_COPY)
    for marker, number, title in (
        ("attachmentDormitoryMark", 1, "《职工宿舍免责协议书》"),
        ("attachmentDutyMark", 2, "《岗位职责说明书》"),
        ("attachmentHandbookMark", 3, "《员工手册》"),
        ("attachmentServiceReceiptMark", 4, "《劳务合同书签收单》"),
    ):
        _replace_checkbox_first(
            document, f"□{number}", "${" + marker + "}", f" {number}  {title}"
        )
    dormitory_intro = _fill_underlined_prefix_first(
        document, "为了方便员工", "甲方（提供方）：${companyName}。"
    )
    dormitory_intro.paragraph_format.keep_together = True
    _rewrite_contract_attachments(document)
    _polish_contract_layout(document)


def _edit_service_receipt(document: DocxDocument) -> None:
    identity = next(
        paragraph for paragraph in _iter_paragraphs(document)
        if _has_anchor(paragraph.text, "本人（姓名：")
    )
    _set_mixed_checkbox_paragraph(
        identity,
        (
            ("本人（姓名：${employeeName}，身份证号：${employeeIdCard}），"
             "系${companyName}劳务人员（", False),
            ("${serviceStudentMark}", True),
            (" 在校实习生 / ", False),
            ("${serviceRetiredMark}", True),
            (" 退休返聘人员）。", False),
        ),
    )
    _replace_first(
        document,
        "本人劳务报酬标准",
        "本人劳务报酬标准为每月人民币${salaryTotal}元。具体构成为：",
    )
    insurance = next(
        paragraph for paragraph in _iter_paragraphs(document)
        if _has_anchor(paragraph.text, "甲方已为本人购买")
    )
    _set_mixed_checkbox_paragraph(
        insurance,
        (
            ("2、甲方已为本人购买（", False),
            ("${insuranceCommercialAccidentMark}", True),
            (" 商业意外保险 / ", False),
            ("${insuranceEmployerLiabilityMark}", True),
            (" 雇主责任险），用于覆盖本人在提供劳务期间可能发生的人身意外风险。", False),
        ),
    )
    for label, placeholder in (
        ("底薪", "baseSalary"),
        ("综合岗位津贴", "postSalary"),
        ("综合驻外补贴", "fieldAllowance"),
        ("月度绩效津贴", "performanceSalary"),
    ):
        _set_amount_by_label(document, label, placeholder)
    _replace_first(
        document,
        "甲方代表签字",
        f"甲方代表：____________    乙方{SYSTEM_SIGNATURE_COPY}",
    )
    _replace_first(document, "日期：", SYSTEM_DATE_COPY)
    _polish_confirmation_tables(document)


def _edit_confidential(document: DocxDocument) -> None:
    _replace_first(document, "乙方（员工）", "乙方（员工）：${employeeName}")
    _replace_first(document, "身份证号：", "身份证号：${employeeIdCard}")
    _replace_first(document, "岗位：", "岗位：${postName} 职级：${postLevel}级")
    _insert_after_first(
        document,
        "岗位：",
        "联系电话：${employeePhone}    通讯地址：${employeeAddress}    "
        "合同起始日：${contractStartDate}",
    )
    _replace_first(
        document,
        "甲方（盖章）",
        f"甲方盖章：____________    乙方{SYSTEM_SIGNATURE_COPY}",
    )
    _replace_first(
        document,
        "授权代表（签字）",
        f"甲方代表：____________    {SYSTEM_DATE_COPY}",
    )
    _insert_after_first(document, "甲方代表：", EVIDENCE_REFERENCE_COPY)
    _polish_short_document_layout(document)
    _request_field_update(document)


def _edit_minor_declaration(document: DocxDocument) -> None:
    _replace_first(
        document,
        "本人（姓名：",
        "本人（姓名：${employeeName}，身份证号：${employeeIdCard}），郑重声明如下：",
    )
    _replace_first(
        document,
        "本人自",
        "本人自${incomeStartYearMonth}起，以个人劳动收入作为主要生活来源，"
        "经济独立，无需依赖父母或他人供养。",
    )
    _remove_empty_paragraphs_before_first(document, "声明人签字")
    _replace_first(document, "声明人签字", f"声明人{SYSTEM_SIGNATURE_COPY}")
    _replace_first(document, "日期：", SYSTEM_DATE_COPY)
    # The source uses 1.5-line spacing and pushes only the signature/date onto
    # an otherwise empty second page.  A local 1.3 setting keeps the complete
    # declaration together without changing any clause or font size.
    for paragraph in document.paragraphs:
        paragraph.paragraph_format.line_spacing = 1.3


SPECS: Final[tuple[TemplateSpec, ...]] = (
    TemplateSpec(
        "2-3入职承诺书.docx",
        "05_ONBOARD_COMMITMENT.docx",
        "ONBOARD_COMMITMENT",
        ("employeeName", "employeeIdCard", "signDate"),
        _edit_commitment,
        review_notes=("同时适用于劳动合同和劳务合同，作为入职套餐通用文件",),
    ),
    TemplateSpec(
        "2-4劳动合同.docx",
        "09_ONBOARD_LABOR_CONTRACT.docx",
        "ONBOARD_LABOR_CONTRACT",
        (
            "employeeName",
            "employeeIdCard",
            "employeePhone",
            "employeeAddress",
            "companyName",
            "companyAddress",
            "companyLegalRepresentative",
            "contractStartDate",
            "contractEndDate",
            "probationStartDate",
            "probationEndDate",
            "contractTermSelection",
            "contractTermFixedMark",
            "contractTermOpenEndedMark",
            "postName",
            "baseSalary",
            "signDate",
            *ATTACHMENT_MARKERS,
        ),
        _edit_labor_contract,
        employment_type="劳动合同",
        review_notes=(
            "试用期底薪仍沿用源模板固定值2800元；系统暂无probationBaseSalary字段，启用前需确认",
        ),
        seal_recommended=True,
    ),
    TemplateSpec(
        "2-5员工手册签收确认书.docx",
        "10_ONBOARD_HANDBOOK_RECEIPT.docx",
        "ONBOARD_HANDBOOK_RECEIPT",
        ("employeeName", "employeeIdCard", "signDate"),
        _edit_handbook_receipt,
        employment_type="劳动合同",
    ),
    TemplateSpec(
        "2-6薪酬结构确认书（A版）.docx",
        "11_ONBOARD_SALARY_CONFIRM_A.docx",
        "ONBOARD_SALARY_CONFIRM",
        (
            "employeeName",
            "employeeIdCard",
            "baseSalary",
            "postSalary",
            "fieldAllowance",
            "performanceSalary",
            "salaryTotal",
            "signDate",
        ),
        _edit_salary_confirmation,
        salary_version="A",
        social_type="SOCIAL_UNINSURED",
        employment_type="劳动合同",
    ),
    TemplateSpec(
        "2-6薪酬结构确认书（B版）.docx",
        "12_ONBOARD_SALARY_CONFIRM_B.docx",
        "ONBOARD_SALARY_CONFIRM",
        (
            "employeeName",
            "employeeIdCard",
            "baseSalary",
            "postSalary",
            "fieldAllowance",
            "performanceSalary",
            "salaryTotal",
            "signDate",
        ),
        _edit_salary_confirmation,
        salary_version="B",
        social_type="SOCIAL_INSURED",
        employment_type="劳动合同",
    ),
    TemplateSpec(
        "2-7劳务合同.docx",
        "13_ONBOARD_SERVICE_CONTRACT.docx",
        "ONBOARD_SERVICE_CONTRACT",
        (
            "employeeName",
            "employeeIdCard",
            "employeePhone",
            "employeeAddress",
            "companyName",
            "companyAddress",
            "companyLegalRepresentative",
            *SERVICE_PERSON_MARKERS,
            "contractStartDate",
            "contractEndDate",
            "postName",
            "baseSalary",
            "signDate",
            *SERVICE_ATTACHMENT_MARKERS,
        ),
        _edit_service_contract,
        employment_type="劳务合同",
        seal_recommended=True,
    ),
    TemplateSpec(
        "2-8劳务合同书签收单.docx",
        "14_ONBOARD_SERVICE_RECEIPT.docx",
        "ONBOARD_SERVICE_RECEIPT",
        (
            "employeeName",
            "employeeIdCard",
            "companyName",
            *SERVICE_PERSON_MARKERS,
            "baseSalary",
            "postSalary",
            "fieldAllowance",
            "performanceSalary",
            "salaryTotal",
            *INSURANCE_MARKERS,
            "signDate",
        ),
        _edit_service_receipt,
        employment_type="劳务合同",
    ),
    TemplateSpec(
        "2-9保密与竞业限制协议（7级及以上）.docx",
        "15_ONBOARD_CONFIDENTIAL_NONCOMPETE.docx",
        "ONBOARD_CONFIDENTIAL_NONCOMPETE",
        (
            "employeeName",
            "employeeIdCard",
            "employeePhone",
            "employeeAddress",
            "postName",
            "postLevel",
            "contractStartDate",
            "signDate",
        ),
        _edit_confidential,
        post_level_scope="7级及以上",
        review_notes=("同时适用于劳动合同和劳务合同，仅按员工职级 7 级及以上匹配",),
        seal_recommended=True,
    ),
    TemplateSpec(
        "2-10非在校生及未成年工入职声明书.docx",
        "16_ONBOARD_MINOR_NONSTUDENT_DECLARATION.docx",
        "ONBOARD_MINOR_NONSTUDENT_DECLARATION",
        ("employeeName", "employeeIdCard", "incomeStartYearMonth", "signDate"),
        _edit_minor_declaration,
        age_scope="16-18",
        student_status_scope="NON_STUDENT",
        review_notes=("仅匹配签约时年龄满 16 周岁且未满 18 周岁、并经 HR 审核为非在校的员工",),
    ),
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _clean_core_properties(document: DocxDocument) -> None:
    properties = document.core_properties
    for name in (
        "author",
        "last_modified_by",
        "comments",
        "identifier",
        "keywords",
        "category",
        "subject",
        "title",
    ):
        setattr(properties, name, "")
    properties.revision = 1
    # python-docx requires datetime values for these two fields.  Retain the
    # source timestamps; unlike author/last_modified_by they contain no name.


def _normalize_docx_archive(path: Path) -> None:
    """Normalize ZIP metadata so an unchanged source builds byte-for-byte."""
    normalized = path.with_suffix(path.suffix + ".normalized")
    with ZipFile(path, "r") as source, ZipFile(
        normalized, "w", compression=ZIP_DEFLATED, compresslevel=9
    ) as target:
        for original in source.infolist():
            info = ZipInfo(original.filename, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            data = source.read(original.filename)
            if original.filename == "docProps/custom.xml":
                # WPS may store account-linked KSOTemplateDocerSaveRecord and
                # ICV values here.  Keep the valid part/relationship but empty
                # every custom property in the review candidate.
                data = EMPTY_CUSTOM_PROPERTIES_XML
            target.writestr(info, data)
    os.replace(normalized, path)


def _find_source(source_dir: Path, name: str) -> Path:
    matches = [path for path in source_dir.rglob(name) if path.is_file()]
    if len(matches) != 1:
        raise FileNotFoundError(
            f"expected exactly one {name!r} below {source_dir}, found {len(matches)}"
        )
    return matches[0]


def _document_text(document: DocxDocument) -> str:
    return "\n".join(paragraph.text for paragraph in _iter_paragraphs(document))


def _validate_candidate(document: DocxDocument, spec: TemplateSpec) -> None:
    text = _document_text(document)
    missing = [name for name in spec.required if "${" + name + "}" not in text]
    if missing:
        raise ValueError(
            f"{spec.candidate_name} is missing required placeholders: {', '.join(missing)}"
        )
    remaining = [value for value in COMPANY_VALUE_REPLACEMENTS if value in text]
    if remaining:
        raise ValueError(
            f"{spec.candidate_name} still contains hard-coded company values: {remaining}"
        )
    if SYSTEM_SIGNATURE_COPY not in text or SYSTEM_DATE_COPY not in text:
        raise ValueError(f"{spec.candidate_name} is missing system signing copy")
    if "见附加页" in text:
        raise ValueError(f"{spec.candidate_name} still contains ambiguous appendix signing copy")
    if spec.seal_recommended and EVIDENCE_REFERENCE_COPY not in text:
        raise ValueError(f"{spec.candidate_name} is missing the evidence-page reference")
    header_text = "\n".join(
        paragraph.text
        for section in document.sections
        for paragraph in section.header.paragraphs
    )
    if "${companyName}" in header_text or any(
        company_name in header_text for company_name in COMPANY_VALUE_REPLACEMENTS
    ):
        raise ValueError(f"{spec.candidate_name} still contains a company header")


def _prepare_one(source: Path, target: Path, spec: TemplateSpec) -> None:
    document = Document(source)
    spec.editor(document)
    _replace_literals(document)
    _clear_company_header_placeholders(document)
    _request_field_update(document)
    _clean_core_properties(document)
    _validate_candidate(document, spec)

    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_suffix(target.suffix + ".tmp")
    document.save(temporary)
    _normalize_docx_archive(temporary)
    os.replace(temporary, target)

    # Re-open the serialized artifact; validation before save alone can miss a
    # packaging/core-properties regression.
    serialized = Document(target)
    _validate_candidate(serialized, spec)
    properties = serialized.core_properties
    if any(
        getattr(properties, name)
        for name in (
            "author",
            "last_modified_by",
            "comments",
            "identifier",
            "keywords",
            "category",
            "subject",
            "title",
        )
    ) or properties.revision != 1:
        raise ValueError(f"core properties were not sanitized in {target}")


def build(source_dir: Path, output_dir: Path, manifest_path: Path) -> Path:
    source_dir = source_dir.resolve()
    output_dir = output_dir.resolve()
    manifest_path = manifest_path.resolve()
    if not source_dir.is_dir():
        raise NotADirectoryError(source_dir)
    if output_dir == source_dir or source_dir in output_dir.parents:
        raise ValueError("output directory must be outside the source directory")
    if manifest_path == source_dir or source_dir in manifest_path.parents:
        raise ValueError("manifest must be outside the source directory")

    resolved = [(spec, _find_source(source_dir, spec.source_name)) for spec in SPECS]
    source_hashes = {spec.source_name: _sha256(path) for spec, path in resolved}
    output_dir.mkdir(parents=True, exist_ok=True)

    artifacts: list[dict[str, object]] = []
    for spec, source in resolved:
        target = output_dir / spec.candidate_name
        _prepare_one(source, target, spec)
        artifact: dict[str, object] = {
            "source": {
                "file": spec.source_name,
                "sha256": source_hashes[spec.source_name],
                "size": source.stat().st_size,
            },
            "candidate": {
                "file": spec.candidate_name,
                "sha256": _sha256(target),
                "size": target.stat().st_size,
            },
            "version": VERSION,
            # In this system 1 means disabled.  Every generated artifact is
            # review-only until a separate, authorized publication step.
            "status": "1",
            "type": spec.template_type,
            "required": list(spec.required),
            "supportStatus": "supported" if spec.supported else "unsupported",
            "activationStatus": "disabled" if spec.supported else "blocked",
            "blockedReasons": list(spec.blocked_reasons),
            "signaturePolicy": {
                "mode": "INLINE_ORIGINAL_POSITION",
                "evidencePageMode": "FIXED_PLACEHOLDER",
                "evidencePageCount": 1,
            },
            "sealPolicy": {
                "status": "pending",
                "recommended": spec.seal_recommended,
            },
        }
        if spec.salary_version is not None:
            artifact["salaryVersion"] = spec.salary_version
        if spec.social_type is not None:
            artifact["socialType"] = spec.social_type
        if spec.employment_type is not None:
            artifact["employmentType"] = spec.employment_type
        if spec.post_level_scope is not None:
            artifact["postLevelScope"] = spec.post_level_scope
        if spec.age_scope is not None:
            artifact["ageScope"] = spec.age_scope
        if spec.student_status_scope is not None:
            artifact["studentStatusScope"] = spec.student_status_scope
        if spec.review_notes:
            artifact["reviewNotes"] = list(spec.review_notes)
        artifacts.append(artifact)

    changed_sources = [
        spec.source_name
        for spec, path in resolved
        if _sha256(path) != source_hashes[spec.source_name]
    ]
    if changed_sources:
        raise RuntimeError(f"source files changed during build: {changed_sources}")

    manifest = {
        "manifestVersion": 1,
        "version": VERSION,
        "statusMeaning": {"1": "disabled"},
        "registrationPayloadGenerated": False,
        "releaseBlockers": [
            "当前可追溯源集缺少《员工手册》正文原始模板；发布新方案前必须另行绑定并校验 ONBOARD_HANDBOOK",
        ],
        "bindingConstraints": [
            {
                "type": "ONBOARD_SALARY_CONFIRM",
                "rule": "有社保固定匹配B版，无社保固定匹配A版；不允许越过社保状态手工选版",
                "selection": "derivedFromSocialType",
                "variants": ["A", "B"],
                "mapping": {
                    "SOCIAL_INSURED": "B",
                    "SOCIAL_UNINSURED": "A",
                },
            }
        ],
        "packageBindingRules": {
            "common": ["ONBOARD_COMMITMENT"],
            "labor": [
                "ONBOARD_LABOR_CONTRACT",
                "ONBOARD_HANDBOOK",
                "ONBOARD_HANDBOOK_RECEIPT",
                "ONBOARD_SALARY_CONFIRM",
            ],
            "service": [
                "ONBOARD_SERVICE_CONTRACT",
                "ONBOARD_SERVICE_RECEIPT",
            ],
            "conditional": [
                {
                    "type": "ONBOARD_CONFIDENTIAL_NONCOMPETE",
                    "employmentTypes": ["劳动合同", "劳务合同"],
                    "minimumJobGrade": 7,
                },
                {
                    "type": "ONBOARD_MINOR_NONSTUDENT_DECLARATION",
                    "studentStatus": "NON_STUDENT",
                    "minimumAgeInclusive": 16,
                    "maximumAgeExclusive": 18,
                    "requiredFact": "incomeStartYearMonth",
                },
            ],
        },
        "attachmentChecklistPolicy": {
            "source": "actualPackageDocuments",
            "allowedMarks": ["☑", "□"],
            "markers": sorted(set(ATTACHMENT_MARKERS + SERVICE_ATTACHMENT_MARKERS)),
            "showHash": False,
            "showTemplateVersion": False,
        },
        "evidencePagePolicy": {
            "mode": "FIXED_PLACEHOLDER",
            "pageCount": 1,
            "placeholderSuppliedBy": "PDF_LAYER",
            "docxAddsBlankPage": False,
        },
        "pageNumberPolicy": {
            "mode": "PDF_LAYER_DYNAMIC",
            "templateTypes": ["ONBOARD_LABOR_CONTRACT"],
            "format": "第 {page} 页 共 {pages} 页",
            "fontSizePoints": 9,
            "baselinePoints": 36,
            "docxFooter": "RESERVED_EMPTY",
        },
        "artifacts": artifacts,
    }
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_manifest = manifest_path.with_suffix(manifest_path.suffix + ".tmp")
    temporary_manifest.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    os.replace(temporary_manifest, manifest_path)
    return manifest_path


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source-dir",
        required=True,
        type=Path,
        help="directory containing the nine source DOCX files",
    )
    parser.add_argument(
        "--output-dir",
        required=True,
        type=Path,
        help="separate directory for disabled candidates",
    )
    parser.add_argument(
        "--manifest",
        required=True,
        type=Path,
        help="explicit path for the review manifest",
    )
    return parser.parse_args()


def main() -> int:
    args = _parse_args()
    manifest = build(args.source_dir, args.output_dir, args.manifest)
    print(f"prepared {len(SPECS)} disabled candidates")
    print(manifest)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
