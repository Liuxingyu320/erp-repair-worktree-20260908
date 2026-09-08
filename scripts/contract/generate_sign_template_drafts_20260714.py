#!/usr/bin/env python3
"""Build disabled contract-signing template drafts for local review.

The script never overwrites currently enabled template files. It creates:
- an onboarding v3 copy set with the fixed legal entity replaced by ${companyName};
- 15 explicitly watermarked technical drafts for regularize/transfer/renewal/offboard;
- a hash/size manifest for idempotent database registration.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
from dataclasses import dataclass
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE_ONBOARD = ROOT / "uploadPath/sign-template/onboard-20260713-v2"
TARGET_ONBOARD = ROOT / "uploadPath/sign-template/onboard-20260714-v3-draft"
TARGET_SCENARIOS = ROOT / "uploadPath/sign-template/scenario-draft-20260714"
DESKTOP_ROOT = Path(
    "/Users/liuxingyu/Desktop/入离调转20260702-改造版-20260713/06_五场景模板草案-20260714"
)

FIXED_LEGAL_ENTITY = "舟山茗汇文化传播有限公司"
COMPANY_TOKEN = "${companyName}"
DRAFT_WARNING = "系统测试草案｜未经法务与HR书面确认，不得启用、发送或作为正式文件"
CHINESE_FONT = "Arial Unicode MS"

# Only mark real data-table headers. The other three tables in the offer notice are
# form-layout tables, so marking their first rows as headers would be semantically wrong.
ONBOARD_HEADER_TABLES = {
    "01_ONBOARD_HANDBOOK.docx": (0,),
    "04_ONBOARD_OFFER_NOTICE.docx": (2,),
    "12_ONBOARD_SALARY_CONFIRM_B.docx": (0,),
    "14_ONBOARD_SERVICE_RECEIPT.docx": (0,),
}
PAGE_NUMBER_ALT = "页脚页码：当前页及总页数"


@dataclass(frozen=True)
class DraftSpec:
    template_type: str
    filename: str
    title: str
    scenario: str
    employment_type: str | None
    employee_sign_required: bool
    required_placeholders: tuple[str, ...]
    sections: tuple[tuple[str, tuple[str, ...]], ...]


ONBOARD_FILES = (
    (19, "01_ONBOARD_HANDBOOK.docx", "ONBOARD_HANDBOOK", "员工手册"),
    (22, "04_ONBOARD_OFFER_NOTICE.docx", "ONBOARD_OFFER_NOTICE", "录用通知书"),
    (23, "05_ONBOARD_COMMITMENT.docx", "ONBOARD_COMMITMENT", "入职承诺书"),
    (24, "06_ONBOARD_POST_DUTY_2_4.docx", "ONBOARD_POST_DUTY", "岗位职责确认书（2-4级）"),
    (25, "07_ONBOARD_POST_DUTY_5_6.docx", "ONBOARD_POST_DUTY", "岗位职责确认书（5-6级）"),
    (26, "08_ONBOARD_POST_DUTY_7_8.docx", "ONBOARD_POST_DUTY", "岗位职责确认书（7-8级）"),
    (27, "09_ONBOARD_LABOR_CONTRACT.docx", "ONBOARD_LABOR_CONTRACT", "劳动合同"),
    (28, "10_ONBOARD_HANDBOOK_RECEIPT.docx", "ONBOARD_HANDBOOK_RECEIPT", "员工手册签收确认书"),
    (29, "11_ONBOARD_SALARY_CONFIRM_A.docx", "ONBOARD_SALARY_CONFIRM", "薪酬结构确认书（A版）"),
    (30, "12_ONBOARD_SALARY_CONFIRM_B.docx", "ONBOARD_SALARY_CONFIRM", "薪酬结构确认书（B版）"),
    (31, "13_ONBOARD_SERVICE_CONTRACT.docx", "ONBOARD_SERVICE_CONTRACT", "劳务合同"),
    (32, "14_ONBOARD_SERVICE_RECEIPT.docx", "ONBOARD_SERVICE_RECEIPT", "劳务合同签收单"),
    (33, "15_ONBOARD_CONFIDENTIAL_NONCOMPETE.docx", "ONBOARD_CONFIDENTIAL_NONCOMPETE", "保密与竞业限制协议（7级及以上）"),
)


SPECS = (
    DraftSpec(
        "REGULARIZE_CONFIRMATION", "01_REGULARIZE_CONFIRMATION.docx", "转正确认书（系统测试草案）",
        "regularize", "劳动合同", True,
        ("employeeName", "employeeIdCard", "actualRegularizationDate", "postName", "postLevel", "signDate"),
        (("一、主体与人员", ("用人法律主体：${companyName}", "员工：${employeeName}，证件号：${employeeIdCard}")),
         ("二、转正信息", ("实际转正日期：${actualRegularizationDate}", "岗位：${postName}，岗位等级：${postLevel}", "正式文本应由HR核对试用期考核、生效日和其他条件后发布。"))),
    ),
    DraftSpec(
        "REGULARIZE_POST_DUTY", "02_REGULARIZE_POST_DUTY.docx", "转正岗位职责确认书（系统测试草案）",
        "regularize", "劳动合同", True,
        ("employeeName", "postName", "postLevel", "actualRegularizationDate", "signDate"),
        (("一、基本信息", ("用人法律主体：${companyName}", "员工：${employeeName}", "岗位：${postName}，岗位等级：${postLevel}", "转正生效日：${actualRegularizationDate}")),
         ("二、职责确认", ("具体岗位职责、考核口径与授权边界须由业务负责人和HR在正式版本中确认。",))),
    ),
    DraftSpec(
        "REGULARIZE_SALARY_CONFIRM", "03_REGULARIZE_SALARY_CONFIRM.docx", "转正薪资确认书（系统测试草案）",
        "regularize", "劳动合同", True,
        ("employeeName", "employeeIdCard", "actualRegularizationDate", "baseSalary", "postSalary", "fieldAllowance", "performanceSalary", "salaryTotal", "salaryVersion", "signDate"),
        (("一、基本信息", ("用人法律主体：${companyName}", "员工：${employeeName}，证件号：${employeeIdCard}", "转正生效日：${actualRegularizationDate}，薪酬版本：${salaryVersion}")),
         ("二、薪酬结构", ("基本工资：${baseSalary}", "岗位工资：${postSalary}", "外勤补贴：${fieldAllowance}", "绩效工资：${performanceSalary}", "合计：${salaryTotal}"))),
    ),
    DraftSpec(
        "TRANSFER_CONFIRMATION", "04_TRANSFER_CONFIRMATION.docx", "调岗确认书（系统测试草案）",
        "transfer", None, True,
        ("employeeName", "employeeIdCard", "transferEffectiveDate", "beforeDeptName", "afterDeptName", "beforePostName", "afterPostName", "signDate"),
        (("一、主体与人员", ("用人法律主体：${companyName}", "员工：${employeeName}，证件号：${employeeIdCard}")),
         ("二、调岗信息", ("生效日：${transferEffectiveDate}", "调整前：${beforeDeptName} / ${beforePostName}", "调整后：${afterDeptName} / ${afterPostName}", "正式文本须由HR确认工作地点、职责、薪酬及其他变更项。"))),
    ),
    DraftSpec(
        "TRANSFER_POST_DUTY", "05_TRANSFER_POST_DUTY.docx", "调岗岗位职责确认书（系统测试草案）",
        "transfer", None, True,
        ("employeeName", "afterDeptName", "afterPostName", "postLevel", "transferEffectiveDate", "signDate"),
        (("一、基本信息", ("用人法律主体：${companyName}", "员工：${employeeName}", "调整后部门：${afterDeptName}", "调整后岗位：${afterPostName}，岗位等级：${postLevel}", "生效日：${transferEffectiveDate}")),
         ("二、职责确认", ("具体职责、考核指标和授权边界须在正式版本中确认。",))),
    ),
    DraftSpec(
        "TRANSFER_SALARY_CONFIRM", "06_TRANSFER_SALARY_CONFIRM.docx", "调岗薪资确认书（系统测试草案）",
        "transfer", None, True,
        ("employeeName", "employeeIdCard", "transferEffectiveDate", "baseSalary", "postSalary", "fieldAllowance", "performanceSalary", "salaryTotal", "salaryVersion", "signDate"),
        (("一、基本信息", ("用人法律主体：${companyName}", "员工：${employeeName}，证件号：${employeeIdCard}", "调岗生效日：${transferEffectiveDate}，薪酬版本：${salaryVersion}")),
         ("二、薪酬结构", ("基本工资：${baseSalary}", "岗位工资：${postSalary}", "外勤补贴：${fieldAllowance}", "绩效工资：${performanceSalary}", "合计：${salaryTotal}"))),
    ),
    DraftSpec(
        "RENEWAL_LABOR_CONTRACT", "07_RENEWAL_LABOR_CONTRACT.docx", "劳动合同续签协议（系统测试草案）",
        "renewal", "劳动合同", True,
        ("employeeName", "employeeIdCard", "employeePhone", "employeeAddress", "companyName", "previousContractEndDate", "previousEmploymentType", "previousRenewalCount", "renewalCount", "contractStartDate", "contractEndDate", "postName", "baseSalary", "signDate"),
        (("一、双方信息", ("甲方：${companyName}", "乙方：${employeeName}，证件号：${employeeIdCard}", "联系电话：${employeePhone}，联系地址：${employeeAddress}")),
         ("二、续签快照", ("原合同类型：${previousEmploymentType}", "原合同结束日：${previousContractEndDate}", "原续签次数：${previousRenewalCount}，本次续签次数：${renewalCount}")),
         ("三、新合同信息", ("期限：${contractStartDate} 至 ${contractEndDate}", "岗位：${postName}，基本工资：${baseSalary}", "其他条款必须由法务与HR以正式文本确认，本草案仅验证系统字段和签署链路。"))),
    ),
    DraftSpec(
        "RENEWAL_SERVICE_CONTRACT", "08_RENEWAL_SERVICE_CONTRACT.docx", "劳务协议续签书（系统测试草案）",
        "renewal", "劳务合同", True,
        ("employeeName", "employeeIdCard", "employeePhone", "employeeAddress", "companyName", "previousContractEndDate", "previousEmploymentType", "previousRenewalCount", "renewalCount", "servicePersonType", "contractStartDate", "contractEndDate", "postName", "baseSalary", "signDate"),
        (("一、双方信息", ("甲方：${companyName}", "乙方：${employeeName}，证件号：${employeeIdCard}", "联系电话：${employeePhone}，联系地址：${employeeAddress}")),
         ("二、续签快照", ("原合同类型：${previousEmploymentType}", "原合同结束日：${previousContractEndDate}", "原续签次数：${previousRenewalCount}，本次续签次数：${renewalCount}")),
         ("三、新协议信息", ("人员类型：${servicePersonType}", "期限：${contractStartDate} 至 ${contractEndDate}", "服务岗位：${postName}，基本报酬：${baseSalary}", "正式劳务条款须经法务与HR确认。"))),
    ),
    DraftSpec(
        "RENEWAL_SALARY_CONFIRM", "09_RENEWAL_SALARY_CONFIRM.docx", "续签薪酬确认书（系统测试草案）",
        "renewal", None, True,
        ("employeeName", "employeeIdCard", "companyName", "previousRenewalCount", "renewalCount", "contractStartDate", "contractEndDate", "baseSalary", "postSalary", "fieldAllowance", "performanceSalary", "salaryTotal", "salaryVersion", "signDate"),
        (("一、基本信息", ("法律主体：${companyName}", "员工：${employeeName}，证件号：${employeeIdCard}", "原续签次数：${previousRenewalCount}，本次续签次数：${renewalCount}", "新期限：${contractStartDate} 至 ${contractEndDate}", "薪酬版本：${salaryVersion}")),
         ("二、薪酬结构", ("基本工资：${baseSalary}", "岗位工资：${postSalary}", "外勤补贴：${fieldAllowance}", "绩效工资：${performanceSalary}", "合计：${salaryTotal}"))),
    ),
    DraftSpec(
        "OFFBOARD_CONFIRMATION", "10_OFFBOARD_CONFIRMATION.docx", "离职确认书（系统测试草案）",
        "offboard", None, True,
        ("employeeName", "employeeIdCard", "entryDate", "leaveDate", "offboardingType", "leaveReason", "signDate"),
        (("一、主体与人员", ("用人法律主体：${companyName}", "员工：${employeeName}，证件号：${employeeIdCard}")),
         ("二、离职信息", ("入职日：${entryDate}", "离职日：${leaveDate}", "离职类型：${offboardingType}", "离职原因：${leaveReason}"))),
    ),
    DraftSpec(
        "OFFBOARD_HANDOVER", "11_OFFBOARD_HANDOVER.docx", "离职交接确认书（系统测试草案）",
        "offboard", None, True,
        ("employeeName", "leaveDate", "postName", "assetHandoverStatus", "signDate"),
        (("一、基本信息", ("法律主体：${companyName}", "员工：${employeeName}，岗位：${postName}", "离职日：${leaveDate}")),
         ("二、交接结果", ("资产交接状态：${assetHandoverStatus}", "正式文本须附交接清单、接收人和异常说明。"))),
    ),
    DraftSpec(
        "OFFBOARD_SETTLEMENT", "12_OFFBOARD_SETTLEMENT.docx", "离职结算确认书（系统测试草案）",
        "offboard", None, True,
        ("employeeName", "employeeIdCard", "leaveDate", "salarySettlementStatus", "compensationAmount", "compensationNote", "signDate"),
        (("一、基本信息", ("法律主体：${companyName}", "员工：${employeeName}，证件号：${employeeIdCard}", "离职日：${leaveDate}")),
         ("二、结算信息", ("薪资结算状态：${salarySettlementStatus}", "补偿金额：${compensationAmount}", "补偿说明：${compensationNote}", "正式文本应附明细和计算口径，由HR/财务/法务确认。"))),
    ),
    DraftSpec(
        "OFFBOARD_CONFIDENTIALITY_NONCOMPETE", "13_OFFBOARD_CONFIDENTIALITY_NONCOMPETE.docx", "离职保密与竞业确认书（系统测试草案）",
        "offboard", None, True,
        ("employeeName", "employeeIdCard", "leaveDate", "nonCompeteDecision", "signDate"),
        (("一、基本信息", ("法律主体：${companyName}", "员工：${employeeName}，证件号：${employeeIdCard}", "离职日：${leaveDate}")),
         ("二、决定与提示", ("竞业限制决定：${nonCompeteDecision}", "保密义务、竞业范围、期限、补偿及违约责任必须由法务在正式文本中单独确认。"))),
    ),
    DraftSpec(
        "OFFBOARD_TERMINATION_NOTICE", "14_OFFBOARD_TERMINATION_NOTICE.docx", "解除/终止通知书（系统测试草案）",
        "offboard", None, False,
        ("employeeName", "employeeIdCard", "entryDate", "leaveDate", "offboardingType", "leaveReason", "companyName", "signDate"),
        (("一、通知对象", ("出具主体：${companyName}", "员工：${employeeName}，证件号：${employeeIdCard}")),
         ("二、事实快照", ("入职日：${entryDate}", "解除/终止日：${leaveDate}", "类型：${offboardingType}", "原因：${leaveReason}", "送达方式、法律依据、异议渠道和具体文案必须由法务审核。"))),
    ),
    DraftSpec(
        "OFFBOARD_LEAVE_CERTIFICATE", "15_OFFBOARD_LEAVE_CERTIFICATE.docx", "离职证明（系统测试草案）",
        "offboard", None, False,
        ("employeeName", "employeeIdCard", "entryDate", "leaveDate", "postName", "companyName", "signDate"),
        (("一、证明事项", ("员工：${employeeName}，证件号：${employeeIdCard}", "于 ${entryDate} 入职，岗位为 ${postName}，于 ${leaveDate} 离职。", "本草案仅用于测试字段渲染，正式证明格式和内容由HR/法务确认。")),
         ("二、出具信息", ("出具主体：${companyName}", "出具日期：${signDate}"))),
    ),
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def set_run_font(run, name: str, size: float, *, bold: bool = False,
                 color: RGBColor | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hint"), "eastAsia")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def configure_document(doc: Document, spec: DraftSpec) -> None:
    section = doc.sections[0]
    # Named legal-template override: Chinese business forms use A4 rather than Letter.
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = CHINESE_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), CHINESE_FONT)
    normal._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.5

    for style_name, size in (("Heading 1", 13), ("Heading 2", 11.5), ("Heading 3", 11)):
        style = doc.styles[style_name]
        style.font.name = CHINESE_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), CHINESE_FONT)
        style._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run(DRAFT_WARNING), CHINESE_FONT, 9, bold=True,
                 color=RGBColor(0x9B, 0x1C, 0x1C))

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    set_run_font(fp.add_run(f"{spec.template_type} ｜ {DRAFT_WARNING}"), CHINESE_FONT, 8,
                 color=RGBColor(0x66, 0x66, 0x66))

    props = doc.core_properties
    props.title = spec.title
    props.subject = "ERP签约系统测试模板草案"
    props.author = "ERP合同系统"
    props.last_modified_by = "ERP合同系统"
    props.comments = DRAFT_WARNING


def add_body_line(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.widow_control = True
    set_run_font(paragraph.add_run(text), CHINESE_FONT, 10.5)


def create_draft(spec: DraftSpec, target: Path) -> None:
    doc = Document()
    configure_document(doc, spec)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(4)
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.keep_with_next = True
    set_run_font(title.add_run(spec.title), CHINESE_FONT, 18, bold=True)

    warning = doc.add_paragraph()
    warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
    warning.paragraph_format.space_after = Pt(12)
    warning.paragraph_format.keep_with_next = True
    set_run_font(warning.add_run(DRAFT_WARNING), CHINESE_FONT, 10, bold=True,
                 color=RGBColor(0x9B, 0x1C, 0x1C))

    for heading, lines in spec.sections:
        paragraph = doc.add_paragraph(style="Heading 1")
        set_run_font(paragraph.add_run(heading), CHINESE_FONT, 13, bold=True,
                     color=RGBColor(0x1F, 0x3A, 0x5F))
        for line in lines:
            add_body_line(doc, line)

    section_numbers = "一二三四五六七八九十"
    signature_number = section_numbers[len(spec.sections)]
    signature_heading = doc.add_paragraph(style="Heading 1")
    set_run_font(signature_heading.add_run(f"{signature_number}、签署/出具"), CHINESE_FONT, 13,
                 bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))
    if spec.employee_sign_required:
        add_body_line(doc, "员工签名：____________________")
        add_body_line(doc, "签署日期：${signDate}")
    else:
        add_body_line(doc, "本文件为公司单方出具文件，不要求员工签名。")
        add_body_line(doc, "出具日期：${signDate}")

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    text = collect_docx_text(target)
    missing = [f"${{{name}}}" for name in spec.required_placeholders
               if f"${{{name}}}" not in text]
    if missing:
        raise RuntimeError(f"{target.name} missing placeholders: {missing}")
    if COMPANY_TOKEN not in text:
        raise RuntimeError(f"{target.name} missing {COMPANY_TOKEN}")


def copy_zip_info(info: ZipInfo) -> ZipInfo:
    copied = ZipInfo(info.filename, date_time=info.date_time)
    copied.compress_type = ZIP_DEFLATED
    copied.comment = info.comment
    copied.extra = info.extra
    copied.create_system = info.create_system
    copied.external_attr = info.external_attr
    copied.internal_attr = info.internal_attr
    copied.flag_bits = info.flag_bits
    return copied


def add_page_number_alt_text(data: bytes) -> tuple[bytes, int]:
    fixes = 0

    def replacement(match: re.Match[bytes]) -> bytes:
        nonlocal fixes
        attributes = match.group(1)
        if b"descr=" in attributes or b"title=" in attributes:
            return match.group(0)
        fixes += 1
        alt = PAGE_NUMBER_ALT.encode("utf-8")
        return b"<wp:docPr" + attributes + b' title="' + alt + b'" descr="' + alt + b'"/>'

    return re.sub(rb"<wp:docPr\b([^>]*)/>", replacement, data), fixes


def mark_data_table_headers(data: bytes, table_indices: tuple[int, ...]) -> tuple[bytes, int]:
    fixes = 0
    for table_index in sorted(table_indices, reverse=True):
        table_starts = [match.start() for match in re.finditer(rb"<w:tbl>", data)]
        if table_index >= len(table_starts):
            raise RuntimeError(f"missing table index {table_index}; found {len(table_starts)}")
        table_start = table_starts[table_index]
        table_end = data.find(b"</w:tbl>", table_start)
        first_row = re.search(rb"<w:tr(?:\s[^>]*)?>", data[table_start:table_end])
        if first_row is None:
            raise RuntimeError(f"table {table_index} has no row")
        row_open_end = table_start + first_row.end()
        first_cell = data.find(b"<w:tc>", row_open_end, table_end)
        row_properties = data.find(b"<w:trPr>", row_open_end, first_cell)
        if row_properties >= 0:
            properties_end = data.find(b"</w:trPr>", row_properties, first_cell)
            if properties_end < 0:
                raise RuntimeError(f"table {table_index} has malformed row properties")
            if b"<w:tblHeader" in data[row_properties:properties_end]:
                continue
            data = data[:properties_end] + b"<w:tblHeader/>" + data[properties_end:]
        else:
            data = (data[:row_open_end] + b"<w:trPr><w:tblHeader/></w:trPr>"
                    + data[row_open_end:])
        fixes += 1
    return data, fixes


def patch_onboard_document(source: Path, target: Path) -> tuple[int, int]:
    target.parent.mkdir(parents=True, exist_ok=True)
    replacements = 0
    accessibility_fixes = 0
    with ZipFile(source, "r") as src, ZipFile(target, "w", compression=ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename.endswith(".xml"):
                count = data.count(FIXED_LEGAL_ENTITY.encode("utf-8"))
                if count:
                    data = data.replace(FIXED_LEGAL_ENTITY.encode("utf-8"),
                                        COMPANY_TOKEN.encode("utf-8"))
                    replacements += count
            if info.filename.startswith("word/footer") and info.filename.endswith(".xml"):
                data, fixes = add_page_number_alt_text(data)
                accessibility_fixes += fixes
            if info.filename == "word/document.xml" and source.name in ONBOARD_HEADER_TABLES:
                data, fixes = mark_data_table_headers(
                    data, ONBOARD_HEADER_TABLES[source.name])
                accessibility_fixes += fixes
            dst.writestr(copy_zip_info(info), data)
    return replacements, accessibility_fixes


def collect_docx_text(path: Path) -> str:
    with ZipFile(path) as archive:
        raw = "".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        )
    return re.sub(r"<[^>]+>", "", raw)


def copy_to_desktop(source: Path, relative: Path) -> Path:
    target = DESKTOP_ROOT / relative
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, target)
    return target


def manifest_row(path: Path, **values: object) -> dict[str, object]:
    return {
        **values,
        "path": str(path),
        "file_url": "/profile/" + str(path.relative_to(ROOT / "uploadPath")),
        "file_name": path.name,
        "file_size": path.stat().st_size,
        "file_hash": sha256(path),
    }


def main() -> None:
    if not SOURCE_ONBOARD.is_dir():
        raise SystemExit(f"missing source directory: {SOURCE_ONBOARD}")
    TARGET_ONBOARD.mkdir(parents=True, exist_ok=True)
    TARGET_SCENARIOS.mkdir(parents=True, exist_ok=True)
    DESKTOP_ROOT.mkdir(parents=True, exist_ok=True)

    rows: list[dict[str, object]] = []
    for source_template_id, filename, template_type, template_name in ONBOARD_FILES:
        source = SOURCE_ONBOARD / filename
        target = TARGET_ONBOARD / filename
        replacement_count, accessibility_fix_count = patch_onboard_document(source, target)
        if FIXED_LEGAL_ENTITY in collect_docx_text(target):
            raise RuntimeError(f"hard-coded legal entity remains in {target}")
        copy_to_desktop(target, Path("01_入职v3法律主体动态版") / filename)
        rows.append(manifest_row(
            target,
            group="onboard_v3_draft",
            source_template_id=source_template_id,
            template_type=template_type,
            template_name=f"{template_name}（法律主体动态草案）",
            scenario="onboard",
            replacement_count=replacement_count,
            accessibility_fix_count=accessibility_fix_count,
            status="1",
        ))

    scenario_dir_by_code = {
        "regularize": "02_转正",
        "transfer": "03_调岗",
        "renewal": "04_续签",
        "offboard": "05_离职",
    }
    for spec in SPECS:
        target = TARGET_SCENARIOS / spec.filename
        create_draft(spec, target)
        copy_to_desktop(target, Path(scenario_dir_by_code[spec.scenario]) / spec.filename)
        rows.append(manifest_row(
            target,
            group="scenario_draft",
            source_template_id=None,
            template_type=spec.template_type,
            template_name=spec.title,
            scenario=spec.scenario,
            employment_type=spec.employment_type,
            required_placeholders=",".join(spec.required_placeholders),
            employee_visible="Y",
            read_confirmation_required="Y" if spec.employee_sign_required else "N",
            employee_sign_required="Y" if spec.employee_sign_required else "N",
            status="1",
        ))

    manifest = {
        "generated_at": "2026-07-14",
        "warning": DRAFT_WARNING,
        "root": str(ROOT),
        "templates": rows,
    }
    manifest_path = DESKTOP_ROOT / "template-draft-manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
                             encoding="utf-8")
    summary = [
        "# 五场景签约模板草案说明（2026-07-14）",
        "",
        f"> {DRAFT_WARNING}",
        "",
        "- 入职v3：13份，不覆盖当前启用文件；已将固定公司名替换为 `${companyName}`，并补充页码替代说明和数据表表头语义。",
        "- 转正：3份系统测试草案。",
        "- 调岗：3份系统测试草案。",
        "- 续签：3份系统测试草案，包含原合同日期/类型和续签次数快照。",
        "- 离职：6份系统测试草案；解除通知和离职证明不要求员工签名。",
        "- 所有新文件只能以停用状态注册，法务/HR复核前不得发布方案。",
        "",
        "完整 SHA-256、文件大小和系统路径见 `template-draft-manifest.json`。",
    ]
    (DESKTOP_ROOT / "请先看我.md").write_text("\n".join(summary) + "\n", encoding="utf-8")
    print(json.dumps({"count": len(rows), "manifest": str(manifest_path)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
