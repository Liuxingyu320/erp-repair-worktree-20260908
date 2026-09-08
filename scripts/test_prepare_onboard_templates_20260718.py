from __future__ import annotations

import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document


REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))

from scripts.contract import prepare_onboard_templates_20260718 as builder
from scripts.contract import build_beijing_sign_safe_templates_v6_20260720 as v6_builder


FIXED_COMPANY_NAME = "舟山茗汇文化传播有限公司"
FIXED_LEGAL_REPRESENTATIVE = "杜翠香"
FIXED_COMPANY_ADDRESS = "浙江省舟山市嵊泗县枸杞乡奇观村育才路9号203室-013工位"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def inject_private_custom_properties(path: Path) -> None:
    with ZipFile(path, "r") as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Override PartName="/docProps/custom.xml" '
        b'ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/>'
        b"</Types>",
    )
    parts["_rels/.rels"] = parts["_rels/.rels"].replace(
        b"</Relationships>",
        b'<Relationship Id="rIdPrivate" '
        b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" '
        b'Target="docProps/custom.xml"/></Relationships>',
    )
    parts["docProps/custom.xml"] = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
        b'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        b'<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" '
        b'name="KSOTemplateDocerSaveRecord"><vt:lpwstr>'
        b'eyJoZGlkIjoicHJpdmF0ZSIsInVzZXJJZCI6IjEyMzQ1In0='
        b'</vt:lpwstr></property><property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" '
        b'pid="3" name="ICV"><vt:lpwstr>PRIVATE_ICV</vt:lpwstr></property></Properties>'
    )
    temporary = path.with_suffix(".private.docx")
    with ZipFile(temporary, "w", compression=ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    temporary.replace(path)


def add_metadata(document: Document) -> None:
    properties = document.core_properties
    properties.author = "应清理的作者"
    properties.last_modified_by = "应清理的修改人"
    properties.comments = "private comment"
    properties.identifier = "private-id"
    properties.keywords = "private keywords"
    properties.category = "private category"
    properties.subject = "private subject"
    properties.title = "private title"
    properties.revision = 9


def add_amount_table(document: Document, labels: list[str]) -> None:
    table = document.add_table(rows=1, cols=3)
    for index, value in enumerate(("项目", "金额（元/月）", "说明")):
        table.cell(0, index).text = value
    for label in labels:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = ""
        cells[2].text = "fixture"


def add_contract_paragraphs(document: Document, service: bool) -> None:
    document.sections[0].header.paragraphs[0].text = FIXED_COMPANY_NAME
    if service:
        document.add_paragraph("乙方（劳务人员）：________________")
    else:
        document.add_paragraph("乙方（劳 动 者）：________________")
    document.add_paragraph("身   份   证   号   码：________________")
    document.add_paragraph("签   订   日   期：____年__月__日")
    company_label = "用人单位" if service else "用工单位"
    document.add_paragraph(f"甲方（{company_label}）：_{FIXED_COMPANY_NAME}_")
    document.add_paragraph(f"法定代表人：{FIXED_LEGAL_REPRESENTATIVE}")
    document.add_paragraph(f"注册地：{FIXED_COMPANY_ADDRESS}")
    document.add_paragraph("姓名：________ 公民身份号码：________")
    document.add_paragraph("家庭住址：________")
    document.add_paragraph("联系电话：________")
    document.add_paragraph("第十五条 送达地址")
    document.add_paragraph("15.1 乙方确定下列地址为唯一有效通讯地址：")
    document.add_paragraph("乙方地址：____________________________________")
    document.add_paragraph("15.2 甲方按上述地址邮寄法律文书。")
    if service:
        document.add_paragraph("乙方为（□ 在校实习生 / □ 退休返聘人员）")
        document.add_paragraph("1.1 本协议期限自____起至____止。")
    else:
        document.add_paragraph("1.1 双方协商同意，劳动合同期限采取下列第____种形式。")
        document.add_paragraph("A、固定期限：自____起至____止。")
        document.add_paragraph(
            "B、无固定期限：自_____年____月____日起至法定的劳动合同终止条件出现时止。"
        )
    document.add_paragraph("3.1 乙方服从甲方工作安排，从事____工作岗位。")
    if service:
        document.add_paragraph("5.1 乙方满勤底薪为____元/月。")
    else:
        document.add_paragraph("5.1 乙方转正后满勤底薪为3300元/月。")
    document.add_paragraph("甲方（盖章）：____ 乙方（签字并按手印）：____")
    document.add_paragraph("授权代表（签字）：____ 签订日期：____")
    if not service:
        document.add_paragraph("附件清单（乙方确认已签收以下文件）：")
        document.add_paragraph("□1 《职工宿舍免责协议书》（文后）")
        document.add_paragraph("□2 《岗位职责说明书》（文后）")
        document.add_paragraph("□3 《员工手册》")
        document.add_paragraph("□4 《薪酬结构确认书》")
        document.add_page_break()
        document.add_paragraph("附件 1")
        document.add_paragraph("职工宿舍免责协议书")
        intro = document.add_paragraph()
        leading = intro.add_run("                       ")
        leading.underline = True
        intro.add_run("为了方便员工，员工可申请住宿。")
        document.add_page_break()
        document.add_paragraph("附件 2")
        document.add_paragraph("员工岗位职责说明书")
        duty_table = document.add_table(rows=3, cols=5)
        for cell in duty_table.rows[0].cells:
            cell.text = "员工岗位职责说明书"
        for index, value in enumerate(("类别", "岗位", "具体内容", "适用岗位", "对应指导文件")):
            duty_table.rows[1].cells[index].text = value
        for index, value in enumerate(("基础工作", "礼仪", "服务标准", "全员", "培训手册")):
            duty_table.rows[2].cells[index].text = value
    else:
        document.add_paragraph("附件清单（乙方确认已签收以下文件）：")
        document.add_paragraph("□1 《职工宿舍免责协议书》（文后）")
        document.add_paragraph("□2 《岗位职责说明书》（文后）")
        document.add_paragraph("□3 《员工手册》")
        document.add_paragraph("□4 《薪酬结构确认书》")
        document.add_page_break()
        document.add_paragraph("附件 1")
        document.add_paragraph("职工宿舍免责协议书")
        intro = document.add_paragraph()
        leading = intro.add_run("                       ")
        leading.underline = True
        intro.add_run("为了方便员工，员工可申请住宿。")
        document.add_page_break()
        document.add_paragraph("附件 2")
        document.add_paragraph("员工岗位职责说明书")
        duty_table = document.add_table(rows=3, cols=5)
        for cell in duty_table.rows[0].cells:
            cell.text = "员工岗位职责说明书"
        for index, value in enumerate(("类别", "岗位", "具体内容", "适用岗位", "对应指导文件")):
            duty_table.rows[1].cells[index].text = value
        for index, value in enumerate(("基础工作", "礼仪", "服务标准", "全员", "培训手册")):
            duty_table.rows[2].cells[index].text = value
    document.add_paragraph("乙方签字：____ 日期：____")
    document.add_paragraph("乙方签字：____ 日期：____")
    document.add_paragraph("确认人（签名及捺印）：")
    document.add_paragraph("身份证号码：")
    document.add_paragraph("日期：")


def save_fixture(path: Path, paragraphs: list[str], table_labels: list[str] | None = None) -> None:
    document = Document()
    add_metadata(document)
    for text in paragraphs:
        document.add_paragraph(text)
    if table_labels:
        add_amount_table(document, table_labels)
    document.save(path)


def write_fixtures(source_dir: Path) -> None:
    source_dir.mkdir(parents=True)

    commitment = Document()
    add_metadata(commitment)
    identity = commitment.add_paragraph()
    first_run = identity.add_run("本人（姓名：")
    first_run.bold = True
    identity.add_run("____，身份证号：____）郑重承诺：")
    commitment.add_paragraph("承诺人签字（并按手印）：____")
    commitment.add_paragraph("日期：____")
    commitment_path = source_dir / "2-3入职承诺书.docx"
    commitment.save(commitment_path)
    inject_private_custom_properties(commitment_path)

    labor = Document()
    add_metadata(labor)
    add_contract_paragraphs(labor, service=False)
    labor.save(source_dir / "2-4劳动合同.docx")

    save_fixture(
        source_dir / "2-5员工手册签收确认书.docx",
        [
            f"本人（姓名：____，身份证号：____）确认：已收到{FIXED_COMPANY_NAME}手册。",
            "员工签字（并按手印）：____",
            "日期：____",
        ],
    )

    salary_labels = ["底薪", "综合岗位津贴", "综合驻外补贴", "月度绩效津贴"]
    for version in ("A", "B"):
        labels = salary_labels + (["综合工资合计"] if version == "A" else [])
        save_fixture(
            source_dir / f"2-6薪酬结构确认书（{version}版）.docx",
            [
                f"致：{FIXED_COMPANY_NAME}",
                "本人（姓名：____，身份证号：____），系贵司员工。",
                "本人月综合工资标准为人民币____元。",
                "员工签字（并按手印）：____",
                "日期：____",
            ],
            labels,
        )

    service = Document()
    add_metadata(service)
    add_contract_paragraphs(service, service=True)
    service.save(source_dir / "2-7劳务合同.docx")

    save_fixture(
        source_dir / "2-8劳务合同书签收单.docx",
        [
            f"本人（姓名：____，身份证号：____），系{FIXED_COMPANY_NAME}（□在校实习生）。",
            "本人劳务报酬标准为每月人民币____元。",
            "2、甲方已为本人购买（□商业意外保险）。",
            "甲方代表签字：____ 乙方签字：____",
            "日期：____",
        ],
        salary_labels,
    )

    save_fixture(
        source_dir / "2-9保密与竞业限制协议（7级及以上）.docx",
        [
            f"甲方（用人单位）：{FIXED_COMPANY_NAME}",
            f"法定代表人：{FIXED_LEGAL_REPRESENTATIVE}",
            f"注册地：{FIXED_COMPANY_ADDRESS}",
            "乙方（员工）：____",
            "身份证号：____",
            "岗位：____职级：____级",
            "甲方（盖章）：____ 乙方（签字并按手印）：____",
            "授权代表（签字）：____ 签订日期：____",
        ],
    )

    save_fixture(
        source_dir / "2-10非在校生及未成年工入职声明书.docx",
        [
            "本人（姓名：____，身份证号：____），郑重声明如下：",
            "本人目前非在校学生。",
            "本人自____年____月起，以个人劳动收入作为主要生活来源。",
            "声明人签字（并按手印）：____",
            "日期：____",
        ],
    )


def document_text(path: Path) -> str:
    document = Document(path)
    return builder._document_text(document)


class PrepareOnboardTemplatesTest(unittest.TestCase):
    def test_builds_repeatable_disabled_candidates_and_review_manifest(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source"
            output = root / "candidates"
            manifest = root / "review" / "candidate-manifest.json"
            write_fixtures(source)
            source_hashes = {path.name: sha256(path) for path in source.glob("*.docx")}

            result = builder.build(source, output, manifest)
            self.assertEqual(result, manifest.resolve())
            self.assertEqual(
                source_hashes,
                {path.name: sha256(path) for path in source.glob("*.docx")},
            )

            payload = json.loads(manifest.read_text(encoding="utf-8"))
            self.assertEqual(payload["version"], builder.VERSION)
            self.assertFalse(payload["registrationPayloadGenerated"])
            self.assertTrue(payload["releaseBlockers"])
            binding_rules = payload["packageBindingRules"]
            self.assertEqual(binding_rules["common"], ["ONBOARD_COMMITMENT"])
            self.assertIn("ONBOARD_HANDBOOK", binding_rules["labor"])
            self.assertEqual(
                binding_rules["service"],
                ["ONBOARD_SERVICE_CONTRACT", "ONBOARD_SERVICE_RECEIPT"],
            )
            salary_binding = payload["bindingConstraints"][0]
            self.assertEqual(salary_binding["selection"], "derivedFromSocialType")
            self.assertEqual(
                salary_binding["mapping"],
                {"SOCIAL_INSURED": "B", "SOCIAL_UNINSURED": "A"},
            )
            self.assertEqual(
                payload["attachmentChecklistPolicy"]["allowedMarks"], ["☑", "□"]
            )
            self.assertEqual(payload["evidencePagePolicy"]["pageCount"], 1)
            self.assertFalse(payload["evidencePagePolicy"]["docxAddsBlankPage"])
            self.assertEqual(
                payload["pageNumberPolicy"],
                {
                    "mode": "PDF_LAYER_DYNAMIC",
                    "templateTypes": ["ONBOARD_LABOR_CONTRACT"],
                    "format": "第 {page} 页 共 {pages} 页",
                    "fontSizePoints": 9,
                    "baselinePoints": 36,
                    "docxFooter": "RESERVED_EMPTY",
                },
            )
            confidential_rule = next(
                item for item in binding_rules["conditional"]
                if item["type"] == "ONBOARD_CONFIDENTIAL_NONCOMPETE"
            )
            self.assertEqual(confidential_rule["minimumJobGrade"], 7)
            self.assertEqual(confidential_rule["employmentTypes"], ["劳动合同", "劳务合同"])
            minor_rule = next(
                item for item in binding_rules["conditional"]
                if item["type"] == "ONBOARD_MINOR_NONSTUDENT_DECLARATION"
            )
            self.assertEqual(minor_rule["studentStatus"], "NON_STUDENT")
            self.assertEqual(
                (minor_rule["minimumAgeInclusive"], minor_rule["maximumAgeExclusive"]),
                (16, 18),
            )
            self.assertEqual(minor_rule["requiredFact"], "incomeStartYearMonth")
            self.assertEqual(len(payload["artifacts"]), 9)
            self.assertEqual(list(output.glob("*.docx")).__len__(), 9)

            by_file = {item["candidate"]["file"]: item for item in payload["artifacts"]}
            for candidate_name, item in by_file.items():
                candidate = output / candidate_name
                text = document_text(candidate)
                self.assertEqual(item["status"], "1")
                self.assertEqual(item["version"], builder.VERSION)
                self.assertEqual(item["candidate"]["sha256"], sha256(candidate))
                self.assertEqual(item["candidate"]["size"], candidate.stat().st_size)
                self.assertEqual(item["signaturePolicy"]["mode"], "INLINE_ORIGINAL_POSITION")
                self.assertEqual(item["signaturePolicy"]["evidencePageCount"], 1)
                self.assertEqual(item["sealPolicy"]["status"], "pending")
                self.assertIn(builder.SYSTEM_SIGNATURE_COPY, text)
                self.assertIn(builder.SYSTEM_DATE_COPY, text)
                self.assertNotIn("见附加页", text)
                for required in item["required"]:
                    self.assertIn("${" + required + "}", text)
                self.assertNotIn(FIXED_COMPANY_NAME, text)
                self.assertNotIn(FIXED_LEGAL_REPRESENTATIVE, text)
                self.assertNotIn(FIXED_COMPANY_ADDRESS, text)
                self.assertTrue(all(
                    not paragraph.text.strip()
                    for section in Document(candidate).sections
                    for paragraph in section.header.paragraphs
                ))

                properties = Document(candidate).core_properties
                for key in (
                    "author",
                    "last_modified_by",
                    "comments",
                    "identifier",
                    "keywords",
                    "category",
                    "subject",
                    "title",
                ):
                    self.assertFalse(getattr(properties, key))
                self.assertEqual(properties.revision, 1)

            labor = by_file["09_ONBOARD_LABOR_CONTRACT.docx"]
            self.assertEqual(labor["employmentType"], "劳动合同")
            labor_text = document_text(output / labor["candidate"]["file"])
            self.assertIn("转正后满勤底薪为${baseSalary}元/月", labor_text)
            self.assertIn("试用期底薪为2800元/月", labor_text)
            self.assertIn("固定休息日加班的全部补偿", labor_text)
            self.assertIn("probationBaseSalary", labor["reviewNotes"][0])
            commitment_manifest = by_file["05_ONBOARD_COMMITMENT.docx"]
            self.assertNotIn("employmentType", commitment_manifest)
            self.assertIn("劳动合同和劳务合同", commitment_manifest["reviewNotes"][0])
            self.assertEqual(
                by_file["10_ONBOARD_HANDBOOK_RECEIPT.docx"]["employmentType"],
                "劳动合同",
            )

            receipt_text = document_text(output / "14_ONBOARD_SERVICE_RECEIPT.docx")
            self.assertIn("劳务报酬标准为每月人民币${salaryTotal}元", receipt_text)
            self.assertIn("${performanceSalary}", receipt_text)
            self.assertIn("${companyName}劳务人员", receipt_text)
            for marker in builder.SERVICE_PERSON_MARKERS + builder.INSURANCE_MARKERS:
                self.assertIn("${" + marker + "}", receipt_text)
            self.assertNotIn("购买（${insuranceType}）", receipt_text)
            self.assertEqual(by_file["14_ONBOARD_SERVICE_RECEIPT.docx"]["employmentType"], "劳务合同")

            salary_a = by_file["11_ONBOARD_SALARY_CONFIRM_A.docx"]
            salary_b = by_file["12_ONBOARD_SALARY_CONFIRM_B.docx"]
            self.assertEqual((salary_a["salaryVersion"], salary_b["salaryVersion"]), ("A", "B"))
            self.assertEqual(salary_a["socialType"], "SOCIAL_UNINSURED")
            self.assertEqual(salary_b["socialType"], "SOCIAL_INSURED")
            self.assertEqual(salary_a["employmentType"], "劳动合同")
            self.assertEqual(salary_b["employmentType"], "劳动合同")
            for salary_file in (
                "11_ONBOARD_SALARY_CONFIRM_A.docx",
                "12_ONBOARD_SALARY_CONFIRM_B.docx",
            ):
                salary_document = Document(output / salary_file)
                self.assertTrue(all(
                    row._tr.trPr is not None and row._tr.trPr.find(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit"
                    ) is not None
                    for table in salary_document.tables for row in table.rows
                ))
                salary_header = salary_document.tables[0].rows[0]
                self.assertIsNotNone(salary_header._tr.trPr.find(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader"
                ))
                signature_index = next(
                    index for index, p in enumerate(salary_document.paragraphs)
                    if "签名：" in p.text
                )
                salary_signature = salary_document.paragraphs[signature_index]
                self.assertTrue(salary_signature.paragraph_format.keep_with_next)
                date_index = next(
                    index for index, p in enumerate(salary_document.paragraphs)
                    if p.text.startswith("日期：")
                )
                self.assertEqual(date_index, signature_index + 1)
                self.assertTrue(all(
                    p.paragraph_format.keep_with_next
                    for p in salary_document.paragraphs[signature_index:date_index]
                ))

            labor_document = Document(output / labor["candidate"]["file"])
            labor_paragraphs = {" ".join(p.text.split()): p for p in labor_document.paragraphs}
            self.assertIn("乙方地址：${employeeAddress}", labor_paragraphs)
            self.assertNotIn("乙方地址：____________________________________", labor_paragraphs)
            for marker in builder.ATTACHMENT_MARKERS:
                self.assertTrue(any(
                    text.startswith("${" + marker + "}") for text in labor_paragraphs
                ))
            self.assertTrue(any(
                "第${contractTermSelection}种形式" in p.text
                for p in labor_document.paragraphs
            ))
            fixed_term = next(
                p for p in labor_document.paragraphs
                if p.text.startswith("${contractTermFixedMark} A、固定期限")
            )
            self.assertTrue(fixed_term.paragraph_format.keep_together)
            self.assertTrue(fixed_term.paragraph_format.keep_with_next)
            self.assertEqual(fixed_term.runs[0].font.name, builder.CHECKBOX_FONT)
            self.assertTrue(any(
                p.text.startswith("${contractTermOpenEndedMark} B、无固定期限")
                for p in labor_document.paragraphs
            ))
            delivery_heading = next(p for p in labor_document.paragraphs if p.text.startswith("第十五条"))
            self.assertTrue(delivery_heading.paragraph_format.keep_with_next)
            for index, paragraph in enumerate(labor_document.paragraphs):
                if not paragraph.text.replace(" ", "").startswith(("附件1", "附件2")):
                    continue
                self.assertTrue(paragraph.paragraph_format.page_break_before)
                previous = labor_document.paragraphs[index - 1]
                self.assertFalse(any(
                    element.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type"
                    ) == "page"
                    for element in previous._p.iter(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br"
                    )
                ))
            hostel_index = next(
                index for index, paragraph in enumerate(labor_document.paragraphs)
                if "为了方便员工" in paragraph.text
            )
            hostel_intro = labor_document.paragraphs[hostel_index]
            self.assertTrue(hostel_intro.text.startswith("${companyName}为了方便员工"))
            self.assertTrue(hostel_intro.paragraph_format.keep_together)
            self.assertEqual(hostel_intro.runs[0].text, "${companyName}")
            self.assertTrue(hostel_intro.runs[0].underline)
            self.assertFalse(any(run.underline for run in hostel_intro.runs[1:]))
            for section in labor_document.sections:
                footer_xml = section.footer._element.xml
                self.assertNotIn("PAGE", footer_xml)
                self.assertNotIn("NUMPAGES", footer_xml)
                self.assertNotIn("txbxContent", footer_xml)
                self.assertNotIn("wp:anchor", footer_xml)
                self.assertTrue(any(
                    paragraph.alignment == 1
                    for paragraph in section.footer.paragraphs
                ))
            self.assertIn(builder.EVIDENCE_REFERENCE_COPY, labor_paragraphs)
            employee_signature_blocks = [
                p for p in labor_document.paragraphs
                if p.text.startswith("乙方：") and "签名：" in p.text
            ]
            self.assertTrue(employee_signature_blocks)
            self.assertTrue(all(
                p.paragraph_format.keep_together for p in employee_signature_blocks
            ))
            update_fields = labor_document.settings._element.find(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}updateFields"
            )
            self.assertIsNotNone(update_fields)
            self.assertEqual(
                update_fields.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"),
                "true",
            )
            duty_table = next(
                table for table in labor_document.tables
                if "员工岗位职责说明书" in table.rows[0].cells[0].text
            )
            regular_rows = [
                row for row in duty_table.rows
                if row.cells[0].text.strip() != "备注"
            ]
            self.assertTrue(all(
                row._tr.trPr is not None and row._tr.trPr.find(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit"
                ) is not None
                for row in regular_rows
            ))
            self.assertTrue(all(
                row._tr.trPr is not None and row._tr.trPr.find(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader"
                ) is not None
                for row in duty_table.rows[:2]
            ))

            service_manifest = by_file["13_ONBOARD_SERVICE_CONTRACT.docx"]
            service_document = Document(output / service_manifest["candidate"]["file"])
            service_text = document_text(output / service_manifest["candidate"]["file"])
            self.assertIn("乙方地址：${employeeAddress}", service_text)
            self.assertIn("甲方（提供方）：${companyName}。为了方便员工", service_text)
            self.assertIn("《劳务合同书签收单》", service_text)
            self.assertNotIn("《薪酬结构确认书》", service_text)
            for marker in builder.SERVICE_PERSON_MARKERS + builder.SERVICE_ATTACHMENT_MARKERS:
                self.assertIn("${" + marker + "}", service_text)
            self.assertTrue(any(
                p.text.startswith("${attachmentServiceReceiptMark} 4")
                for p in service_document.paragraphs
            ))

            confidential = by_file["15_ONBOARD_CONFIDENTIAL_NONCOMPETE.docx"]
            self.assertNotIn("employmentType", confidential)
            self.assertEqual(confidential["postLevelScope"], "7级及以上")
            self.assertIn("劳动合同和劳务合同", confidential["reviewNotes"][0])
            confidential_document = Document(
                output / confidential["candidate"]["file"]
            )
            self.assertTrue(all(
                p.paragraph_format.keep_together
                for p in confidential_document.paragraphs if p.text.strip()
            ))

            minor = by_file["16_ONBOARD_MINOR_NONSTUDENT_DECLARATION.docx"]
            self.assertEqual(minor["supportStatus"], "supported")
            self.assertEqual(minor["activationStatus"], "disabled")
            self.assertEqual(minor["status"], "1")
            self.assertEqual(minor["blockedReasons"], [])
            self.assertEqual(minor["ageScope"], "16-18")
            self.assertEqual(minor["studentStatusScope"], "NON_STUDENT")
            minor_text = document_text(output / minor["candidate"]["file"])
            self.assertIn("本人自${incomeStartYearMonth}起", minor_text)
            self.assertNotIn("${workStartDate}", minor_text)

            commitment = Document(output / "05_ONBOARD_COMMITMENT.docx")
            identity = next(p for p in commitment.paragraphs if "${employeeName}" in p.text)
            self.assertTrue(identity.runs[0].bold)
            with ZipFile(output / "05_ONBOARD_COMMITMENT.docx", "r") as archive:
                custom_properties = archive.read("docProps/custom.xml")
            self.assertEqual(custom_properties, builder.EMPTY_CUSTOM_PROPERTIES_XML)
            self.assertNotIn(b"KSOTemplateDocerSaveRecord", custom_properties)
            self.assertNotIn(b"userId", custom_properties)
            self.assertNotIn(b"ICV", custom_properties)

            first_hashes = {path.name: sha256(path) for path in output.glob("*.docx")}
            first_manifest_hash = sha256(manifest)
            builder.build(source, output, manifest)
            self.assertEqual(
                first_hashes,
                {path.name: sha256(path) for path in output.glob("*.docx")},
            )
            self.assertEqual(first_manifest_hash, sha256(manifest))

            release_dir = root / "release" / "templates"
            qa_dir = root / "release" / "qa"
            release_manifest = root / "release" / "manifest-v6.json"
            self.assertEqual(
                v6_builder.build(
                    output,
                    manifest,
                    release_dir,
                    qa_dir,
                    release_manifest,
                ),
                release_manifest.resolve(),
            )
            release_payload = json.loads(release_manifest.read_text(encoding="utf-8"))
            self.assertFalse(release_payload["productionMutation"])
            self.assertFalse(release_payload["registrationPayloadGenerated"])
            self.assertEqual(release_payload["qaProfiles"], ["normal", "database-max"])
            self.assertEqual(len(release_payload["artifacts"]), 9)
            self.assertEqual(len(release_payload["qaVariants"]), 18)
            self.assertTrue(release_payload["releaseBlockers"])
            self.assertEqual(
                release_payload["pageNumberPolicy"]["mode"],
                "PDF_LAYER_DYNAMIC",
            )
            self.assertEqual(
                release_payload["pathBase"], "manifest-directory"
            )
            portable_paths = [
                release_payload["sourceCandidateManifest"],
                *(
                    value
                    for item in release_payload["artifacts"]
                    for value in (item["file"], item["source"])
                ),
                *(item["file"] for item in release_payload["qaVariants"]),
            ]
            self.assertTrue(all(not Path(value).is_absolute()
                                for value in portable_paths))
            self.assertTrue(all(
                item["signaturePlacementMode"] == "INLINE_ORIGINAL_POSITION"
                and item["evidencePage"]["count"] == 1
                and not item["evidencePage"]["docxAddsBlankPage"]
                for item in release_payload["artifacts"]
            ))
            labor_normal = next(
                (release_manifest.parent / item["file"]).resolve()
                for item in release_payload["qaVariants"]
                if item["profile"] == "normal"
                and item["templateType"] == "ONBOARD_LABOR_CONTRACT"
            )
            labor_normal_text = document_text(labor_normal)
            self.assertIn("☑ A、固定期限", labor_normal_text)
            self.assertNotIn("${", labor_normal_text)
            self.assertEqual(labor_normal_text.count("☑"), 5)
            with self.assertRaisesRegex(RuntimeError, "refusing to overwrite"):
                v6_builder.build(
                    output,
                    manifest,
                    release_dir,
                    qa_dir,
                    release_manifest,
                )

    def test_contract_attachment_count_is_guarded(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source"
            write_fixtures(source)
            labor_path = source / "2-4劳动合同.docx"
            labor = Document(labor_path)
            matches = [p for p in labor.paragraphs if "乙方签字：" in p.text]
            matches[-1].text = "签署位被意外删除"
            labor.save(labor_path)

            with self.assertRaisesRegex(ValueError, "expected 2 occurrences"):
                builder.build(source, root / "output", root / "candidate-manifest.json")


if __name__ == "__main__":
    unittest.main()
