#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path
from typing import Dict, List, Tuple

import openpyxl

import erp_employee_importer as importer


REGION_SHEETS = {"北京区域", "浙江一区", "浙江二三区", "江苏区域", "上海区域"}


def load_original_rows(excel_path: str) -> Tuple[Dict[Tuple[str, str], dict], Dict[str, str], importer.ImportPlan]:
    xlrd = importer.load_xlrd()
    book = xlrd.open_workbook(excel_path, formatting_info=True)
    rows_by_key: Dict[Tuple[str, str], dict] = {}
    for sheet_name in ["名田", *REGION_SHEETS]:
        rows = importer.parse_sheet(book, sheet_name, fill_down_departments=(sheet_name == "名田"))
        for row in rows:
            rows_by_key[(sheet_name, str(row.get("_row", "")))] = row
    sheet1_projects = importer.sheet1_project_map(importer.parse_sheet1(book))
    plan = importer.build_plan_from_workbook(excel_path)
    return rows_by_key, sheet1_projects, plan


def path_for_missing_phone_row(row: dict, sheet_name: str, sheet1_projects: Dict[str, str]) -> List[Tuple[str, ...]]:
    if sheet_name == "名田":
        return [importer.mitian_project_path(row, sheet1_projects)]
    if sheet_name in REGION_SHEETS:
        paths = importer.region_project_paths(row, sheet_name)
        name = importer.clean_name(row.get("姓名", ""))
        project = importer.normalize_project_name(row.get("4级部门", ""))
        if project == "上海瑞吉&麒麟餐厅" and name in {"侯梦璃月", "胡熠熠"}:
            return [
                ("金英灵韵", "上海区域", "上海区域运营", "上海瑞吉"),
                ("金英灵韵", "上海区域", "上海区域运营", "麒麟餐厅"),
            ]
        return paths
    return []


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--employee-excel", default=importer.DEFAULT_EXCEL_PATH)
    parser.add_argument("--missing-excel", default="无手机号员工明细_含负责人.xlsx")
    args = parser.parse_args()

    rows_by_key, sheet1_projects, plan = load_original_rows(args.employee_excel)
    wb = openpyxl.load_workbook(args.missing_excel, data_only=True)
    ws = wb["无手机号明细"]
    headers = [cell.value for cell in ws[1]]
    for raw_row in ws.iter_rows(min_row=2, values_only=True):
        data = dict(zip(headers, raw_row))
        sheet_name = str(data.get("来源工作表") or "").strip()
        row_no = str(data.get("原行号") or "").strip()
        name = importer.clean_name(data.get("姓名", ""))
        if not name:
            continue
        row = rows_by_key.get((sheet_name, row_no), {})
        paths = [p for p in path_for_missing_phone_row(row, sheet_name, sheet1_projects) if p]
        if not paths:
            leader = importer.clean_name(data.get("负责人（整理）", ""))
            leader_user = plan.users_by_phone.get(next((phone for phone, user in plan.users_by_phone.items() if user.name == leader), ""), None)
            if leader_user:
                paths = [leader_user.main_dept_path]
        print("\t".join([
            sheet_name,
            row_no,
            name,
            importer.clean_name(data.get("负责人（整理）", "")),
            " || ".join(" / ".join(path) for path in paths),
        ]))


if __name__ == "__main__":
    main()
