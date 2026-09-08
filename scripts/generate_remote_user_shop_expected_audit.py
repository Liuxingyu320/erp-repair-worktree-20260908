#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path
from typing import Iterable, List, Tuple

import openpyxl

import erp_employee_importer as importer
from build_missing_phone_reconcile_plan import load_original_rows, path_for_missing_phone_row


ExpectedRow = Tuple[str, str, str, str, str, str]


def sql_quote(value: str) -> str:
    return "'" + (value or "").replace("\\", "\\\\").replace("'", "''") + "'"


def values_rows(rows: Iterable[ExpectedRow]) -> str:
    return ",\n".join(
        "(" + ", ".join(sql_quote(item) for item in row) + ")"
        for row in rows
    )


def full_path(path: Tuple[str, ...]) -> str:
    return "金英灵韵集团 / " + " / ".join(path)


def main_expected_rows(plan: importer.ImportPlan) -> List[ExpectedRow]:
    rows: List[ExpectedRow] = []
    for phone, user in sorted(plan.users_by_phone.items()):
        for path in sorted(user.project_paths, key=lambda item: (item != user.main_dept_path, item)):
            rows.append(
                (
                    "phone",
                    phone,
                    user.name,
                    full_path(path),
                    "Y" if path == user.main_dept_path else "N",
                    "员工Excel",
                )
            )
    return rows


def missing_phone_expected_rows(employee_excel: Path, missing_excel: Path, plan: importer.ImportPlan) -> List[ExpectedRow]:
    rows_by_key, sheet1_projects, _ = load_original_rows(str(employee_excel))
    wb = openpyxl.load_workbook(missing_excel, data_only=True)
    ws = wb["无手机号明细"]
    headers = [cell.value for cell in ws[1]]
    result: List[ExpectedRow] = []
    for raw_row in ws.iter_rows(min_row=2, values_only=True):
        data = dict(zip(headers, raw_row))
        name = importer.clean_name(data.get("姓名", ""))
        if not name:
            continue
        sheet_name = str(data.get("来源工作表") or "").strip()
        row_no = str(data.get("原行号") or "").strip()
        source_row = rows_by_key.get((sheet_name, row_no), {})
        paths = [path for path in path_for_missing_phone_row(source_row, sheet_name, sheet1_projects) if path]
        if not paths:
            leader = importer.clean_name(data.get("负责人（整理）", ""))
            leader_user = next((user for user in plan.users_by_phone.values() if user.name == leader), None)
            if leader_user:
                paths = [leader_user.main_dept_path]
        for index, path in enumerate(paths):
            result.append(
                (
                    "missing_name",
                    name,
                    name,
                    full_path(path),
                    "Y" if index == 0 else "N",
                    "无手机号员工明细_含负责人.xlsx",
                )
            )
    return result


def build_expected_rows(employee_excel: Path, missing_excel: Path) -> List[ExpectedRow]:
    plan = importer.build_plan_from_workbook(str(employee_excel))
    rows = main_expected_rows(plan)
    rows.extend(missing_phone_expected_rows(employee_excel, missing_excel, plan))
    return rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--employee-excel", default=importer.DEFAULT_EXCEL_PATH)
    parser.add_argument("--missing-excel", default="无手机号员工明细_含负责人.xlsx")
    args = parser.parse_args()

    rows = build_expected_rows(Path(args.employee_excel), Path(args.missing_excel))
    if not rows:
        raise SystemExit("no expected rows generated")

    print("#!/usr/bin/env bash")
    print("set -euo pipefail")
    print('DB="BossERP_NEW"')
    print('MYSQL_ROOT_PASS="$(cat /root/.erp-mysql-root-pass)"')
    print('mysql --protocol=tcp -uroot -p"${MYSQL_ROOT_PASS}" "${DB}" --batch --raw <<\'SQL\'')
    print("CREATE TEMPORARY TABLE expected_user_shop_scope (")
    print("  match_type varchar(32) not null,")
    print("  match_key varchar(128) not null,")
    print("  expected_name varchar(128) not null,")
    print("  expected_full_path varchar(1000) not null,")
    print("  expected_default char(1) not null,")
    print("  source_name varchar(128) not null")
    print(");")
    print("INSERT INTO expected_user_shop_scope(match_type, match_key, expected_name, expected_full_path, expected_default, source_name) VALUES")
    print(values_rows(rows) + ";")
    print(r"""
CREATE TEMPORARY TABLE expected_user_shop_scope_lookup AS
SELECT * FROM expected_user_shop_scope;

CREATE TEMPORARY TABLE active_dept_paths AS
WITH RECURSIVE dept_tree AS (
    SELECT dept_id, parent_id, dept_name, CAST(dept_name AS CHAR(1000)) AS full_path
    FROM sys_dept
    WHERE dept_id = 100 AND del_flag = '0'
    UNION ALL
    SELECT d.dept_id, d.parent_id, d.dept_name, CONCAT(t.full_path, ' / ', d.dept_name) AS full_path
    FROM sys_dept d
    JOIN dept_tree t ON t.dept_id = d.parent_id
    WHERE d.del_flag = '0'
)
SELECT dept_id, full_path
FROM dept_tree;

CREATE TEMPORARY TABLE actual_user_shop_scope AS
SELECT
    CASE WHEN u.create_by = 'missing_phone_reseed' THEN 'missing_name' ELSE 'phone' END AS match_type,
    CASE WHEN u.create_by = 'missing_phone_reseed' THEN u.nick_name ELSE COALESCE(NULLIF(u.phonenumber, ''), u.user_name) END AS match_key,
    u.user_id,
    u.nick_name,
    p.full_path,
    us.is_default,
    u.create_by
FROM sys_user u
JOIN sys_user_shop us ON us.user_id = u.user_id
JOIN active_dept_paths p ON p.dept_id = us.dept_id
WHERE u.del_flag = '0'
  AND u.create_by IN ('employee_xls_import', 'missing_phone_reseed');

select 'shop_scope_counts' as section;
select count(*) as expected_rows,
       sum(source_name='员工Excel') as expected_employee_rows,
       sum(source_name='无手机号员工明细_含负责人.xlsx') as expected_missing_phone_rows
from expected_user_shop_scope;

select 'actual_counts' as section;
select count(*) as actual_rows,
       sum(create_by='employee_xls_import') as actual_employee_rows,
       sum(create_by='missing_phone_reseed') as actual_missing_phone_rows
from actual_user_shop_scope;

select 'missing_expected_shop_scope' as section;
select e.source_name, e.match_type, e.match_key, e.expected_name, e.expected_full_path, e.expected_default
from expected_user_shop_scope e
where not exists (
    select 1
    from actual_user_shop_scope a
    where a.match_type = e.match_type
      and a.match_key = e.match_key
      and a.full_path = e.expected_full_path
      and a.is_default = e.expected_default
)
order by e.source_name, e.expected_name, e.expected_full_path;

select 'extra_actual_shop_scope' as section;
select a.create_by, a.match_type, a.match_key, a.user_id, a.nick_name, a.full_path, a.is_default
from actual_user_shop_scope a
where not exists (
    select 1
    from expected_user_shop_scope_lookup e
    where e.match_type = a.match_type
      and e.match_key = a.match_key
      and e.expected_full_path = a.full_path
      and e.expected_default = a.is_default
)
order by a.create_by, a.nick_name, a.full_path;

select 'default_count_errors' as section;
select a.create_by, a.match_type, a.match_key, max(a.nick_name) as nick_name,
       count(*) as shop_count,
       sum(a.is_default='Y') as default_count,
       group_concat(concat(a.full_path, ':', a.is_default) order by a.is_default desc, a.full_path separator ' | ') as shops
from actual_user_shop_scope a
group by a.create_by, a.match_type, a.match_key
having sum(a.is_default='Y') <> 1;

select 'invalid_active_user_shop_depts' as section;
select count(*) as invalid_count
from sys_user_shop us
join sys_user u on u.user_id = us.user_id and u.del_flag = '0'
left join sys_dept d on d.dept_id = us.dept_id and d.del_flag = '0'
where d.dept_id is null;
SQL
""")


if __name__ == "__main__":
    main()
