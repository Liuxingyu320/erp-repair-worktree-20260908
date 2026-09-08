from __future__ import annotations

import hashlib
import json
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook


SOURCE_DIR = Path("/Users/liuxingyu/Desktop/入离调转20260702/入离调转")
OUTPUT_DIR = Path("/Users/liuxingyu/Desktop/入离调转20260702/签约系统模板")
RUNTIME_DIR = Path("/Users/liuxingyu/Desktop/备份/ERP-NEW/uploadPath/sign-template/onboard-20260702")
PUBLIC_PREFIX = "/profile/sign-template/onboard-20260702"
REPORT_PATH = OUTPUT_DIR / "占位符校验报告.md"
PAYLOAD_PATH = OUTPUT_DIR / "sign_template_registration_payload.json"


@dataclass(frozen=True)
class TemplateSpec:
    source: str
    runtime_name: str
    template_type: str
    template_name: str
    required: tuple[str, ...]
    sign_required: bool
    sort_order: int
    employment_type: str = ""
    social_type: str = ""
    post_level_scope: str = ""
    salary_version: str = ""
    editor: Callable[[Path], None] | None = field(default=None, compare=False)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    for i in range(len(paragraph.runs) - 1, -1, -1):
        paragraph._p.remove(paragraph.runs[i]._r)
    paragraph.add_run(text)


def replace_first_paragraph(doc: DocxDocument, needle: str, text: str) -> None:
    for paragraph in doc.paragraphs:
        if needle in paragraph.text:
            set_paragraph_text(paragraph, text)
            return
    raise RuntimeError(f"paragraph not found: {needle}")


def replace_all_paragraphs(doc: DocxDocument, needle: str, text_factory: Callable[[str], str]) -> None:
    matched = False
    for paragraph in doc.paragraphs:
        if needle in paragraph.text:
            set_paragraph_text(paragraph, text_factory(paragraph.text))
            matched = True
    if not matched:
        raise RuntimeError(f"paragraph not found: {needle}")


def append_after_first_paragraph(doc: DocxDocument, needle: str, text: str) -> None:
    for paragraph in doc.paragraphs:
        if needle in paragraph.text:
            new_p = OxmlElement("w:p")
            paragraph._p.addnext(new_p)
            inserted = Paragraph(new_p, paragraph._parent)
            inserted.add_run(text)
            return
    raise RuntimeError(f"paragraph not found: {needle}")


def set_cell(table, row: int, col: int, text: str) -> None:
    table.cell(row, col).text = text


def edit_archive(path: Path) -> None:
    doc = Document(path)
    replace_first_paragraph(doc, "入职人员：", "入职人员：${employeeName}               入职时间：${entryDate}")
    replace_first_paragraph(doc, "入职部门：", "入职部门：${employeeDeptName}               入职岗位：${postName}")
    doc.save(path)


def edit_application(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook.active
    sheet["C3"] = "${employeeName}"
    sheet["G3"] = "${employeeIdCard}"
    sheet["C7"] = "${employeePhone}"
    sheet["H2"] = "${postName}"
    sheet["K2"] = "${entryDate}"
    sheet["C33"] = "${employeeName}"
    sheet["E33"] = "${employeePhone}"
    sheet["C34"] = "${entryDate}"
    sheet["E34"] = "${postName}"
    workbook.save(path)


def edit_background(path: Path) -> None:
    doc = Document(path)
    table = doc.tables[0]
    set_cell(table, 0, 1, "${employeeName}")
    set_cell(table, 0, 4, "${postName}")
    doc.save(path)


def edit_offer(path: Path) -> None:
    doc = Document(path)
    replace_first_paragraph(doc, "尊敬的", "尊敬的 ${employeeName} 先生/女士：")
    replace_first_paragraph(doc, "______年", "${signDate}")
    table = doc.tables[0]
    set_cell(table, 0, 1, "${employeeName}")
    set_cell(table, 1, 1, "${employeeIdCard}")
    set_cell(table, 1, 3, "${employeePhone}")
    set_cell(table, 2, 1, "${employeeAddress}")
    table = doc.tables[1]
    set_cell(table, 0, 1, "${postName}")
    set_cell(table, 1, 1, "${employeeDeptName}")
    set_cell(table, 2, 1, "${entryDate}")
    table = doc.tables[2]
    set_cell(table, 1, 1, "${baseSalary}")
    set_cell(table, 5, 1, "${salaryTotal}")
    doc.save(path)


def edit_commitment(path: Path) -> None:
    doc = Document(path)
    replace_first_paragraph(
        doc,
        "本人（姓名：",
        "本人（姓名：${employeeName}，身份证号：${employeeIdCard}）郑重承诺：",
    )
    replace_first_paragraph(doc, "承诺人签字", "承诺人签字（并按手印）：${employeeName}")
    replace_first_paragraph(doc, "日期：", "日期：${signDate}")
    doc.save(path)


def edit_post_duty(path: Path) -> None:
    doc = Document(path)
    replace_first_paragraph(
        doc,
        "员工确认：",
        "员工确认：本人 ${employeeName} 已找到并确认自己的岗位 ${postName}（职级 ${postLevel}），并理解、同意履行上述职责。",
    )
    replace_first_paragraph(doc, "员工签字", "员工签字（并按手印）：${employeeName}")
    replace_first_paragraph(doc, "日期：", "日期：${signDate}")
    doc.save(path)


def edit_labor_contract(path: Path) -> None:
    doc = Document(path)
    replace_first_paragraph(doc, "乙方（劳 动 者）", "乙方（劳 动 者）：${employeeName}")
    replace_first_paragraph(doc, "身 份 证 号", "身 份 证 号  码：${employeeIdCard}")
    replace_first_paragraph(doc, "签   订  日", "签   订  日  期：${signDate}")
    replace_first_paragraph(doc, "姓名：", "姓名：${employeeName} 公民身份号码：${employeeIdCard}")
    replace_all_paragraphs(doc, "家庭住址：", lambda _: "家庭住址：${employeeAddress}")
    replace_all_paragraphs(doc, "联系电话：", lambda _: "联系电话：${employeePhone}")
    replace_first_paragraph(
        doc,
        "A、固定期限",
        "A、固定期限：自${contractStartDate}起至${contractEndDate}止，其中试用期从${probationStartDate}至${probationEndDate}止。",
    )
    replace_first_paragraph(
        doc,
        "3.1 乙方服从甲方工作安排",
        "3.1 乙方服从甲方工作安排，从事${postName}工作岗位。具体任务、职责及工作标准详见甲方的《岗位职责说明书》、操作规范等文件，以及甲方管理人员的安排和要求。",
    )
    replace_first_paragraph(doc, "甲方（盖章）", "甲方（盖章）：____________       乙方（签字并按手印）：${employeeName}")
    replace_first_paragraph(doc, "授权代表（签字）", "授权代表（签字）：________       签订日期：${signDate}")
    replace_first_paragraph(doc, "乙方签字：", "乙方签字：${employeeName} 日期：${signDate}")
    doc.save(path)


def edit_handbook_receipt(path: Path) -> None:
    doc = Document(path)
    replace_first_paragraph(
        doc,
        "本人（姓名：",
        "本人（姓名：${employeeName}，身份证号：${employeeIdCard}）确认：",
    )
    replace_first_paragraph(doc, "员工签字", "员工签字（并按手印）：${employeeName}")
    replace_first_paragraph(doc, "日期：", "日期：${signDate}")
    doc.save(path)


def edit_salary_confirm(path: Path) -> None:
    doc = Document(path)
    replace_first_paragraph(
        doc,
        "本人（姓名：",
        "本人（姓名：${employeeName}（手印），身份证号：${employeeIdCard}），系贵司员工。现就本人薪酬结构事宜，经与公司平等协商，确认如下：",
    )
    replace_first_paragraph(doc, "本人月综合工资标准", "本人月综合工资标准为人民币${salaryTotal}元。")
    table = doc.tables[0]
    set_cell(table, 1, 1, "${baseSalary}")
    set_cell(table, 2, 1, "${postSalary}")
    set_cell(table, 3, 1, "${fieldAllowance}")
    if len(table.rows) > 4:
        set_cell(table, 4, 1, "${performanceSalary}")
    if len(table.rows) > 5:
        set_cell(table, 5, 1, "${salaryTotal}")
    replace_first_paragraph(doc, "员工签字", "员工签字（并按手印）：${employeeName}")
    replace_first_paragraph(doc, "日期：", "日期：${signDate}")
    doc.save(path)


def edit_service_contract(path: Path) -> None:
    doc = Document(path)
    replace_first_paragraph(doc, "乙方（劳务人员）", "乙方（劳务人员）：${employeeName}")
    replace_first_paragraph(doc, "身 份 证 号", "身 份 证 号  码：${employeeIdCard}")
    replace_first_paragraph(doc, "签   订  日", "签   订  日  期：${signDate}")
    replace_first_paragraph(doc, "姓名：", "姓名：${employeeName} 公民身份号码：${employeeIdCard}")
    replace_all_paragraphs(doc, "家庭住址：", lambda _: "家庭住址：${employeeAddress}")
    replace_all_paragraphs(doc, "联系电话：", lambda _: "联系电话：${employeePhone}")
    replace_first_paragraph(doc, "乙方为（", "乙方为（${servicePersonType}），不具备与甲方建立法定劳动关系的主体资格。")
    replace_first_paragraph(doc, "1.1 本协议期限", "1.1 本协议期限自${contractStartDate}起至${contractEndDate}止。")
    replace_first_paragraph(
        doc,
        "3.1 乙方服从甲方工作安排",
        "3.1 乙方服从甲方工作安排，从事${postName}工作岗位。具体任务、职责及工作标准详见甲方的相关文件及管理人员的安排和要求。",
    )
    replace_first_paragraph(doc, "5.1 乙方满勤底薪", "5.1 乙方满勤底薪为${baseSalary}元/月。底薪已包含固定休息日加班的全部补偿。")
    replace_first_paragraph(doc, "甲方（盖章）", "甲方（盖章）：____________       乙方（签字并按手印）：${employeeName}")
    replace_first_paragraph(doc, "授权代表（签字）", "授权代表（签字）：________       签订日期：${signDate}")
    replace_first_paragraph(doc, "乙方签字：", "乙方签字：${employeeName} 日期：${signDate}")
    doc.save(path)


def edit_service_receipt(path: Path) -> None:
    doc = Document(path)
    replace_first_paragraph(
        doc,
        "本人（姓名：",
        "本人（姓名：${employeeName}，身份证号：${employeeIdCard}），系舟山茗汇文化传播有限公司（${servicePersonType}）。",
    )
    replace_first_paragraph(doc, "本人劳务报酬标准", "本人劳务报酬标准为每月人民币${salaryTotal}元。具体构成为：")
    replace_first_paragraph(doc, "甲方已为本人购买", "2、甲方已为本人购买（${insuranceType}），用于覆盖本人在提供劳务期间可能发生的人身意外风险。")
    table = doc.tables[0]
    set_cell(table, 1, 1, "${baseSalary}")
    set_cell(table, 2, 1, "${postSalary}")
    set_cell(table, 3, 1, "${fieldAllowance}")
    if len(table.rows) > 4:
        set_cell(table, 4, 1, "${salaryTotal}")
    replace_first_paragraph(doc, "甲方代表签字", "甲方代表签字：_____________     乙方签字（并按手印）：${employeeName}")
    replace_first_paragraph(doc, "日期：", "日期：${signDate}      日期：${signDate}")
    doc.save(path)


def edit_confidential(path: Path) -> None:
    doc = Document(path)
    replace_first_paragraph(doc, "乙方（员工）", "乙方（员工）：${employeeName}")
    replace_first_paragraph(doc, "身份证号：", "身份证号：${employeeIdCard}")
    replace_first_paragraph(doc, "岗位：", "岗位：${postName} 职级：${postLevel}级")
    append_after_first_paragraph(
        doc,
        "岗位：",
        "联系电话：${employeePhone}    通讯地址：${employeeAddress}    合同起始日：${contractStartDate}",
    )
    replace_first_paragraph(doc, "甲方（盖章）", "甲方（盖章）：____________       乙方（签字并按手印）：${employeeName}")
    replace_first_paragraph(doc, "授权代表（签字）", "授权代表（签字）：________       签订日期：${signDate}")
    doc.save(path)


SPECS: tuple[TemplateSpec, ...] = (
    TemplateSpec("0 员工档案目录.docx", "00_ONBOARD_ARCHIVE_CATALOG.docx", "ONBOARD_ARCHIVE_CATALOG", "员工档案目录", ("employeeName", "employeeDeptName", "entryDate", "postName"), False, 10, editor=edit_archive),
    TemplateSpec("01员工手册.docx", "01_ONBOARD_HANDBOOK.docx", "ONBOARD_HANDBOOK", "员工手册", (), False, 20),
    TemplateSpec("1-1应聘登记表.xlsx", "02_ONBOARD_APPLICATION_FORM.xlsx", "ONBOARD_APPLICATION_FORM", "应聘登记表", ("employeeName", "employeeIdCard", "employeePhone", "entryDate", "postName"), False, 30, editor=edit_application),
    TemplateSpec("1-8背景调查报告.docx", "03_ONBOARD_BACKGROUND_CHECK.docx", "ONBOARD_BACKGROUND_CHECK", "背景调查报告", ("employeeName", "postName"), False, 40, editor=edit_background),
    TemplateSpec("2-1录用通知书.docx", "04_ONBOARD_OFFER_NOTICE.docx", "ONBOARD_OFFER_NOTICE", "录用通知书", ("employeeName", "employeePhone", "employeeDeptName", "postName", "entryDate", "baseSalary", "signDate"), False, 50, editor=edit_offer),
    TemplateSpec("2-3入职承诺书.docx", "05_ONBOARD_COMMITMENT.docx", "ONBOARD_COMMITMENT", "入职承诺书", ("employeeName", "employeeIdCard", "signDate"), True, 60, editor=edit_commitment),
    TemplateSpec("2-4岗位职责确认书（2-4级）.docx", "06_ONBOARD_POST_DUTY_2_4.docx", "ONBOARD_POST_DUTY", "岗位职责确认书（2-4级）", ("employeeName", "postName", "postLevel", "signDate"), True, 70, post_level_scope="2-4", editor=edit_post_duty),
    TemplateSpec("2-4岗位职责确认书（5-6级）.docx", "07_ONBOARD_POST_DUTY_5_6.docx", "ONBOARD_POST_DUTY", "岗位职责确认书（5-6级）", ("employeeName", "postName", "postLevel", "signDate"), True, 80, post_level_scope="5-6", editor=edit_post_duty),
    TemplateSpec("2-4岗位职责确认书（7-8级）.docx", "08_ONBOARD_POST_DUTY_7_8.docx", "ONBOARD_POST_DUTY", "岗位职责确认书（7-8级）", ("employeeName", "postName", "postLevel", "signDate"), True, 90, post_level_scope="7-8", editor=edit_post_duty),
    TemplateSpec("2-5劳动合同.docx", "09_ONBOARD_LABOR_CONTRACT.docx", "ONBOARD_LABOR_CONTRACT", "劳动合同", ("employeeName", "employeeIdCard", "employeePhone", "employeeAddress", "contractStartDate", "contractEndDate", "probationStartDate", "probationEndDate", "postName", "signDate"), True, 100, employment_type="劳动合同", editor=edit_labor_contract),
    TemplateSpec("2-6员工手册签收确认书.docx", "10_ONBOARD_HANDBOOK_RECEIPT.docx", "ONBOARD_HANDBOOK_RECEIPT", "员工手册签收确认书", ("employeeName", "employeeIdCard", "signDate"), True, 110, editor=edit_handbook_receipt),
    TemplateSpec("2-7薪酬结构确认书（A版）.docx", "11_ONBOARD_SALARY_CONFIRM_A.docx", "ONBOARD_SALARY_CONFIRM", "薪酬结构确认书（A版）", ("employeeName", "employeeIdCard", "baseSalary", "postSalary", "fieldAllowance", "salaryTotal", "signDate"), True, 120, salary_version="A", editor=edit_salary_confirm),
    TemplateSpec("2-7薪酬结构确认书（B版）.docx", "12_ONBOARD_SALARY_CONFIRM_B.docx", "ONBOARD_SALARY_CONFIRM", "薪酬结构确认书（B版）", ("employeeName", "employeeIdCard", "baseSalary", "postSalary", "fieldAllowance", "salaryTotal", "signDate"), True, 130, salary_version="B", editor=edit_salary_confirm),
    TemplateSpec("2-8劳务合同.docx", "13_ONBOARD_SERVICE_CONTRACT.docx", "ONBOARD_SERVICE_CONTRACT", "劳务合同", ("employeeName", "employeeIdCard", "employeePhone", "employeeAddress", "servicePersonType", "contractStartDate", "contractEndDate", "postName", "baseSalary", "signDate"), True, 140, employment_type="劳务合同", editor=edit_service_contract),
    TemplateSpec("2-9劳务合同签收单.docx", "14_ONBOARD_SERVICE_RECEIPT.docx", "ONBOARD_SERVICE_RECEIPT", "劳务合同签收单", ("employeeName", "employeeIdCard", "servicePersonType", "baseSalary", "postSalary", "fieldAllowance", "salaryTotal", "insuranceType", "signDate"), True, 150, employment_type="劳务合同", editor=edit_service_receipt),
    TemplateSpec("2-10保密与竞业限制协议（7级及以上）.docx", "15_ONBOARD_CONFIDENTIAL_NONCOMPETE.docx", "ONBOARD_CONFIDENTIAL_NONCOMPETE", "保密与竞业限制协议（7级及以上）", ("employeeName", "employeeIdCard", "employeePhone", "employeeAddress", "postName", "postLevel", "contractStartDate", "signDate"), True, 160, post_level_scope="7-8", editor=edit_confidential),
)


def collect_docx_text(path: Path) -> str:
    doc = Document(path)
    parts: list[str] = []
    parts.extend(paragraph.text for paragraph in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(parts)


def collect_xlsx_text(path: Path) -> str:
    workbook = load_workbook(path, data_only=False)
    values: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    values.append(str(cell.value))
    return "\n".join(values)


def collect_text(path: Path) -> str:
    if path.suffix.lower() == ".xlsx":
        return collect_xlsx_text(path)
    return collect_docx_text(path)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(8192), b""):
            digest.update(chunk)
    return digest.hexdigest()


def convert() -> tuple[list[dict], list[dict]]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUNTIME_DIR.mkdir(parents=True, exist_ok=True)
    audit: list[dict] = []
    payload: list[dict] = []
    for spec in SPECS:
        source = SOURCE_DIR / spec.source
        output = OUTPUT_DIR / spec.source
        runtime = RUNTIME_DIR / spec.runtime_name
        if not source.exists():
            raise FileNotFoundError(source)
        shutil.copy2(source, output)
        if spec.editor:
            spec.editor(output)
        shutil.copy2(output, runtime)
        text = collect_text(output)
        missing = [name for name in spec.required if f"${{{name}}}" not in text]
        audit.append({
            "source": spec.source,
            "templateType": spec.template_type,
            "runtimeName": spec.runtime_name,
            "required": list(spec.required),
            "missing": missing,
            "ok": not missing,
        })
        payload.append({
            "templateType": spec.template_type,
            "templateName": spec.template_name,
            "templateVersion": "20260702-v1",
            "scenario": "onboard",
            "employmentType": spec.employment_type,
            "socialType": spec.social_type,
            "postLevelScope": spec.post_level_scope,
            "salaryVersion": spec.salary_version,
            "fileUrl": f"{PUBLIC_PREFIX}/{spec.runtime_name}",
            "fileName": spec.source,
            "fileSize": runtime.stat().st_size,
            "fileHash": sha256(runtime),
            "requiredPlaceholders": ",".join(spec.required),
            "employeeSignRequired": "Y" if spec.sign_required else "N",
            "sortOrder": spec.sort_order,
            "status": "0",
            "remark": "20260702入离调转签约模板转换导入",
        })
    return audit, payload


def write_report(audit: list[dict], payload: list[dict]) -> None:
    lines = [
        "# 签约系统模板占位符校验报告",
        "",
        f"- 源目录：`{SOURCE_DIR}`",
        f"- 转换目录：`{OUTPUT_DIR}`",
        f"- 运行目录：`{RUNTIME_DIR}`",
        f"- 注册配置：`{PAYLOAD_PATH}`",
        "",
        "| 文件 | 模板类型 | 必需占位符 | 校验 |",
        "| --- | --- | --- | --- |",
    ]
    for item in audit:
        required = ", ".join(f"${{{name}}}" for name in item["required"]) or "无"
        status = "通过" if item["ok"] else "缺少 " + ", ".join(f"${{{name}}}" for name in item["missing"])
        lines.append(f"| {item['source']} | {item['templateType']} | {required} | {status} |")
    lines.extend([
        "",
        "## 待注册模板",
        "",
        "| 顺序 | 名称 | 文件URL | 签署 | 规则 |",
        "| --- | --- | --- | --- | --- |",
    ])
    for item in payload:
        rules = []
        for key in ("employmentType", "postLevelScope", "salaryVersion"):
            if item.get(key):
                rules.append(f"{key}={item[key]}")
        lines.append(
            f"| {item['sortOrder']} | {item['templateName']} | {item['fileUrl']} | {item['employeeSignRequired']} | {', '.join(rules) or '-'} |"
        )
    REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    audit, payload = convert()
    PAYLOAD_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    write_report(audit, payload)
    failed = [item for item in audit if not item["ok"]]
    print(f"converted={len(audit)}")
    print(f"payload={PAYLOAD_PATH}")
    print(f"report={REPORT_PATH}")
    if failed:
        print(json.dumps(failed, ensure_ascii=False, indent=2))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
